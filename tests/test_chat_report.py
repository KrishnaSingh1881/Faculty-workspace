import asyncio
import io
from types import SimpleNamespace
import httpx
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool
from exam_assistant.api import app, current_user, db_session, require_access
from exam_assistant.db import Subject, Paper, Question, Comparison, PairResult, Conversation, ChatMessage


def test_report_is_scoped_to_saved_chat_and_keeps_reasons():
    engine=create_engine('sqlite://',poolclass=StaticPool,connect_args={'check_same_thread':False})
    for model in (Subject,Paper,Question,Comparison,PairResult,Conversation,ChatMessage):
        model.__table__.create(engine)
    with Session(engine) as db:
        db.add(Subject(id=1,code='T',name='Test',tree={}))
        db.add_all([Paper(id=1,subject_id=1,title='Base 2025'),Paper(id=2,subject_id=1,title='Past 2024')])
        db.add_all([Question(id=1,paper_id=1,number='1a',text='Explain the topic',unit=1,confirmed=True),Question(id=2,paper_id=2,number='1a',text='Describe the topic',unit=1,confirmed=True)])
        db.add_all([Comparison(id=1,subject_id=1,title='In this chat',target_ids=[1],reference_ids=[2],status='complete',progress=1,total=1),Comparison(id=2,subject_id=1,title='Another chat',target_ids=[1],reference_ids=[2],status='complete',progress=1,total=1)])
        db.add_all([PairResult(comparison_id=1,target_id=1,reference_id=2,label='B',reason='Step 3: Q1a is partial.',evidence=[{'id':'q1','text':'Explain the topic'},{'id':'block','unit':1,'paper_id':2,'questions':[{'id':2,'number':'1a','text':'Describe the topic','revision':1}]}],source='test',cache_key='x'),PairResult(comparison_id=2,target_id=1,reference_id=2,label='C',reason='Different chat result',evidence=[],source='test',cache_key='x')])
        db.add_all([Conversation(id=1,user_id=7,subject_id=1,title='First'),Conversation(id=2,user_id=7,subject_id=1,title='Second')])
        db.add_all([ChatMessage(conversation_id=1,role='assistant',text='Comparison',payload={'tool_trace':[{'name':'start_comparison','result':{'success':True,'comparison_id':1}}]}),ChatMessage(conversation_id=2,role='assistant',text='Other',payload={'tool_trace':[{'name':'start_comparison','result':{'success':True,'comparison_id':2}}]})])
        db.commit()
    async def access(): current_user.set(SimpleNamespace(id=7,role='teacher',subject_ids=[1]))
    async def session():
        with Session(engine) as db: yield db
    async def run():
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app),base_url='http://test') as client:
            response=await client.get('/api/conversations/1/report?format=xlsx')
            assert response.status_code==200, response.text
            book=load_workbook(io.BytesIO(response.content),data_only=False)
            assert book.sheetnames==['Comparison 1','Reasons']
            assert [book['Comparison 1'].cell(1,col).value for col in range(1,6)]==['Q. No','Base 2025','Q. No','Past 2024','Classification']
            assert book['Comparison 1']['E3'].value=='B'
            assert book['Reasons']['G2'].value=='Step 3: Q1a is partial.'
            assert book['Reasons'].max_row==2
            text=await client.get('/api/conversations/1/report?format=txt')
            assert 'Step 3: Q1a is partial.' in text.text and 'Different chat result' not in text.text
            saved=await client.post('/api/conversations/items',json={'subject_id':1,'kind':'comparison','item_id':1,'text':'Saved comparison'})
            assert saved.status_code==200
            new_id=saved.json()['conversation_id']
            reopened=await client.get(f'/api/conversations/{new_id}')
            assert reopened.json()['messages'][0]['tool_trace'][0]['result']['comparison_id']==1
            assert (await client.get(f'/api/conversations/{new_id}/report?format=xlsx')).status_code==200
    previous=app.dependency_overrides.copy();app.dependency_overrides.update({require_access:access,db_session:session})
    try: asyncio.run(run())
    finally: app.dependency_overrides.clear();app.dependency_overrides.update(previous);engine.dispose()
