import asyncio
import hashlib
import io
import json
import logging
import os
import re
import secrets
import textwrap
from contextlib import asynccontextmanager, suppress
from contextvars import ContextVar
from datetime import datetime, timedelta, timezone
from pathlib import Path
from uuid import uuid4
from zipfile import BadZipFile
from types import SimpleNamespace

import pymupdf
from fastapi import Depends, FastAPI, File, Form, HTTPException, Request, Response as FastAPIResponse, UploadFile
from fastapi.responses import FileResponse, Response, StreamingResponse
from starlette.background import BackgroundTask
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, Field
from sqlalchemy import delete, func, or_, select
from sqlalchemy.orm import Session as DbSession

from .db import AuditEvent, BackgroundJob, ChatMessage, Comparison, ComparisonActivity, Conversation, LoginSession, Material, PairResult, Paper, Question, QuestionAnalysis, ROOT, Session, Subject, SyllabusImport, Template, TreeRevision, User, engine, initialize
from .documents import draft_questions, expand_subquestions, extract_text, suggest_paper_pattern
from .block_comparison import STRATEGY as BLOCK_VERSION, VERSION as BLOCK_PROMPT_VERSION, JUDGE_MODEL as BLOCK_JUDGE_MODEL, block_key, classify_block
from .engine import MODEL, ModelUnavailable, classifier_model_status, classifier_models_ready, judge_pair, model_json, pair_key
from .comparison_agents import run_tool_agent
from .model_runtime import chat_context_tokens, inspect_local_models, model_session
from .chat_tools import SYSTEM as CHAT_SYSTEM, bind_chat_tools
from .chat_routing import compact_trace, select_chat_tools
from .generation import PROMPT_VERSION as GENERATION_PROMPT_VERSION, draft_slot, draft_batch_with_fallback, screen_batch, semantic_topic_hints
from .replacements import replacement_report
from .knowledge import unit_for_topic, validate_tree
from .material_enrichment import PROMPT_VERSION as MATERIAL_PROMPT_VERSION, chunks_for as material_chunks_for, map_material
from .question_analysis import ANALYSIS_PROMPT_VERSION, analysis_cache_key, analyze_questions
from .jobs import JobCancelled, enqueue_job, ensure_not_cancelled, run_worker, set_progress
from .security import hash_password, verify_password
from .backup import create_backup
from .chat_report import build_chat_report
from .syllabus_import import PROMPT_VERSION as SYLLABUS_PROMPT_VERSION, build_tree
from .activity import activity_sink, collect, publish, sse
from .embeddings import index_paper, semantic_search, schedule_missing_index_jobs

Db = DbSession
logger = logging.getLogger(__name__)
chat_tasks: set[asyncio.Task] = set()
chat_runs: dict[str, tuple[int, asyncio.Task]] = {}


@asynccontextmanager
async def lifespan(_: FastAPI):
    initialize()
    await schedule_missing_index_jobs()
    worker = asyncio.create_task(run_worker(JOB_HANDLERS))
    try:
        yield
    finally:
        for task in list(chat_tasks):
            task.cancel()
        if chat_tasks:
            await asyncio.gather(*list(chat_tasks), return_exceptions=True)
        worker.cancel()
        with suppress(asyncio.CancelledError):
            await worker


app = FastAPI(title="Faculty workspace API", lifespan=lifespan)
current_user: ContextVar[User | SimpleNamespace | None] = ContextVar("current_user", default=None)


async def db_session():
    with Session() as db:
        yield db


async def require_access(request: Request, db: DbSession = Depends(db_session)):
    key = os.getenv("API_KEY")
    if key and secrets.compare_digest(request.headers.get("x-api-key", ""), key):
        user = SimpleNamespace(id=0, role="admin", subject_ids=[], display_name="Service administrator")
    elif os.getenv("AUTH_MODE", "local") == "shared":
        token = request.cookies.get("faculty_session", "")
        digest = hashlib.sha256(token.encode()).hexdigest()
        session = db.scalar(select(LoginSession).where(LoginSession.token_hash == digest)) if token else None
        expires = session.expires_at.replace(tzinfo=timezone.utc) if session and session.expires_at.tzinfo is None else session.expires_at if session else None
        user = db.get(User, session.user_id) if session and expires > datetime.now(timezone.utc) else None
        if not user or not user.active:
            raise HTTPException(401, "Sign in to continue")
    else:
        if request.client and request.client.host not in ("127.0.0.1", "::1", "testclient"):
            raise HTTPException(403, "Local mode accepts loopback connections only")
        user = SimpleNamespace(id=0, role="admin", subject_ids=[], display_name="Local administrator")
    current_user.set(user)
    return user


async def require_admin(user=Depends(require_access)):
    if user.role != "admin":
        raise HTTPException(403, "Administrator role required")
    return user


def user_out(user) -> dict:
    return {"id": user.id, "email": getattr(user, "email", None), "display_name": user.display_name, "role": user.role, "subject_ids": user.subject_ids, "active": getattr(user, "active", True)}


class LoginInput(BaseModel):
    email: str
    password: str


@app.post("/api/auth/login")
async def login(body: LoginInput, response: FastAPIResponse, db: DbSession = Depends(db_session)):
    if os.getenv("AUTH_MODE", "local") != "shared":
        raise HTTPException(409, "Sign-in is disabled in local mode")
    user = db.scalar(select(User).where(User.email == body.email.lower().strip()))
    if not user or not user.active or not verify_password(body.password, user.password_hash):
        raise HTTPException(401, "Incorrect email or password")
    token = secrets.token_urlsafe(32)
    db.add(LoginSession(user_id=user.id, token_hash=hashlib.sha256(token.encode()).hexdigest(), expires_at=datetime.now(timezone.utc) + timedelta(hours=12)))
    db.add(AuditEvent(action="login", details={"user_id": user.id}))
    db.commit()
    response.set_cookie("faculty_session", token, httponly=True, secure=os.getenv("COOKIE_SECURE") == "1", samesite="strict", max_age=12 * 3600)
    return user_out(user)


@app.get("/api/auth/me")
async def me(user=Depends(require_access)):
    return user_out(user)


@app.post("/api/auth/logout")
async def logout(request: Request, response: FastAPIResponse, db: DbSession = Depends(db_session)):
    token = request.cookies.get("faculty_session")
    if token:
        session = db.scalar(select(LoginSession).where(LoginSession.token_hash == hashlib.sha256(token.encode()).hexdigest()))
        if session:
            db.delete(session)
            db.commit()
    response.delete_cookie("faculty_session")
    return {"signed_out": True}


class UserInput(BaseModel):
    email: str = Field(min_length=3, max_length=250)
    display_name: str = Field(min_length=1, max_length=150)
    password: str = Field(min_length=12)
    role: str = "teacher"
    subject_ids: list[int] = Field(default_factory=list)


@app.get("/api/users", dependencies=[Depends(require_admin)])
async def list_users(db: DbSession = Depends(db_session)):
    return [user_out(user) for user in db.scalars(select(User).order_by(User.email))]


@app.post("/api/users", dependencies=[Depends(require_admin)])
async def create_user(body: UserInput, db: DbSession = Depends(db_session)):
    if body.role not in ("teacher", "admin"):
        raise HTTPException(422, "Role must be teacher or admin")
    if any(not db.get(Subject, sid) for sid in body.subject_ids):
        raise HTTPException(422, "Unknown subject assignment")
    email = body.email.lower().strip()
    if db.scalar(select(User).where(User.email == email)):
        raise HTTPException(409, "Email already exists")
    user = User(email=email, display_name=body.display_name, role=body.role, password_hash=hash_password(body.password), subject_ids=body.subject_ids)
    db.add(user)
    db.commit()
    return user_out(user)


class ChatInput(BaseModel):
    subject_id: int
    text: str = Field(min_length=1, max_length=4000)
    conversation_id: int | None = None
    comparison_id: int | None = None
    result_id: int | None = None


@app.get("/api/conversations", dependencies=[Depends(require_access)])
async def conversations(subject_id: int, db: Db = Depends(db_session)):
    subject_or_404(db, subject_id)
    user = current_user.get()
    return [{"id": c.id, "title": c.title, "created_at": c.created_at.isoformat()} for c in db.scalars(select(Conversation).where(Conversation.subject_id == subject_id, Conversation.user_id == user.id).order_by(Conversation.id.desc()))]


@app.get("/api/conversations/{cid}", dependencies=[Depends(require_access)])
async def conversation_messages(cid: int, db: Db = Depends(db_session)):
    conversation = db.get(Conversation, cid)
    if not conversation or conversation.user_id != current_user.get().id:
        raise HTTPException(404, "Conversation not found")
    subject_or_404(db, conversation.subject_id)
    return {"id": cid, "subject_id": conversation.subject_id, "messages": [{"id": m.id, "role": m.role, "text": m.text, "card": m.card, **(m.payload or {})} for m in db.scalars(select(ChatMessage).where(ChatMessage.conversation_id == cid).order_by(ChatMessage.id))]}


class ConversationItemInput(BaseModel):
    subject_id: int
    conversation_id: int | None = None
    kind: str
    item_id: int
    text: str = Field(min_length=1, max_length=2000)


@app.post("/api/conversations/items", dependencies=[Depends(require_access)])
async def save_conversation_item(body: ConversationItemInput, db: Db = Depends(db_session)):
    subject_or_404(db, body.subject_id)
    if body.kind == "comparison":
        item = comparison_or_404(db, body.item_id)
        event = {"name": "start_comparison", "result": {"success": True, "comparison_id": item.id}}
        item_subject = item.subject_id
    elif body.kind == "paper":
        item = paper_or_404(db, body.item_id)
        if not item.generated:
            raise HTTPException(422, "Select a generated paper")
        event = {"name": "generate_paper_chat", "result": {"paper_id": item.id}}
        item_subject = item.subject_id
    else:
        raise HTTPException(422, "Unknown chat item")
    if item_subject != body.subject_id:
        raise HTTPException(404, "Item not found in this subject")
    user = current_user.get()
    if body.conversation_id is None:
        conversation = Conversation(user_id=user.id, subject_id=body.subject_id, title=body.text[:80])
        db.add(conversation)
        db.flush()
    else:
        conversation = db.get(Conversation, body.conversation_id)
        if not conversation or conversation.user_id != user.id or conversation.subject_id != body.subject_id:
            raise HTTPException(404, "Conversation not found")
    db.add(ChatMessage(conversation_id=conversation.id, role="assistant", text=body.text, payload={"tool_trace": [event]}))
    db.commit()
    return {"conversation_id": conversation.id, "tool_trace": [event]}


def conversation_comparisons(db: Db, conversation: Conversation):
    ids = []
    for message in db.scalars(select(ChatMessage).where(ChatMessage.conversation_id == conversation.id).order_by(ChatMessage.id)):
        for event in (message.payload or {}).get("tool_trace", []):
            if not isinstance(event, dict):
                continue
            data = event.get("result") or {}
            if not isinstance(data, dict):
                continue
            if event.get("name") in {"start_comparison", "rerun_comparison"} and data.get("success"):
                ids.append(data.get("comparison_id"))
            elif event.get("name") in {"list_comparison_results", "inspect_comparison_result"} and not data.get("error"):
                ids.append(data.get("comparison_id"))
            elif event.get("name") == "generate_paper_chat":
                paper = db.get(Paper, data.get("paper_id"))
                if paper and paper.subject_id == conversation.subject_id:
                    ids.append((paper.generation or {}).get("comparison_id"))
    return [comparison for cid in dict.fromkeys(ids) if isinstance(cid, int)
            and (comparison := db.get(Comparison, cid)) and comparison.subject_id == conversation.subject_id]


@app.get("/api/conversations/{cid}/report", dependencies=[Depends(require_access)])
async def export_conversation_report(cid: int, format: str = "xlsx", db: Db = Depends(db_session)):
    conversation = db.get(Conversation, cid)
    if not conversation or conversation.user_id != current_user.get().id:
        raise HTTPException(404, "Conversation not found")
    subject_or_404(db, conversation.subject_id)
    comparisons = conversation_comparisons(db, conversation)
    if not comparisons:
        raise HTTPException(404, "No saved comparisons in this conversation")
    if not any(db.scalar(select(PairResult.id).where(PairResult.comparison_id == c.id).limit(1)) for c in comparisons):
        raise HTTPException(409, "No completed results in this conversation yet")
    if format == "xlsx":
        data = build_chat_report(comparisons, db, Question, PairResult, report_row_values)
        media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    elif format == "txt":
        sections = [report_text(c, db.scalars(select(PairResult).where(PairResult.comparison_id == c.id).order_by(PairResult.id)).all(), db) for c in comparisons]
        data, media = "\n\n".join(sections).encode(), "text/plain"
    else:
        raise HTTPException(422, "Format must be xlsx or txt")
    return Response(data, media_type=media, headers={"Content-Disposition": f'attachment; filename="conversation-{cid}-report.{format}"'})


@app.delete("/api/conversations/{cid}", dependencies=[Depends(require_access)])
async def delete_conversation(cid: int, db: Db = Depends(db_session)):
    conversation = db.get(Conversation, cid)
    if not conversation or conversation.user_id != current_user.get().id:
        raise HTTPException(404, "Conversation not found")
    subject_or_404(db, conversation.subject_id)
    db.execute(delete(ChatMessage).where(ChatMessage.conversation_id == cid))
    db.delete(conversation)
    db.add(AuditEvent(action="conversation_delete", details={"conversation_id": cid, "by": current_user.get().id}))
    db.commit()
    return {"deleted": cid}


def resolve_delete_target(db: Db, subject_id: int, kind: str, target: str) -> int | None:
    value = target.removeprefix("#").strip()
    if value.isdigit():
        return int(value)
    if kind == "paper":
        matches = [p.id for p in db.scalars(select(Paper).where(Paper.subject_id == subject_id)) if p.title.casefold() == value.casefold()]
    elif kind == "material":
        matches = [m.id for m in db.scalars(select(Material).where(Material.subject_id == subject_id)) if m.filename.casefold() == value.casefold()]
    elif kind == "template":
        matches = [t.id for t in db.scalars(select(Template)) if t.subject_id in (None, subject_id) and t.name.casefold() == value.casefold()]
    elif kind == "syllabus":
        user = current_user.get()
        matches = [item.id for item in db.scalars(select(SyllabusImport)) if item.subject_id in (None, subject_id) and (not user or user.role == "admin" or item.submitted_by == user.id) and item.filename.casefold() == value.casefold()]
    else:
        matches = [q.id for q in db.scalars(select(Question).join(Paper).where(Paper.subject_id == subject_id)) if q.number.casefold() == value.casefold()]
    return matches[0] if len(matches) == 1 else None


@app.post("/api/chat", dependencies=[Depends(require_access)])
async def chat(body: ChatInput, db: Db = Depends(db_session)):
    subject_or_404(db, body.subject_id)
    user = current_user.get()
    if body.conversation_id is None:
        conversation = Conversation(user_id=user.id, subject_id=body.subject_id, title=body.text[:80])
        db.add(conversation)
        db.flush()
    else:
        conversation = db.get(Conversation, body.conversation_id)
        if not conversation or conversation.user_id != user.id or conversation.subject_id != body.subject_id:
            raise HTTPException(404, "Conversation not found")
    text = body.text.strip()
    lower = text.casefold()
    answer, card, view, file_url, tool_trace = "", None, None, None, []
    activity_events = []
    db.commit()
    publish({"agent": "workspace", "event": "conversation", "payload": {"conversation_id": conversation.id}})
    deletion = re.fullmatch(r"(?:please\s+)?(?:permanently\s+)?(?:delete|remove)\s+(?:the\s+)?(paper|question|material|template|syllabus upload)\s+(?:named\s+|called\s+)?(.+?)(?:\s+please)?", text, re.I)
    if deletion:
        kind, target = deletion.group(1).lower(), deletion.group(2)
        kind = "syllabus" if kind == "syllabus upload" else kind
        resolved = resolve_delete_target(db, body.subject_id, kind, target)
        if resolved is None:
            answer = f"I found no unique {kind} named “{target}”. Use its ID or exact name."
        else:
            if kind == "paper":
                await delete_paper(resolved, db)
            elif kind == "question":
                await delete_question(resolved, db)
            elif kind == "material":
                await delete_material(resolved, db)
            elif kind == "syllabus":
                await delete_syllabus_import(resolved, db)
            else:
                await delete_template(resolved, db)
            answer = f"Deleted {kind} {resolved}. The deletion is final and recorded."
    elif MODEL:
        try:
            prior_messages = list(db.scalars(select(ChatMessage).where(ChatMessage.conversation_id == conversation.id).order_by(ChatMessage.id.desc()).limit(12)).all())[::-1]
            prior = [{"role": item.role, "content": item.text[:1000]} for item in prior_messages]
            prior_tool_trace = [event for item in prior_messages if item.role == "assistant" for event in (item.payload or {}).get("tool_trace", []) if isinstance(event, dict)]
            if prior_tool_trace:
                prior.append({"role": "assistant", "content": "Recent recorded Workspace result IDs and labels: " + compact_trace(prior_tool_trace)})
            inventory = await inspect_local_models()
            active_user = current_user.get()
            registry = bind_chat_tools(db, subject_or_404(db, body.subject_id), model_configured=bool(MODEL), comparison_ready=classifier_models_ready() and inventory.get("roles_available", False), prior_tool_trace=prior_tool_trace, selected_result_id=body.result_id, is_admin=bool(active_user and active_user.role == "admin"))
            def router_event(event):
                record = {"agent": "router", "event": event["event"], "payload": event["payload"]}
                publish(record)
                collect(activity_events, record)
            route_context = [{"role": item.role, "content": item.text[:400]} for item in prior_messages[-4:]]
            selected_schemas, selected_families = await select_chat_tools(MODEL, text, route_context, [event.get("name", "") for event in prior_tool_trace], on_event=router_event)
            chosen_names = {item["function"]["name"] for item in selected_schemas}
            registry = {name: handler for name, handler in registry.items() if name in chosen_names}
            selection_event = {"agent": "router", "event": "finished", "payload": {"families": selected_families, "tools_available": len(selected_schemas)}}
            publish(selection_event)
            collect(activity_events, selection_event)
            response_data, tool_trace, _ = await run_tool_agent(
                model=MODEL,
                system=CHAT_SYSTEM,
                payload={"subject": {"id": body.subject_id, "code": subject_or_404(db, body.subject_id).code, "name": subject_or_404(db, body.subject_id).name}, "conversation": prior, "user_request": text},
                registry=registry,
                require_tool_call=False,
                max_calls=8,
                tool_schemas=selected_schemas,
                agent="workspace",
                session_per_request=True,
                num_ctx=chat_context_tokens(),
                on_event=lambda event: collect(activity_events, event),
            )
            answer = response_data.get("answer") if isinstance(response_data.get("answer"), str) else "I checked the Workspace data. See the recorded tool results for details."
            view = response_data.get("view") if response_data.get("view") in {"chat", "papers", "knowledge", "materials", "templates", "reports", "operations"} else None
            for event in tool_trace:
                result = event.get("result") or {}
                if result.get("success") and "file_url" in result:
                    file_url = result["file_url"]
                if event.get("name") == "create_paper" and result.get("success"):
                    view = "papers"
                    answer += f" Created {result['paper']['title']} as a {result['paper']['status']} paper. Attach its file here or ask me to add questions."
                elif event.get("name") == "create_question" and result.get("success"):
                    view = "papers"
                    question = result["question"]
                    answer += f" Added question {question['number']} to {result['paper']}."
                    if not question["confirmed"]:
                        answer += " It needs a valid syllabus topic tag before it can be compared."
                elif event.get("name") == "set_paper_status" and result.get("success"):
                    view = "papers"
                    answer += f" Paper {result['paper_id']} is now {result['status']}."
                elif event.get("name") == "update_paper" and result.get("success"):
                    view = "papers"
                    answer += f" Updated {', '.join(result.get('changed_fields', []))} for {result['title']}."
                elif event.get("name") == "set_paper_archived" and result.get("success"):
                    view = "papers"
                    answer += f" Paper {result['paper_id']} was {'archived' if result['archived'] else 'restored'}."
                elif event.get("name") == "save_template" and result.get("success"):
                    view = "templates"
                    answer += f" Saved template {result['name']} as version {result['version']}."
                if event.get("name") == "rerun_comparison" and result.get("success"):
                    view = "reports"
                    answer += f" Refresh comparison {result['comparison_id']} is queued from #{result['source_comparison_id']} for {result['pairs']} pairs. The earlier run is preserved."
                elif event.get("name") == "cancel_comparison" and result.get("success"):
                    view = "reports"
                    answer += f" Cancellation for {result['title']} is { 'requested' if result['cancel_requested'] else 'stopped' }."
                elif event.get("name") == "resume_comparison" and result.get("success"):
                    view = "reports"
                    answer += f" Resumed {result['title']}; it is queued to continue."
                elif event.get("name") == "start_comparison" and result.get("success"):
                    view = "reports"
                    answer += f" Comparison {result['comparison_id']} is queued for {result['pairs']} same-unit pairs."
                    if result.get("skipped_cross_unit_pairs"):
                        answer += f" I skipped {result['skipped_cross_unit_pairs']} cross-unit pair(s), as required by the syllabus boundary."
                elif event.get("name") == "list_comparison_results" and result.get("results") is not None:
                    view = "reports"
                    answer += f" Found {result['total_matches']} saved result(s) for comparison {result['comparison_id']}."
                elif event.get("name") == "inspect_comparison_result" and result.get("id"):
                    view = "reports"
                elif event.get("name") == "map_material" and result.get("success"):
                    view = "chat"
                    matches = result.get("topic_matches", [])
                    lines = "\n".join(f"{item['topic']} — page {', '.join(map(str, item['pages']))}" for item in matches)
                    answer = (f"Mapped {result['filename']} to {len(matches)} syllabus topics.\n\n{lines or 'No topic matches found.'}"
                              "\n\nSay 'go ahead, enrich' to add these matches to the tree.")
                elif event.get("name") == "apply_material_mapping" and result.get("success"):
                    view = "chat"
                    answer = f"Added {len(result['topics'])} material topic matches to the syllabus tree."
                elif event.get("name") == "generate_questions" and result.get("success"):
                    view = "papers"
                    answer += f" Draft {result['title']} is generating {result['candidate_count']} questions as job {result['job_id']}; redundancy screening will follow."
                elif event.get("name") == "suggest_replacements" and result.get("success"):
                    view = "papers"
                    answer += f" Replacement draft {result['title']} is queued as job {result['job_id']}. The source question is unchanged."
                elif event.get("name") == "replacement_analysis" and result.get("success"):
                    view = "papers"
                elif event.get("name") == "apply_replacement" and result.get("success"):
                    view = "papers"
                    answer += f" Applied candidate {result['candidate_id']} to question {result['source_question_id']}; revision {result['new_revision']} is recorded."
                elif event.get("name") == "update_question" and result.get("success"):
                    view = "papers"
                    answer += f" Question {result['question']['number']} was updated; earlier comparison records now show stale status where applicable."
                elif event.get("name") == "update_syllabus_topic" and result.get("success"):
                    view = "knowledge"
                    answer += f" Updated {', '.join(result.get('changed_fields', []))} for {result['label']} in syllabus version {result['tree_version']}."
                elif event.get("name") == "create_syllabus_topic" and result.get("success"):
                    view = "knowledge"
                    answer += f" Added {result['label']} under {result['parent']} in syllabus version {result['tree_version']}."
                elif event.get("name") == "delete_syllabus_topic" and result.get("success"):
                    view = "knowledge"
                    answer += f" Permanently removed {result['label']} from the syllabus. {result['questions_retagged']} question(s) lost that topic tag and may need retagging."
                elif event.get("name") == "delete_record" and result.get("success"):
                    view = "papers" if result["kind"] in {"paper", "question"} else "materials" if result["kind"] == "material" else "templates"
                    answer += f" Permanently deleted {result['kind']} “{result['title']}”. The deletion is recorded."
                elif event.get("name") == "resolve_selected_result" and result.get("success"):
                    view = "reports"
                    answer += f" Recorded your case-specific {result['label']} ruling. It applies only to this pair and will not be propagated."
                elif event.get("name") == "note_on_selected_result" and result.get("success"):
                    view = "reports"
                    answer += " Recorded your feedback with this comparison case. It does not change the saved result or train the model."
                elif event.get("name") == "link_syllabus_topics" and result.get("success"):
                    view = "knowledge"
                    labels = result.get("topic_labels", result.get("topic_ids", ["the selected topics"]))
                    answer += f" The distinction between {labels[0]} and {labels[1]} was saved in syllabus version {result['tree_version']}."
                elif event.get("name") == "prepare_report" and result.get("success"):
                    file_url = result["file_url"]
                elif event.get("name") == "list_jobs" and result.get("jobs") is not None:
                    view = "operations" if active_user and active_user.role == "admin" else "chat"
                    answer += f" Found {len(result['jobs'])} matching background job(s)."
                elif event.get("name") == "manage_job" and result.get("success"):
                    view = "operations" if active_user and active_user.role == "admin" else "chat"
                    job_state = result.get("status", "unknown")
                    answer += f" Job {result['id']} ({result['kind']}) is {job_state} after the {result['action']} request."
                    if result.get("file_url"):
                        file_url = result["file_url"]
                        answer += " The export is ready to download."
                    elif result.get("artifact_expired"):
                        answer += " Its export file has expired; queue a new export from the comparison."
                elif event.get("name") == "workspace_operations" and result.get("success"):
                    view = "operations"
                    upload_mb = result.get("upload_bytes", 0) / (1024 * 1024)
                    export_mb = result.get("export_bytes", 0) / (1024 * 1024)
                    answer += f" Workspace operations: {result.get('database')} database, {result.get('auth_mode')} access, {upload_mb:.1f} MB of uploads and {export_mb:.1f} MB of exports."
        except Exception as exc:
            answer = f"I couldn't complete that request ({str(exc)[:300]}). You can clarify or retry it here. "
            if re.match(r"^(?:please\s+)?(?:upload|attach|import|add paper)\b", lower):
                answer += "Attach the file in this chat and choose what it contains; the upload is stored immediately."
                card = "new-paper"
    elif "upload" in lower or "add paper" in lower:
        answer, card = "Create a paper and attach its file. The upload is stored immediately.", "new-paper"
    elif any(word in lower for word in ("compare", "repeat", "redundan", "check draft")):
        answer, card = "Choose target and reference questions. Only pairs in the same subject and unit will run.", "compare"
    elif any(word in lower for word in ("export", "download", "report")):
        comparison = db.get(Comparison, body.comparison_id) if body.comparison_id else None
        if comparison:
            subject_or_404(db, comparison.subject_id)
            format_match = re.search(r"\b(md|markdown|txt|text|pdf|xlsx|excel)\b", lower)
            fmt = {"markdown": "md", "text": "txt", "excel": "xlsx"}.get(format_match.group(1), format_match.group(1)) if format_match else None
            if fmt:
                file_url = f"/api/comparisons/{comparison.id}/export?format={fmt}"
                answer = f"Your {fmt.upper()} report is ready to download from the saved comparison."
            else:
                answer, card = "Choose MD, TXT, PDF or XLSX for this comparison.", "export"
        else:
            answer, view = "Select a saved comparison first.", "reports"
    elif any(word in lower for word in ("why", "explain", "evidence")) and body.result_id:
        result = db.get(PairResult, body.result_id)
        if not result:
            raise HTTPException(404, "Result not found")
        comparison_or_404(db, result.comparison_id)
        cited = ", ".join(e.get("id", "?") for e in result.evidence if e.get("cited"))
        answer = f"{result.label}: {result.reason} Recorded evidence: {cited}."
    elif "paper" in lower and any(word in lower for word in ("list", "show", "which")):
        found = db.scalars(select(Paper).where(Paper.subject_id == body.subject_id).order_by(Paper.year.desc(), Paper.id.desc())).all()
        answer = "Papers: " + ("; ".join(f"{p.id}: {p.title} ({p.status})" for p in found) if found else "none uploaded yet")
    elif "bank" in lower and any(word in lower for word in ("search", "find")):
        query = re.sub(r"(?i)\b(search|find|bank|for|questions?)\b", "", text).strip()
        matches = [q for q in db.scalars(select(Question).join(Paper).where(Paper.subject_id == body.subject_id, Paper.status == "bank", Paper.archived.is_(False))) if query.casefold() in q.text.casefold()][:10] if query else []
        answer = "Bank matches: " + ("; ".join(f"{q.id}: {q.text[:120]}" for q in matches) if matches else "none for that phrase")
    elif any(word in lower for word in ("syllabus", "tree", "topic")):
        answer, view = "The active syllabus tree and its interpretation notes are in Knowledge.", "knowledge"
    elif any(word in lower for word in ("template", "pattern")):
        answer, view = "Paper patterns, question slots, comparison scopes and report layouts are in Template studio.", "templates"
    elif any(word in lower for word in ("notes", "slides", "material", "textbook")):
        answer, view = "Teaching materials are stored immediately. Select a topic when you want cited enrichment.", "materials"
    else:
        answer = "Tell me what you want to find, compare, analyze, edit, organize, map, draft, export or delete. I’ll use the Workspace tools and show what I did here."
    db.add(ChatMessage(conversation_id=conversation.id, role="user", text=text))
    db.add(ChatMessage(conversation_id=conversation.id, role="assistant", text=answer, card=card, payload={"view": view, "file_url": file_url, "tool_trace": tool_trace, "activity": activity_events}))
    db.commit()
    return {"conversation_id": conversation.id, "text": answer, "card": card, "view": view, "file_url": file_url, "tool_trace": tool_trace, "activity": activity_events}


@app.post("/api/chat/runs/{run_id}/stop", dependencies=[Depends(require_access)])
async def stop_chat(run_id: str):
    run = chat_runs.get(run_id)
    if not run:
        return {"status": "finished"}
    if run[0] != current_user.get().id:
        raise HTTPException(404, "Active response not found")
    task = run[1]
    if not task.done() and not task.cancelling():
        task.cancel()
    return {"status": "stopping"}


@app.post("/api/chat/stream", dependencies=[Depends(require_access)])
async def stream_chat(body: ChatInput, db: Db = Depends(db_session)):
    subject_or_404(db, body.subject_id)
    if body.conversation_id is not None:
        conversation = db.get(Conversation, body.conversation_id)
        if not conversation or conversation.user_id != current_user.get().id or conversation.subject_id != body.subject_id:
            raise HTTPException(404, "Conversation not found")

    async def events():
        queue = asyncio.Queue()
        connected = True
        run_id = uuid4().hex
        owner_id = current_user.get().id
        saved_cid = body.conversation_id
        recorded = []

        def deliver(event):
            nonlocal saved_cid
            if event.get("event") == "conversation":
                saved_cid = event["payload"]["conversation_id"]
            if event.get("event") not in {"complete", "error"}:
                collect(recorded, event)
            if connected:
                queue.put_nowait(event)

        async def execute():
            token = activity_sink.set(deliver)
            try:
                with Session() as task_db:
                    result = await chat(body, task_db)
                deliver({"event": "complete", "payload": result})
            except asyncio.CancelledError:
                # Explicit stop cancels inference; committed tool actions remain intact.
                try:
                    with Session() as stopped_db:
                        conversation = stopped_db.get(Conversation, saved_cid) if saved_cid else None
                        if not conversation:
                            conversation = Conversation(user_id=owner_id, subject_id=body.subject_id, title=body.text[:80])
                            stopped_db.add(conversation)
                            stopped_db.flush()
                        text = "Response stopped. Any completed workspace actions were kept."
                        collect(recorded, {"agent": "workspace", "event": "stopped", "payload": {}})
                        stopped_db.add(ChatMessage(conversation_id=conversation.id, role="user", text=body.text))
                        stopped_db.add(ChatMessage(conversation_id=conversation.id, role="assistant", text=text, payload={"activity": recorded, "stopped": True}))
                        stopped_db.commit()
                        deliver({"event": "complete", "payload": {"conversation_id": conversation.id, "text": text, "activity": recorded, "stopped": True}})
                except Exception as exc:
                    logger.exception("Could not save stopped chat")
                    deliver({"event": "error", "payload": {"message": "Generation stopped, but the interrupted turn could not be saved. Reopen the conversation to check completed actions."}})
            except Exception as exc:
                logger.exception("Streaming chat failed")
                deliver({"event": "error", "payload": {"message": str(getattr(exc, "detail", exc))[:500]}})
            finally:
                activity_sink.reset(token)

        task = asyncio.create_task(execute())
        # A disconnected browser must not replay already-executed actions. Finish and
        # persist the original turn so reopening the conversation retrieves the result.
        chat_tasks.add(task)
        task.add_done_callback(chat_tasks.discard)
        chat_runs[run_id] = (owner_id, task)
        task.add_done_callback(lambda _: chat_runs.pop(run_id, None))
        try:
            await asyncio.sleep(0)  # Enter the cancellation handler before exposing Stop.
            yield sse("activity", {"agent": "workspace", "event": "started", "payload": {"run_id": run_id}})
            while True:
                try:
                    event = await asyncio.wait_for(queue.get(), timeout=15)
                except TimeoutError:
                    yield ": keepalive\n\n"
                    continue
                kind = event.get("event")
                yield sse(kind if kind in {"complete", "error"} else "activity", event["payload"] if kind in {"complete", "error"} else event)
                if kind in {"complete", "error"}:
                    break
        finally:
            connected = False

    return StreamingResponse(events(), media_type="text/event-stream", headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


@app.post("/api/chat/upload", dependencies=[Depends(require_access)])
async def chat_upload(
    subject_id: int = Form(...),
    kind: str = Form(...),
    file: UploadFile = File(...),
    conversation_id: int | None = Form(None),
    status: str = Form("reference"),
    year: int | None = Form(None),
    message: str = Form(""),
    db: Db = Depends(db_session),
):
    subject_or_404(db, subject_id)
    user = current_user.get()
    if conversation_id is None:
        conversation = Conversation(user_id=user.id, subject_id=subject_id, title=f"Upload: {Path(file.filename or 'file').name}"[:80])
        db.add(conversation)
        db.flush()
    else:
        conversation = db.get(Conversation, conversation_id)
        if not conversation or conversation.user_id != user.id or conversation.subject_id != subject_id:
            raise HTTPException(404, "Conversation not found")
    filename = Path(file.filename or "upload").name
    user_message = message.strip() or f"Upload {filename} as {kind}."
    db.add(ChatMessage(conversation_id=conversation.id, role="user", text=user_message))
    db.commit()
    view, result, answer = "chat", {}, ""
    try:
        if kind == "paper":
            # Chat is the primary paper workflow: extract the questions first so
            # they can be shown immediately, then tag them in the durable queue.
            imported = await bulk_upload_papers([file], subject_id, status, year, db, auto_tag=True)
            outcome = imported["outcomes"][0]
            if outcome["status"] != "imported":
                raise HTTPException(422, outcome.get("error", "Paper upload failed"))
            result = outcome
            drafts = outcome["paper"]["questions"]
            answer = (f"I extracted {len(drafts)} question{'s' if len(drafts) != 1 else ''} from {filename}. "
                      + ("I’m assigning each question to its syllabus unit and topic now; you can review the extracted text below while that runs. "
                       "They’ll become available for comparison as soon as tagging finishes."
                       if outcome.get("tagging_job_id") else
                       "The extracted questions are below. Automatic syllabus tagging could not start because no local tagging model is configured."))
            if not drafts:
                answer = f"I saved {filename}, but couldn’t identify any questions. Check the PDF text extraction or add the questions manually."
        elif kind == "material":
            data = await file.read(25 * 1024 * 1024 + 1)
            await file.seek(0)
            duplicate = db.scalar(select(Material).where(Material.subject_id == subject_id, Material.file_hash == hashlib.sha256(data).hexdigest()))
            if duplicate:
                material = duplicate
                result = {"id": material.id, "subject_id": subject_id, "filename": material.filename, "warnings": material.warnings, "pages": len(material.extracted_text.split("\f")), "proposals": material.proposals, "enrichment": material.enrichment, "duplicate": True}
            else:
                result = await upload_material(subject_id, file, db)
                material = db.get(Material, result["id"])
            pages = material.extracted_text.split("\f") if material else []
            page_samples = "\n\n".join(f"Page {index}: {page.strip()[:700]}" for index, page in enumerate(pages[:3], 1) if page.strip())
            opening = f"This material was already stored as material {material.id}, so I reopened it here." if result.get("duplicate") else f"I extracted {len(pages)} pages or sections from {filename}."
            answer = (f"{opening} Here is a preview:\n\n{page_samples or 'No readable text found.'}\n\n"
                      "Expand the pages below to read the extraction. Ask me to map it across this subject's syllabus; I will show the topic and page matches here. Say 'go ahead, enrich' to add those matches to the tree.")
        elif kind == "syllabus":
            result = await import_syllabus(file, db, subject_id=subject_id)
            view = "knowledge"
            answer = f"Stored {filename}; syllabus extraction and validation are queued as job {result['job_id']}."
        elif kind == "template":
            result = await import_paper_sample(subject_id, file, db)
            view = "templates"
            answer = f"Saved the sample as editable template {result['name']}; syllabus slots and choice rules still need configuration."
        else:
            raise HTTPException(422, "Upload kind must be paper, material, syllabus or template")
    except HTTPException as exc:
        answer = f"Upload failed: {exc.detail}"
        result = {"success": False, "error": str(exc.detail)}
    except Exception:
        logger.exception("Chat upload failed for %s", filename)
        answer = "Upload failed while processing the file. The error was recorded in the application log."
        result = {"success": False, "error": "File processing failed"}
    trace_result = {**result, "paper_id": result["paper"]["id"]} if kind == "paper" and result.get("paper") else result
    if trace_result is not result:
        trace_result.pop("paper")
    trace = [{"id": "upload_1", "name": f"upload_{kind}", "arguments": {"filename": filename, "subject_id": subject_id}, "result": trace_result}]
    db.add(ChatMessage(conversation_id=conversation.id, role="assistant", text=answer, payload={"view": view, "tool_trace": trace, "upload": result if kind == "paper" else None}))
    db.commit()
    return {"conversation_id": conversation.id, "text": answer, "view": view, "tool_trace": trace, "upload": result}


class UserUpdate(BaseModel):
    display_name: str | None = Field(default=None, min_length=1, max_length=150)
    password: str | None = Field(default=None, min_length=12)
    role: str | None = None
    active: bool | None = None
    subject_ids: list[int] | None = None


@app.put("/api/users/{uid}", dependencies=[Depends(require_admin)])
async def update_user(uid: int, body: UserUpdate, db: DbSession = Depends(db_session)):
    user = db.get(User, uid)
    if not user:
        raise HTTPException(404, "User not found")
    if body.role is not None and body.role not in ("teacher", "admin"):
        raise HTTPException(422, "Invalid role")
    if body.subject_ids is not None and any(not db.get(Subject, sid) for sid in body.subject_ids):
        raise HTTPException(422, "Unknown subject assignment")
    if (body.active is False or body.role == "teacher") and user.role == "admin":
        others = db.scalars(select(User).where(User.id != uid, User.role == "admin", User.active == True)).all()
        if not others:
            raise HTTPException(409, "Keep at least one active administrator")
    for field in ("display_name", "role", "active", "subject_ids"):
        value = getattr(body, field)
        if value is not None:
            setattr(user, field, value)
    if body.password:
        user.password_hash = hash_password(body.password)
        for session in db.scalars(select(LoginSession).where(LoginSession.user_id == uid)):
            db.delete(session)
    db.add(AuditEvent(action="user_update", details={"user_id": uid, "by": current_user.get().id}))
    db.commit()
    return user_out(user)


def subject_or_404(db: Db, sid: int) -> Subject:
    subject = db.get(Subject, sid)
    if not subject:
        raise HTTPException(404, "Subject not found")
    user = current_user.get()
    if user and user.role == "teacher" and sid not in user.subject_ids:
        raise HTTPException(403, "Subject is not assigned to this teacher")
    return subject


def paper_or_404(db: Db, pid: int) -> Paper:
    paper = db.get(Paper, pid)
    if not paper:
        raise HTTPException(404, "Paper not found")
    subject_or_404(db, paper.subject_id)
    return paper


def comparison_or_404(db: Db, cid: int) -> Comparison:
    comparison = db.get(Comparison, cid)
    if not comparison:
        raise HTTPException(404, "Comparison not found")
    subject_or_404(db, comparison.subject_id)
    return comparison


def question_out(q: Question) -> dict:
    return {"id": q.id, "paper_id": q.paper_id, "number": q.number, "text": q.text, "marks": q.marks, "structure": q.structure or {"parts": [], "choice_groups": [], "warnings": []}, "unit": q.unit, "topic_ids": q.topic_ids, "confirmed": q.confirmed, "revision": q.revision}


def paper_out(p: Paper) -> dict:
    return {"id": p.id, "subject_id": p.subject_id, "title": p.title, "year": p.year, "status": p.status, "archived": p.archived, "filename": p.filename, "generated": p.generated, "generation": p.generation, "questions": [question_out(q) for q in p.questions]}


def comparison_out(c: Comparison) -> dict:
    return {"id": c.id, "subject_id": c.subject_id, "title": c.title, "target_ids": c.target_ids, "reference_ids": c.reference_ids, "rerun_of_id": c.rerun_of_id, "template_snapshot": c.template_snapshot or {}, "strategy": (c.template_snapshot or {}).get("strategy"), "status": c.status, "progress": c.progress, "total": c.total, "error": c.error}


@app.get("/api/health")
async def health():
    return {"status": "ok", "model_configured": bool(MODEL), "model": MODEL or None, "comparison_models": classifier_model_status()}


@app.get("/api/models/status", dependencies=[Depends(require_access)])
async def model_status():
    return await inspect_local_models()


@app.get("/api/jobs/{job_id}", dependencies=[Depends(require_access)])
async def job_detail(job_id: int, db: Db = Depends(db_session)):
    job = db.get(BackgroundJob, job_id)
    if not job:
        raise HTTPException(404, "Job not found")
    payload = job.payload or {}
    if job.kind == "comparison":
        comparison_or_404(db, int(payload["comparison_id"]))
    elif job.kind == "comparison_export":
        comparison_or_404(db, int(payload["comparison_id"]))
    elif job.kind in {"tag_paper", "question_generation", "index_paper"}:
        paper_or_404(db, int(payload["paper_id"]))
    elif job.kind == "material_enrichment":
        material = db.get(Material, int(payload["material_id"]))
        if not material:
            raise HTTPException(404, "Job not found")
        subject_or_404(db, material.subject_id)
    elif job.kind == "syllabus_import":
        source = db.get(SyllabusImport, int(payload["source_id"]))
        if not source:
            raise HTTPException(404, "Job not found")
        user = current_user.get()
        if source.subject_id is not None:
            subject_or_404(db, source.subject_id)
        elif user and user.role != "admin" and source.submitted_by != user.id:
            raise HTTPException(404, "Job not found")
    else:
        raise HTTPException(404, "Job not found")
    return {"id": job.id, "kind": job.kind, "status": job.status, "progress": job.progress, "total": job.total, "attempts": job.attempts, "error": job.error, "created_at": job.created_at, "started_at": job.started_at, "finished_at": job.finished_at, "tagging": payload.get("tagging") or {"pattern": (payload.get("pattern") or {}).get("template_name")}, "paper": paper_out(paper_or_404(db, int(payload["paper_id"]))) if job.kind == "tag_paper" else None, "artifact_expired": bool((job.payload or {}).get("artifact_expired")), "download_url": f"/api/jobs/{job.id}/download" if job.kind == "comparison_export" and job.status == "complete" and (job.payload or {}).get("artifact") else None}


@app.get("/api/jobs/{job_id}/download", dependencies=[Depends(require_access)])
async def download_job_artifact(job_id: int, db: Db = Depends(db_session)):
    job = db.get(BackgroundJob, job_id)
    if not job or job.kind != "comparison_export":
        raise HTTPException(404, "Export job not found")
    comparison_or_404(db, int((job.payload or {})["comparison_id"]))
    if job.status != "complete":
        raise HTTPException(409, "The export is not ready")
    if (job.payload or {}).get("artifact_expired"):
        raise HTTPException(410, "The export file has expired; queue it again to create a new copy")
    filename = (job.payload or {}).get("artifact")
    if not isinstance(filename, str) or Path(filename).name != filename:
        raise HTTPException(404, "Export artifact not found")
    path = ROOT / "data" / "exports" / filename
    if not path.is_file():
        raise HTTPException(404, "Export artifact expired or was removed")
    return FileResponse(path, filename=filename, media_type=(job.payload or {}).get("media_type", "application/octet-stream"))


@app.get("/api/admin/operations", dependencies=[Depends(require_admin)])
async def operations_status(db: Db = Depends(db_session)):
    jobs = db.scalars(select(BackgroundJob).order_by(BackgroundJob.id.desc()).limit(100)).all()
    counts = {}
    for job in db.scalars(select(BackgroundJob)):
        key = f"{job.kind}:{job.status}"
        counts[key] = counts.get(key, 0) + 1
    upload_root = ROOT / "data" / "uploads"
    upload_bytes = sum(path.stat().st_size for path in upload_root.iterdir() if path.is_file()) if upload_root.exists() else 0
    export_root = ROOT / "data" / "exports"
    export_bytes = sum(path.stat().st_size for path in export_root.iterdir() if path.is_file()) if export_root.exists() else 0
    return {
        "database": engine.dialect.name,
        "auth_mode": os.getenv("AUTH_MODE", "local"),
        "job_counts": counts,
        "upload_bytes": upload_bytes,
        "export_bytes": export_bytes,
        "recent_jobs": [{"id": job.id, "kind": job.kind, "status": job.status, "progress": job.progress, "total": job.total, "attempts": job.attempts, "max_attempts": job.max_attempts, "error": job.error, "created_at": job.created_at, "finished_at": job.finished_at} for job in jobs],
    }


@app.get("/api/admin/backup", dependencies=[Depends(require_admin)])
async def download_backup():
    backup_root = ROOT / "data" / "backups"
    path = backup_root / f"workspace-{datetime.now(timezone.utc):%Y%m%dT%H%M%SZ}-{uuid4().hex[:8]}.zip"
    try:
        await asyncio.to_thread(create_backup, path)
    except Exception as exc:
        path.unlink(missing_ok=True)
        logger.exception("Workspace backup failed")
        raise HTTPException(500, f"Backup failed: {str(exc)[:300]}") from exc
    return FileResponse(path, filename=path.name, media_type="application/zip", background=BackgroundTask(lambda: path.unlink(missing_ok=True)))


@app.post("/api/admin/jobs/{job_id}/retry", dependencies=[Depends(require_admin)])
async def retry_job(job_id: int, db: Db = Depends(db_session)):
    job = db.get(BackgroundJob, job_id)
    if not job:
        raise HTTPException(404, "Job not found")
    if job.status not in ("failed", "queued"):
        raise HTTPException(409, "Only failed or queued jobs can be manually retried")
    if job.kind not in JOB_HANDLERS:
        raise HTTPException(409, f"No worker is registered for {job.kind}")
    job.status, job.error, job.attempts, job.started_at, job.finished_at = "queued", None, 0, None, None
    if job.kind == "comparison":
        comparison = db.get(Comparison, (job.payload or {}).get("comparison_id"))
        if comparison:
            comparison.status, comparison.error = "queued", None
    elif job.kind == "syllabus_import":
        source = db.get(SyllabusImport, (job.payload or {}).get("source_id"))
        if source:
            source.status, source.error = "queued", None
    elif job.kind == "material_enrichment":
        material = db.get(Material, (job.payload or {}).get("material_id"))
        if material:
            material.enrichment = {**(material.enrichment or {}), "status": "queued", "error": None}
    elif job.kind == "question_generation":
        paper = db.get(Paper, (job.payload or {}).get("paper_id"))
        if paper:
            paper.generation = {**(paper.generation or {}), "status": "queued", "error": None}
    db.add(AuditEvent(action="job_retry", details={"job_id": job.id, "kind": job.kind, "by": current_user.get().id}))
    db.commit()
    return {"id": job.id, "status": job.status, "kind": job.kind}


@app.post("/api/admin/jobs/{job_id}/cancel", dependencies=[Depends(require_admin)])
async def cancel_job(job_id: int, db: Db = Depends(db_session)):
    job = db.get(BackgroundJob, job_id)
    if not job:
        raise HTTPException(404, "Job not found")
    if job.status not in ("queued", "running", "cancel_requested"):
        raise HTTPException(409, "Only queued or running jobs can be cancelled")
    if job.status != "cancel_requested":
        immediate = job.status == "queued"
        job.status = "cancelled" if immediate else "cancel_requested"
        if immediate:
            job.finished_at = datetime.now(timezone.utc)
            payload = job.payload or {}
            if job.kind == "comparison":
                comparison = db.get(Comparison, payload.get("comparison_id"))
                if comparison:
                    comparison.status, comparison.error = "interrupted", "Cancelled before processing"
            elif job.kind == "material_enrichment":
                material = db.get(Material, payload.get("material_id"))
                if material:
                    material.enrichment = {**(material.enrichment or {}), "status": "cancelled", "error": "Cancelled before processing"}
            elif job.kind == "syllabus_import":
                source = db.get(SyllabusImport, payload.get("source_id"))
                if source:
                    source.status, source.error = "cancelled", "Cancelled before processing"
            elif job.kind == "question_generation":
                paper = db.get(Paper, payload.get("paper_id"))
                if paper:
                    paper.generation = {**(paper.generation or {}), "status": "cancelled", "error": "Cancelled before processing"}
        db.add(AuditEvent(action="job_cancel", details={"job_id": job.id, "kind": job.kind, "by": current_user.get().id, "immediate": immediate}))
        db.commit()
    return {"id": job.id, "kind": job.kind, "status": job.status}


@app.post("/api/comparisons/{cid}/cancel", dependencies=[Depends(require_access)])
async def cancel_comparison(cid: int, db: Db = Depends(db_session)):
    comparison = comparison_or_404(db, cid)
    if comparison.status not in ("queued", "running"):
        raise HTTPException(409, "Only queued or running comparisons can be cancelled")
    jobs = db.scalars(select(BackgroundJob).where(BackgroundJob.kind == "comparison").order_by(BackgroundJob.id.desc())).all()
    job = next((item for item in jobs if (item.payload or {}).get("comparison_id") == cid and item.status in ("queued", "running", "cancel_requested")), None)
    if not job:
        raise HTTPException(409, "No active worker job exists for this comparison")
    if job.status == "queued":
        job.status = "cancelled"
        job.finished_at = datetime.now(timezone.utc)
        comparison.status = "interrupted"
        comparison.error = "Cancelled before the worker started; resume to continue."
    else:
        job.status = "cancel_requested"
        comparison.error = "Cancellation requested; the current question pair will finish before stopping."
    db.add(AuditEvent(action="comparison_cancel", details={"comparison_id": cid, "job_id": job.id, "by": current_user.get().id}))
    db.commit()
    return {**comparison_out(comparison), "job_id": job.id, "cancel_requested": job.status == "cancel_requested"}


@app.get("/api/subjects", dependencies=[Depends(require_access)])
async def subjects(db: Db = Depends(db_session)):
    user = current_user.get()
    return [{"id": s.id, "code": s.code, "name": s.name, "active": True, "tree_version": s.tree_version, "review_questions": len([q for q in s.tree.get("review_questions", []) if q.get("status") != "resolved"])} for s in db.scalars(select(Subject).order_by(Subject.name)) if not user or user.role == "admin" or s.id in user.subject_ids]


@app.get("/api/subjects/{sid}/tree", dependencies=[Depends(require_access)])
async def get_tree(sid: int, db: Db = Depends(db_session)):
    return subject_or_404(db, sid).tree


def activate_tree(db: Db, tree: dict, reason: str) -> Subject:
    course = tree["course"]
    subject = db.scalar(select(Subject).where(Subject.code == course["code"]))
    if subject:
        db.add(TreeRevision(subject_id=subject.id, version=subject.tree_version, tree=subject.tree, reason=f"Before: {reason}"))
        subject.name = course["title"]
        subject.tree = tree
        subject.tree_version += 1
    else:
        subject = Subject(code=course["code"], name=course["title"], tree=tree, tree_version=1)
        db.add(subject)
        db.flush()
    valid_topics = {node["id"]: unit_for_topic(node["id"]) for node in tree.get("nodes", []) if node.get("kind") == "topic"}
    invalidated = 0
    for question in db.scalars(select(Question).join(Paper).where(Paper.subject_id == subject.id)):
        if question.confirmed and (not question.topic_ids or question.unit not in range(1, 6) or any(topic_id not in valid_topics or valid_topics[topic_id] != question.unit for topic_id in question.topic_ids)):
            before = question_out(question)
            question.topic_ids = []
            question.unit = None
            question.confirmed = False
            question.revision += 1
            db.add(AuditEvent(action="question_tree_invalidated", details={"question_id": question.id, "subject_id": subject.id, "revision": question.revision, "before": before, "after": question_out(question), "reason": reason}))
            invalidated += 1
    if invalidated:
        db.add(AuditEvent(action="tree_edit_invalidated_tags", details={"subject_id": subject.id, "question_count": invalidated, "reason": reason}))
    db.add(TreeRevision(subject_id=subject.id, version=subject.tree_version, tree=tree, reason=reason))
    return subject


class TreeImport(BaseModel):
    tree: dict


@app.put("/api/subjects/{sid}/tree", dependencies=[Depends(require_access)])
async def update_subject_tree(sid: int, body: TreeImport, db: Db = Depends(db_session)):
    subject = subject_or_404(db, sid)
    errors = validate_tree(body.tree)
    if errors:
        raise HTTPException(422, errors)
    course = body.tree.get("course", {})
    if course.get("code") != subject.code:
        raise HTTPException(422, "A direct tree edit must keep the subject course code")
    updated = activate_tree(db, body.tree, "direct syllabus tree edit")
    db.add(AuditEvent(action="tree_edit", details={"subject_id": sid, "tree_version": updated.tree_version}))
    db.commit()
    return {"id": updated.id, "tree_version": updated.tree_version, "active": True}


@app.get("/api/subjects/{sid}/tree/export", dependencies=[Depends(require_access)])
async def export_subject_tree(sid: int, db: Db = Depends(db_session)):
    tree = subject_or_404(db, sid).tree
    return Response(json.dumps(tree, ensure_ascii=False, indent=2), media_type="application/json", headers={"Content-Disposition": f'attachment; filename="{sid}-syllabus-tree.json"'})


@app.post("/api/subjects/import", dependencies=[Depends(require_access)])
async def import_tree(body: TreeImport, db: Db = Depends(db_session)):
    errors = validate_tree(body.tree)
    if errors:
        raise HTTPException(422, errors)
    course = body.tree.get("course", {})
    if not course.get("code") or not course.get("title"):
        raise HTTPException(422, "Tree needs a course code and title")
    subject = activate_tree(db, body.tree, "structured tree import")
    db.add(AuditEvent(action="tree_import", details={"code": course["code"]}))
    db.commit()
    return {"id": subject.id, "tree_version": subject.tree_version, "active": True}


@app.post("/api/subjects/import-syllabus", dependencies=[Depends(require_access)])
async def import_syllabus(file: UploadFile = File(...), db: Db = Depends(db_session), subject_id: int | None = Form(None)):
    if subject_id is not None:
        subject_or_404(db, subject_id)
    filename = Path(file.filename or "syllabus").name
    data = await file.read(25 * 1024 * 1024 + 1)
    try:
        text, warnings = await extract_text(filename, data)
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc
    digest = hashlib.sha256(data).hexdigest()
    duplicate = db.scalar(select(SyllabusImport.id).where(SyllabusImport.file_hash == digest))
    if duplicate:
        raise HTTPException(409, f"This syllabus file was already uploaded (import {duplicate})")
    storage = ROOT / "data" / "uploads"
    storage.mkdir(parents=True, exist_ok=True)
    target = storage / digest
    if not target.exists():
        target.write_bytes(data)
    source = SyllabusImport(filename=filename, file_hash=digest, extracted_text=text, warnings=warnings, status="queued", submitted_by=getattr(current_user.get(), "id", None))
    db.add(source)
    db.add(AuditEvent(action="syllabus_upload", details={"filename": filename, "hash": digest}))
    db.commit()
    db.refresh(source)
    job = enqueue_job(db, "syllabus_import", {"source_id": source.id}, total=1)
    return {"id": source.id, "filename": source.filename, "status": source.status, "warnings": source.warnings, "job_id": job.id}


@app.get("/api/syllabus-imports/{source_id}", dependencies=[Depends(require_access)])
async def syllabus_import_status(source_id: int, db: Db = Depends(db_session)):
    source = db.get(SyllabusImport, source_id)
    if not source:
        raise HTTPException(404, "Syllabus upload not found")
    if source.subject_id:
        subject_or_404(db, source.subject_id)
    elif current_user.get() and current_user.get().role != "admin" and source.submitted_by != current_user.get().id:
        raise HTTPException(404, "Syllabus upload not found")
    job = next((item for item in db.scalars(select(BackgroundJob).where(BackgroundJob.kind == "syllabus_import").order_by(BackgroundJob.id.desc())) if item.payload.get("source_id") == source_id), None)
    return {"id": source.id, "filename": source.filename, "subject_id": source.subject_id, "status": source.status, "error": source.error, "warnings": source.warnings, "job": {"id": job.id, "status": job.status, "progress": job.progress, "total": job.total, "error": job.error} if job else None}


@app.post("/api/syllabus-imports/{source_id}/retry", dependencies=[Depends(require_access)])
async def retry_syllabus_import(source_id: int, db: Db = Depends(db_session)):
    source = db.get(SyllabusImport, source_id)
    if not source:
        raise HTTPException(404, "Syllabus upload not found")
    if source.subject_id:
        subject_or_404(db, source.subject_id)
    elif current_user.get() and current_user.get().role != "admin" and source.submitted_by != current_user.get().id:
        raise HTTPException(404, "Syllabus upload not found")
    if source.status != "failed":
        raise HTTPException(409, "Only failed syllabus imports can be retried")
    source.status, source.error = "queued", None
    db.commit()
    job = enqueue_job(db, "syllabus_import", {"source_id": source.id}, total=1)
    return {"id": source.id, "status": source.status, "job_id": job.id}


@app.delete("/api/syllabus-imports/{source_id}", dependencies=[Depends(require_access)])
async def delete_syllabus_import(source_id: int, db: Db = Depends(db_session)):
    source = db.get(SyllabusImport, source_id)
    if not source:
        raise HTTPException(404, "Syllabus upload not found")
    if source.subject_id:
        subject_or_404(db, source.subject_id)
    elif current_user.get() and current_user.get().role != "admin" and source.submitted_by != current_user.get().id:
        raise HTTPException(403, "Only the uploader or an administrator can delete a pending syllabus import")
    digest = source.file_hash
    db.add(AuditEvent(action="syllabus_source_delete", details={"source_id": source.id, "filename": source.filename, "hash": digest, "subject_id": source.subject_id}))
    db.delete(source)
    db.commit()
    unlink_if_unused(db, digest)
    return {"deleted": source_id}


@app.get("/api/subjects/{sid}/tree/revisions", dependencies=[Depends(require_access)])
async def tree_revisions(sid: int, db: Db = Depends(db_session)):
    subject_or_404(db, sid)
    revisions = db.scalars(select(TreeRevision).where(TreeRevision.subject_id == sid).order_by(TreeRevision.version.desc(), TreeRevision.id.desc()))
    return [{"id": item.id, "version": item.version, "reason": item.reason, "created_at": item.created_at} for item in revisions]


@app.get("/api/subjects/{sid}/tree/revisions/{revision_id}/diff", dependencies=[Depends(require_access)])
async def tree_revision_diff(sid: int, revision_id: int, db: Db = Depends(db_session)):
    subject = subject_or_404(db, sid)
    revision = db.get(TreeRevision, revision_id)
    if not revision or revision.subject_id != sid:
        raise HTTPException(404, "Tree revision not found")
    before = {node.get("id"): node for node in revision.tree.get("nodes", []) if isinstance(node, dict) and isinstance(node.get("id"), str)}
    after = {node.get("id"): node for node in subject.tree.get("nodes", []) if isinstance(node, dict) and isinstance(node.get("id"), str)}
    added = [{"id": key, "node": after[key]} for key in sorted(after.keys() - before.keys())]
    removed = [{"id": key, "node": before[key]} for key in sorted(before.keys() - after.keys())]
    changed = []
    for node_id in sorted(before.keys() & after.keys()):
        fields = [{"field": key, "before": before[node_id].get(key), "after": after[node_id].get(key)} for key in sorted(set(before[node_id]) | set(after[node_id])) if before[node_id].get(key) != after[node_id].get(key)]
        if fields:
            changed.append({"id": node_id, "fields": fields})
    return {"revision_id": revision.id, "revision_version": revision.version, "current_version": subject.tree_version, "added": added, "removed": removed, "changed": changed}


@app.post("/api/subjects/{sid}/tree/revisions/{revision_id}/restore", dependencies=[Depends(require_access)])
async def restore_tree_revision(sid: int, revision_id: int, db: Db = Depends(db_session)):
    subject = subject_or_404(db, sid)
    revision = db.get(TreeRevision, revision_id)
    if not revision or revision.subject_id != sid:
        raise HTTPException(404, "Tree revision not found")
    errors = validate_tree(revision.tree)
    if errors:
        raise HTTPException(422, errors)
    subject = activate_tree(db, revision.tree, f"restore revision {revision.version}")
    db.add(AuditEvent(action="tree_restore", details={"subject_id": sid, "revision_id": revision_id, "new_version": subject.tree_version}))
    db.commit()
    return {"id": subject.id, "tree_version": subject.tree_version, "active": True}


class PaperCreate(BaseModel):
    subject_id: int
    title: str = Field(min_length=1, max_length=250)
    year: int | None = Field(default=None, ge=1900, le=2200)
    status: str = "draft"


@app.get("/api/papers", dependencies=[Depends(require_access)])
async def papers(subject_id: int | None = None, include_archived: bool = False, db: Db = Depends(db_session)):
    query = select(Paper).order_by(Paper.created_at.desc())
    if subject_id:
        subject_or_404(db, subject_id)
        query = query.where(Paper.subject_id == subject_id)
    if not include_archived:
        query = query.where(Paper.archived.is_(False))
    user = current_user.get()
    return [paper_out(p) for p in db.scalars(query) if not user or user.role == "admin" or p.subject_id in user.subject_ids]


@app.get("/api/questions/search", dependencies=[Depends(require_access)])
async def search_question_bank(subject_id: int, query: str = "", unit: int | None = None, limit: int = 50, semantic: bool = False, db: Db = Depends(db_session)):
    subject_or_404(db, subject_id)
    if semantic and query.strip():
        if len(query) > 4000 or unit is not None and unit not in range(1, 6):
            raise HTTPException(422, "Use a query of at most 4000 characters and a unit from 1 through 5")
        try:
            matches = await semantic_search(db, subject_id, query.strip(), unit=unit, limit=limit)
        except ModelUnavailable as exc:
            raise HTTPException(503, str(exc))
        return {"results": [{**question_out(item["question"]), "paper_title": item["question"].paper.title, "paper_year": item["question"].paper.year, "similarity": item["similarity"], "matched_part": item["matched_part"]} for item in matches], "query": query, "count": len(matches), "search": "semantic", "note": "Retrieval similarity is not an A/B/C classification. Only current indexed questions are searched."}
    terms = [term.strip() for term in query.split() if term.strip()][:8]
    if not terms or len(" ".join(terms)) < 2:
        return {"results": [], "query": query, "count": 0}
    limit = min(max(limit, 1), 100)
    statement = select(Question).join(Paper).where(Paper.subject_id == subject_id, Paper.status == "bank", Paper.archived.is_(False))
    if unit is not None:
        if unit not in range(1, 6):
            raise HTTPException(422, "Unit must be from 1 through 5")
        statement = statement.where(Question.unit == unit)
    for term in terms:
        escaped = term.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
        statement = statement.where(Question.text.ilike(f"%{escaped}%", escape="\\"))
    questions = db.scalars(statement.order_by(Question.id.desc()).limit(limit)).all()
    return {"results": [{**question_out(question), "paper_title": question.paper.title, "paper_year": question.paper.year} for question in questions], "query": query, "count": len(questions)}


@app.get("/api/papers/{pid}", dependencies=[Depends(require_access)])
async def paper_detail(pid: int, db: Db = Depends(db_session)):
    return paper_out(paper_or_404(db, pid))


@app.get("/api/papers/{pid}/analysis", dependencies=[Depends(require_access)])
async def paper_analysis(pid: int, db: Db = Depends(db_session)):
    paper = paper_or_404(db, pid)
    subject = subject_or_404(db, paper.subject_id)
    paper_questions = list(paper.questions)
    units = []
    topic_nodes = {node.get("id") for node in subject.tree.get("nodes", []) if node.get("kind") == "topic" and isinstance(node.get("id"), str)}
    for number in range(1, 6):
        questions = [q for q in paper_questions if q.confirmed and q.unit == number]
        unit_topics = {topic_id for topic_id in topic_nodes if unit_for_topic(topic_id) == number}
        covered = {topic_id for q in questions for topic_id in q.topic_ids if topic_id in unit_topics}
        units.append({"unit": number, "questions": len(questions), "marks": sum(q.marks or 0 for q in questions), "missing_marks": sum(q.marks is None for q in questions), "covered_topic_ids": sorted(covered), "covered_topics": len(covered), "topic_count": len(unit_topics), "coverage_percent": round(100 * len(covered) / len(unit_topics), 1) if unit_topics else None})
    question_ids = {q.id for q in paper_questions}
    results = list(db.scalars(select(PairResult).join(Comparison, PairResult.comparison_id == Comparison.id).where(Comparison.subject_id == paper.subject_id, or_(PairResult.target_id.in_(question_ids), PairResult.reference_id.in_(question_ids))))) if question_ids else []
    latest_by_pair = {}
    for result in results:
        pair = tuple(sorted((result.target_id, result.reference_id)))
        if pair not in latest_by_pair or result.id > latest_by_pair[pair].id:
            latest_by_pair[pair] = result
    current_results = []
    stale = 0
    for result in latest_by_pair.values():
        left, right = db.get(Question, result.target_id), db.get(Question, result.reference_id)
        if not left or not right:
            stale += 1
            continue
        if result.cache_key != pair_key(left, right, subject.tree_version):
            stale += 1
        else:
            current_results.append(result)
    label_counts = {label: sum(result.label == label for result in current_results) for label in ("A", "B", "C", "review")}
    known_marks = [q.marks for q in paper_questions if q.marks is not None]
    marks_distribution = {}
    for mark in known_marks:
        marks_distribution[str(mark)] = marks_distribution.get(str(mark), 0) + 1
    question_distribution = {str(unit): sum(q.unit == unit for q in paper_questions) for unit in range(1, 6)}
    question_distribution["untagged_or_invalid"] = sum(q.unit not in range(1, 6) for q in paper_questions)
    marked_total = sum(known_marks)
    return {
        "paper_id": pid,
        "title": paper.title,
        "confirmed": sum(q.confirmed for q in paper.questions),
        "unconfirmed": sum(not q.confirmed for q in paper.questions),
        "total_questions": len(paper.questions),
        "total_marks": sum(u["marks"] for u in units),
        "missing_marks": sum(u["missing_marks"] for u in units),
        "marks_summary": {
            "known_question_count": len(known_marks),
            "mean_per_question": round(sum(known_marks) / len(known_marks), 2) if known_marks else None,
            "minimum": min(known_marks) if known_marks else None,
            "maximum": max(known_marks) if known_marks else None,
            "distribution": marks_distribution,
            "unit_share_percent": {str(item["unit"]): round(100 * item["marks"] / marked_total, 1) if marked_total else None for item in units},
        },
        "question_distribution_by_unit": question_distribution,
        "units": units,
        "redundancy": {"A": label_counts["A"], "B": label_counts["B"], "C": label_counts["C"], "review": label_counts["review"], "stale": stale, "saved_pair_results": len(latest_by_pair), "current_pair_results": len(current_results)},
        "diagnostics": {
            "untagged_question_ids": [q.id for q in paper.questions if not q.confirmed],
            "questions_missing_marks": [q.id for q in paper.questions if q.marks is None],
            "units_without_confirmed_questions": [unit["unit"] for unit in units if unit["questions"] == 0],
            "note": "These are direct corpus checks; no overall quality or predictability score is calibrated.",
        },
        "quality_score": None,
        "predictability_score": None,
        "score_status": "not_calibrated",
    }


@app.get("/api/subjects/{sid}/analytics", dependencies=[Depends(require_access)])
async def subject_analytics(sid: int, reference_paper_ids: list[int] | None = None, db: Db = Depends(db_session)):
    subject = subject_or_404(db, sid)
    papers = db.scalars(select(Paper).where(Paper.subject_id == sid, Paper.status.in_(("reference", "bank")), Paper.archived.is_(False))).all()
    if reference_paper_ids is not None:
        unknown = set(reference_paper_ids) - {paper.id for paper in papers}
        if unknown:
            raise HTTPException(422, f"Reference papers are not available: {sorted(unknown)}")
        papers = [paper for paper in papers if paper.id in set(reference_paper_ids)]
    questions = [question for paper in papers for question in paper.questions if question.confirmed]
    units = []
    year_map: dict[int, dict] = {}
    for number in range(1, 6):
        unit_prefix = f"{subject.code}/u{number}/"
        topic_ids = {node["id"] for node in subject.tree.get("nodes", []) if node.get("kind") == "topic" and node.get("id", "").startswith(unit_prefix)}
        unit_questions = [question for question in questions if question.unit == number]
        covered = {topic_id for question in unit_questions for topic_id in question.topic_ids if topic_id in topic_ids}
        units.append({"unit": number, "question_count": len(unit_questions), "marks": sum(question.marks or 0 for question in unit_questions), "covered_topic_count": len(covered), "topic_count": len(topic_ids), "coverage_percent": round(100 * len(covered) / len(topic_ids), 1) if topic_ids else None})
    for paper in papers:
        if paper.year is None:
            continue
        entry = year_map.setdefault(paper.year, {"year": paper.year, "units": {unit: 0 for unit in range(1, 6)}, "topics": {}})
        for question in paper.questions:
            if question.confirmed and question.unit in range(1, 6):
                entry["units"][question.unit] += 1
                for topic_id in question.topic_ids:
                    entry["topics"][topic_id] = entry["topics"].get(topic_id, 0) + 1
    topic_nodes = {node["id"]: node for node in subject.tree.get("nodes", []) if node.get("kind") == "topic"}
    repeated_topics = []
    for topic_id in topic_nodes:
        years = [year for year, entry in year_map.items() if entry["topics"].get(topic_id, 0)]
        if years:
            repeated_topics.append({"topic_id": topic_id, "label": topic_nodes[topic_id].get("label", topic_id), "unit": unit_for_topic(topic_id), "years_present": years, "years_count": len(years), "question_count": sum(entry["topics"].get(topic_id, 0) for entry in year_map.values())})
    recurrence_trend, prior_topics = [], set()
    for year in sorted(year_map):
        current_topics = set(year_map[year]["topics"])
        repeated = current_topics & prior_topics
        recurrence_trend.append({"year": year, "observed_topic_count": len(current_topics), "topics_seen_in_earlier_years": len(repeated), "new_topic_count": len(current_topics - prior_topics), "repeat_share_percent": round(100 * len(repeated) / len(current_topics), 1) if current_topics else None})
        prior_topics.update(current_topics)
    return {"subject_id": sid, "tree_version": subject.tree_version, "reference_papers": [{"id": paper.id, "title": paper.title, "year": paper.year} for paper in papers], "reference_count": len(papers), "confirmed_questions": len(questions), "untagged_questions": sum(not question.confirmed for paper in papers for question in paper.questions), "unit_coverage": units, "coverage_by_year": [year_map[year] for year in sorted(year_map)], "topic_recurrence": sorted(repeated_topics, key=lambda item: (-item["years_count"], -item["question_count"], item["topic_id"])), "historical_topic_recurrence_by_year": recurrence_trend}


@app.post("/api/papers", dependencies=[Depends(require_access)])
async def create_paper(body: PaperCreate, db: Db = Depends(db_session)):
    subject_or_404(db, body.subject_id)
    if body.status not in ("draft", "reference", "bank"):
        raise HTTPException(422, "Invalid paper status")
    paper = Paper(**body.model_dump())
    db.add(paper)
    db.commit()
    db.refresh(paper)
    return paper_out(paper)


class PaperStatusUpdate(BaseModel):
    status: str


@app.put("/api/papers/{pid}/status", dependencies=[Depends(require_access)])
async def update_paper_status(pid: int, body: PaperStatusUpdate, db: Db = Depends(db_session)):
    paper = paper_or_404(db, pid)
    if body.status not in ("draft", "reference", "bank"):
        raise HTTPException(422, "Paper status must be draft, reference or bank")
    old_status = paper.status
    paper.status = body.status
    db.add(AuditEvent(action="paper_status_update", details={"paper_id": pid, "from": old_status, "to": paper.status, "generated": paper.generated}))
    db.commit()
    return paper_out(paper)


@app.post("/api/papers/{pid}/archive", dependencies=[Depends(require_access)])
async def archive_paper(pid: int, db: Db = Depends(db_session)):
    paper = paper_or_404(db, pid)
    if not paper.archived:
        paper.archived = True
        db.add(AuditEvent(action="paper_archive", details={"paper_id": pid, "status": paper.status, "by": current_user.get().id}))
        db.commit()
    return paper_out(paper)


@app.post("/api/papers/{pid}/restore", dependencies=[Depends(require_access)])
async def restore_paper(pid: int, db: Db = Depends(db_session)):
    paper = paper_or_404(db, pid)
    if paper.archived:
        paper.archived = False
        db.add(AuditEvent(action="paper_restore", details={"paper_id": pid, "status": paper.status, "by": current_user.get().id}))
        db.commit()
    return paper_out(paper)


class GenerationInput(BaseModel):
    subject_id: int
    title: str = Field(min_length=1, max_length=250)
    reference_ids: list[int] = Field(default_factory=list, max_length=120)
    insem: bool = False
    template_id: int | None = None
    unit: int | None = Field(default=None, ge=1, le=5)
    topic_ids: list[str] = Field(default_factory=list, max_length=3)
    marks: int | None = Field(default=None, ge=1, le=100)
    count: int = Field(default=3, ge=1, le=10)
    difficulty: str = "intermediate"
    constraints: str = Field(default="", max_length=1200)
    prefer_bc: bool = False


@app.post("/api/generation", dependencies=[Depends(require_access)])
async def start_generation(body: GenerationInput, db: Db = Depends(db_session)):
    subject = subject_or_404(db, body.subject_id)
    if not MODEL:
        raise HTTPException(503, "Configure a local JUDGE_MODEL_NAME before generating questions")
    inventory = await inspect_local_models()
    if not inventory.get("judge_installed"):
        raise HTTPException(503, f"Configured judge model is not installed in local Ollama: {inventory.get('error', inventory.get('configured', {}).get('judge'))}")
    if body.difficulty not in {"introductory", "intermediate", "advanced"}:
        raise HTTPException(422, "Difficulty must be introductory, intermediate or advanced")
    topic_nodes = {node["id"]: node for node in subject.tree.get("nodes", []) if node.get("kind") == "topic"}
    slots = []
    if body.insem:
        labels = {1: ["1a", "1b", "1c", "1d", "1e", "2a", "2b", "2c", "2d", "2e"], 2: ["3a", "3b", "3c", "3d", "3e", "4a", "4b", "4c", "4d", "4e"]}
        for unit, numbers in labels.items():
            available = [topic_id for topic_id in topic_nodes if unit_for_topic(topic_id) == unit]
            if not available:
                raise HTTPException(422, f"Unit {unit} needs syllabus topics before an InSem paper can be generated")
            for index, number in enumerate(numbers):
                slots.append({"section": f"Unit {unit}", "number": number, "unit": unit,
                    "topic_ids": [available[index * len(available) // len(numbers)]], "marks": 4,
                    "count": 1, "difficulty": body.difficulty, "constraints": body.constraints})
        if not body.reference_ids:
            past_papers = db.scalars(select(Paper).where(Paper.subject_id == body.subject_id,
                Paper.archived.is_(False), Paper.generated.is_(False), Paper.file_hash.is_not(None)).order_by(Paper.created_at.desc(), Paper.id.desc())).all()
            reference_paper = next((p for p in past_papers if {1, 2}.issubset({q.unit for q in p.questions if q.confirmed})), None)
            if not reference_paper:
                raise HTTPException(422, "Upload a past paper with tagged questions in both Unit 1 and Unit 2 before generating an InSem draft")
            body.reference_ids = [q.id for q in reference_paper.questions if q.confirmed and q.unit in (1, 2)]
    elif body.template_id:
        template = db.get(Template, body.template_id)
        if not template or template.kind != "paper" or template.subject_id not in (None, body.subject_id):
            raise HTTPException(422, "Select a paper-pattern template available to this subject")
        template_errors = validate_template("paper", template.definition)
        template_errors.extend(validate_template_topics("paper", template.definition, subject))
        if template_errors:
            raise HTTPException(422, template_errors)
        sections = template.definition.get("sections", [])
        if not isinstance(sections, list) or not sections:
            raise HTTPException(422, "Configure at least one section and question slot in the paper template")
        for section in sections:
            if not isinstance(section, dict) or not isinstance(section.get("slots", []), list):
                raise HTTPException(422, "Paper template has an invalid section or slot")
            for slot in section.get("slots", []):
                unit = (slot.get("unit") or body.unit) if isinstance(slot, dict) else None
                topic_ids = (slot.get("topic_ids") or body.topic_ids) if isinstance(slot, dict) else []
                marks = (slot.get("marks") or body.marks) if isinstance(slot, dict) else None
                count = slot.get("count", 1) if isinstance(slot, dict) else 0
                slots.append({"section": str(section.get("name", "Section")), "slot_name": slot.get("name"), "unit": unit, "topic_ids": topic_ids, "marks": marks, "count": count, "choice_group": slot.get("choice_group"), "choice_rule": slot.get("choice_rule"), "choose_count": slot.get("choose_count"), "difficulty": body.difficulty, "constraints": body.constraints})
    elif body.unit is not None and body.topic_ids and body.marks:
        slots = [{"section": "Generated questions", "unit": body.unit, "topic_ids": body.topic_ids, "marks": body.marks, "count": body.count, "difficulty": body.difficulty, "constraints": body.constraints}]
    else:
        raise HTTPException(422, "Choose a paper template or provide unit, topic IDs and marks")
    if not slots or any(not isinstance(slot, dict) for slot in slots):
        raise HTTPException(422, "Generation needs at least one valid paper slot")
    if any(not isinstance(slot.get("count"), int) or isinstance(slot.get("count"), bool) or slot["count"] < 1 for slot in slots):
        raise HTTPException(422, "Every generation slot needs a positive integer question count")
    if sum(slot["count"] for slot in slots) > 20:
        raise HTTPException(422, "Generation is limited to 20 candidate questions per draft")
    if not body.reference_ids:
        raise HTTPException(422, "At least one tagged past question is required for each unit")
    questions = {q.id: q for q in db.scalars(select(Question).where(Question.id.in_(body.reference_ids)))}
    if len(questions) != len(set(body.reference_ids)) or len(body.reference_ids) != len(set(body.reference_ids)) or any(q.paper.subject_id != body.subject_id or q.paper.archived or not q.confirmed or q.unit not in range(1, 6) for q in questions.values()):
        raise HTTPException(422, "Every reference must be an existing confirmed question in this subject")
    for slot in slots:
        if not isinstance(slot.get("unit"), int) or isinstance(slot.get("unit"), bool) or slot["unit"] not in range(1, 6) or not isinstance(slot.get("marks"), int) or isinstance(slot.get("marks"), bool) or slot["marks"] < 1:
            raise HTTPException(422, "Every generation slot needs a valid unit, positive marks and count")
        ids = slot.get("topic_ids")
        if not isinstance(ids, list) or not ids or len(ids) > 3 or any(not isinstance(topic_id, str) for topic_id in ids) or len(set(ids)) != len(ids) or any(topic_id not in topic_nodes or unit_for_topic(topic_id) != slot["unit"] for topic_id in ids):
            raise HTTPException(422, f"Generation slot topics must belong to unit {slot.get('unit')}")
    if body.insem and len({question.paper_id for question in questions.values()}) != 1:
        raise HTTPException(422, "Choose exactly one past paper for the InSem quick screen")
    for slot in slots:
        if not any(question.unit == slot["unit"] for question in questions.values()):
            raise HTTPException(422, f"Provide at least one confirmed reference question from unit {slot['unit']}")
    paper = Paper(subject_id=body.subject_id, title=body.title, status="draft", generated=True, generation={"status": "queued", "prompt_version": GENERATION_PROMPT_VERSION, "reference_ids": body.reference_ids, "slots": slots, "template_id": body.template_id, "prefer_bc": body.prefer_bc, "screen_strategy": "insem_quick_v1" if body.insem else None, "reference_paper_id": next(iter(questions.values())).paper_id if body.insem else None, "reference_paper_title": next(iter(questions.values())).paper.title if body.insem else None})
    db.add(paper)
    total = sum(slot["count"] for slot in slots)
    db.flush()
    job = BackgroundJob(kind="question_generation", payload={"paper_id": paper.id, "reference_ids": body.reference_ids, "slots": slots, "prefer_bc": body.prefer_bc}, total=total)
    db.add(job)
    db.add(AuditEvent(action="question_generation_start", details={"paper_id": paper.id, "slots": slots, "reference_ids": body.reference_ids}))
    db.commit()
    db.refresh(paper)
    db.refresh(job)
    return {"paper": paper_out(paper), "job_id": job.id, "status": "queued", "candidate_count": total}


@app.get("/api/papers/{pid}/replacement-analysis", dependencies=[Depends(require_access)])
async def replacement_analysis(pid: int, db: Db = Depends(db_session)):
    draft = paper_or_404(db, pid)
    source_id = (draft.generation or {}).get("replacement_for_question_id")
    if not draft.generated or not isinstance(source_id, int):
        raise HTTPException(404, "This paper does not contain replacement suggestions")
    source = db.get(Question, source_id)
    if not source or source.paper.subject_id != draft.subject_id:
        raise HTTPException(409, "The source question for these suggestions is no longer available")
    comparison_id = (draft.generation or {}).get("comparison_id")
    comparison = db.get(Comparison, comparison_id) if isinstance(comparison_id, int) else None
    return replacement_report(subject_or_404(db, draft.subject_id), source, draft, comparison)


@app.post("/api/papers/{pid}/generation/retry", dependencies=[Depends(require_access)])
async def retry_generation(pid: int, db: Db = Depends(db_session)):
    paper = paper_or_404(db, pid)
    if not paper.generated or paper.generation.get("status") not in ("failed", "check_failed"):
        raise HTTPException(409, "Only a failed generated draft can be retried")
    if paper.generation.get("status") == "check_failed":
        comparison_id = paper.generation.get("comparison_id")
        comparison = db.get(Comparison, comparison_id) if comparison_id else None
        if not comparison:
            raise HTTPException(409, "The failed screening record is missing")
        comparison.status, comparison.error = "queued", None
        paper.generation = {**paper.generation, "status": "checking", "check_error": None}
        job = BackgroundJob(kind="comparison", payload={"comparison_id": comparison.id}, total=comparison.total)
        db.add(job)
        db.commit()
        db.refresh(job)
        return {"paper_id": pid, "job_id": job.id, "status": "checking"}
    if not MODEL:
        raise HTTPException(503, "Configure a local JUDGE_MODEL_NAME before retrying generation")
    slots = paper.generation.get("slots")
    reference_ids = paper.generation.get("reference_ids")
    if not isinstance(slots, list) or not isinstance(reference_ids, list):
        raise HTTPException(409, "Saved generation settings are incomplete")
    paper.generation = {**paper.generation, "status": "queued", "error": None}
    job = BackgroundJob(kind="question_generation", payload={"paper_id": pid, "reference_ids": reference_ids, "slots": slots, "prefer_bc": bool(paper.generation.get("prefer_bc"))}, total=sum(slot.get("count", 0) for slot in slots))
    db.add(job)
    db.commit()
    db.refresh(job)
    return {"paper_id": pid, "job_id": job.id, "status": "queued"}


@app.post("/api/papers/{pid}/upload", dependencies=[Depends(require_access)])
async def upload_paper(pid: int, file: UploadFile = File(...), db: Db = Depends(db_session)):
    paper = paper_or_404(db, pid)
    if paper.file_hash:
        raise HTTPException(409, "This paper already has an uploaded file. Create a new paper record to keep its questions and comparisons intact.")
    data = await file.read(25 * 1024 * 1024 + 1)
    try:
        text, warnings = await extract_text(file.filename or "", data)
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc
    digest = hashlib.sha256(data).hexdigest()
    duplicates = list(db.scalars(select(Paper).where(Paper.file_hash == digest, Paper.id != pid)))
    if duplicates:
        existing = ", ".join(str(item.id) for item in duplicates)
        raise HTTPException(409, f"This file is already attached to paper(s): {existing}")
    storage = ROOT / "data" / "uploads"
    storage.mkdir(parents=True, exist_ok=True)
    target = storage / digest
    if not target.exists():
        target.write_bytes(data)
    paper.filename = Path(file.filename or "upload").name
    paper.file_hash = digest
    extracted = expand_subquestions(draft_questions(text))
    for index, item in enumerate(extracted, start=1):
        db.add(Question(paper_id=paper.id, number=item.get("number", str(index)), text=item["text"], marks=item.get("marks"), structure=item.get("structure", {})))
    db.add(AuditEvent(action="paper_upload", details={"paper_id": pid, "hash": digest, "questions_drafted": len(extracted)}))
    db.commit()
    apply_tagging_pattern(db, paper)
    tagging_job = enqueue_paper_tagging(db, paper) if MODEL and extracted else None
    return {"paper": paper_out(paper), "warnings": warnings, "drafted_questions": len(extracted), "extracted_text_preview": text[:3000], "tagging_job_id": tagging_job.id if tagging_job else None}


@app.post("/api/papers/bulk-upload", dependencies=[Depends(require_access)])
async def bulk_upload_papers(
    files: list[UploadFile] = File(...),
    subject_id: int = Form(...),
    status: str = Form("reference"),
    year: int | None = Form(None),
    db: Db = Depends(db_session),
    auto_tag: bool = True,
):
    subject_or_404(db, subject_id)
    if status not in ("draft", "reference", "bank"):
        raise HTTPException(422, "Paper status must be draft, reference or bank")
    if not files or len(files) > 20:
        raise HTTPException(422, "Choose between 1 and 20 files per bulk upload")
    outcomes = []
    batch_hashes = set()
    for file in files:
        filename = Path(file.filename or "upload").name
        try:
            data = await file.read(25 * 1024 * 1024 + 1)
            text, warnings = await extract_text(filename, data)
            digest = hashlib.sha256(data).hexdigest()
            duplicate = db.scalar(select(Paper.id).where(Paper.file_hash == digest))
            if digest in batch_hashes or duplicate:
                outcomes.append({"filename": filename, "status": "duplicate", "existing_paper_id": duplicate, "error": "This file is already in the paper corpus or this upload batch."})
                continue
            batch_hashes.add(digest)
            questions = expand_subquestions(draft_questions(text))
            paper = Paper(subject_id=subject_id, title=Path(filename).stem[:250] or "Uploaded paper", year=year, status=status, filename=filename, file_hash=digest)
            db.add(paper)
            db.flush()
            for index, item in enumerate(questions, start=1):
                db.add(Question(paper_id=paper.id, number=item.get("number", str(index)), text=item["text"], marks=item.get("marks"), structure=item.get("structure", {})))
            storage = ROOT / "data" / "uploads"
            storage.mkdir(parents=True, exist_ok=True)
            target = storage / digest
            if not target.exists():
                target.write_bytes(data)
            db.add(AuditEvent(action="paper_bulk_upload", details={"paper_id": paper.id, "hash": digest, "questions_drafted": len(questions)}))
            db.commit()
            db.refresh(paper)
            apply_tagging_pattern(db, paper)
            tagging_job = enqueue_paper_tagging(db, paper) if auto_tag and MODEL and questions else None
            outcomes.append({"filename": filename, "status": "imported", "paper": paper_out(paper), "warnings": warnings, "drafted_questions": len(questions), "tagging_job_id": tagging_job.id if tagging_job else None})
        except Exception as exc:
            db.rollback()
            outcomes.append({"filename": filename, "status": "failed", "error": str(exc)[:500] or "Could not process this file"})
        finally:
            await file.close()
    return {"outcomes": outcomes, "imported": sum(item["status"] == "imported" for item in outcomes), "failed": sum(item["status"] == "failed" for item in outcomes), "duplicates": sum(item["status"] == "duplicate" for item in outcomes)}


class QuestionPartInput(BaseModel):
    id: str = Field(min_length=1, max_length=40)
    label: str = Field(min_length=1, max_length=40)
    text: str = Field(min_length=1, max_length=10000)
    marks: int | None = Field(default=None, ge=1, le=100)


class ChoiceGroupInput(BaseModel):
    id: str = Field(min_length=1, max_length=40)
    rule: str = "choose_any"
    choose_count: int = Field(default=1, ge=1)
    alternatives: list[list[str]] = Field(min_length=2)


class QuestionStructureInput(BaseModel):
    parts: list[QuestionPartInput] = Field(default_factory=list)
    choice_groups: list[ChoiceGroupInput] = Field(default_factory=list)
    warnings: list[str] = Field(default_factory=list)


def validated_question_structure(value, question_text: str) -> dict:
    if value is None:
        extracted = draft_questions("1. " + question_text)
        value = extracted[0]["structure"] if extracted else {"parts": [], "choice_groups": [], "warnings": []}
    paper_context = value.get("paper_context") if isinstance(value, dict) else None
    structure = value.model_dump() if isinstance(value, QuestionStructureInput) else QuestionStructureInput.model_validate(value).model_dump()
    if isinstance(paper_context, dict) and isinstance(paper_context.get("root_number"), str) and isinstance(paper_context.get("choice_groups", []), list):
        structure["paper_context"] = {"root_number": paper_context["root_number"][:50], "choice_groups": [group for group in paper_context["choice_groups"] if isinstance(group, dict) and isinstance(group.get("group"), int) and isinstance(group.get("alternatives"), list) and all(isinstance(option, list) and all(isinstance(number, str) and len(number) <= 30 for number in option) for option in group["alternatives"])]}
    parts = structure["parts"]
    ids = [part["id"] for part in parts]
    if len(ids) != len(set(ids)):
        raise HTTPException(422, "Question part IDs must be unique")
    normalized_source = re.sub(r"\s+", " ", question_text).casefold()
    if any(re.sub(r"\s+", " ", part["text"]).casefold() not in normalized_source for part in parts):
        raise HTTPException(422, "Each structured part must come from the question text")
    seen_groups = set()
    for group in structure["choice_groups"]:
        if group["id"] in seen_groups or group["rule"] != "choose_any":
            raise HTTPException(422, "Choice groups need unique IDs and the choose_any rule")
        seen_groups.add(group["id"])
        alternatives = group["alternatives"]
        if group["choose_count"] > len(alternatives) or any(not alt or len(alt) != len(set(alt)) or any(part_id not in ids for part_id in alt) for alt in alternatives):
            raise HTTPException(422, "Choice alternatives must reference existing parts and have a valid choose count")
    return structure


class QuestionInput(BaseModel):
    number: str = Field(min_length=1, max_length=50)
    text: str = Field(min_length=3)
    marks: int | None = Field(default=None, ge=1)
    unit: int = Field(ge=1, le=5)
    topic_ids: list[str] = Field(default_factory=list)
    structure: QuestionStructureInput | None = None
    confirmed: bool = False


def check_topics(subject: Subject, unit: int, topic_ids: list[str]):
    valid = {n["id"] for n in subject.tree["nodes"]}
    if any(t not in valid or unit_for_topic(t) != unit for t in topic_ids):
        raise HTTPException(422, "Topics must exist in the selected unit")


@app.post("/api/papers/{pid}/questions", dependencies=[Depends(require_access)])
async def add_question(pid: int, body: QuestionInput, db: Db = Depends(db_session)):
    paper = paper_or_404(db, pid)
    check_topics(subject_or_404(db, paper.subject_id), body.unit, body.topic_ids)
    fields = body.model_dump(exclude={"structure"})
    fields.pop("confirmed", None)
    question = Question(paper_id=pid, **fields, structure=validated_question_structure(body.structure, body.text), confirmed=bool(body.topic_ids))
    db.add(question)
    db.flush()
    db.add(AuditEvent(action="question_create", details={"question_id": question.id, "subject_id": paper.subject_id, "after": question_out(question), "by": current_user.get().id}))
    db.commit()
    return question_out(question)


@app.put("/api/questions/{qid}", dependencies=[Depends(require_access)])
async def update_question(qid: int, body: QuestionInput, db: Db = Depends(db_session)):
    question = db.get(Question, qid)
    if not question:
        raise HTTPException(404, "Question not found")
    check_topics(subject_or_404(db, question.paper.subject_id), body.unit, body.topic_ids)
    before = question_out(question)
    fields = body.model_dump(exclude={"structure"})
    fields["structure"] = validated_question_structure(body.structure, body.text)
    for field, value in {**fields, "confirmed": bool(body.topic_ids)}.items():
        setattr(question, field, value)
    question.revision += 1
    db.add(AuditEvent(action="question_edit", details={"question_id": qid, "subject_id": question.paper.subject_id, "revision": question.revision, "before": before, "after": question_out(question), "by": current_user.get().id}))
    db.commit()
    return question_out(question)


class ApplyReplacementInput(BaseModel):
    candidate_id: int


@app.post("/api/questions/{qid}/apply-replacement", dependencies=[Depends(require_access)])
async def apply_replacement(qid: int, body: ApplyReplacementInput, db: Db = Depends(db_session)):
    source = db.get(Question, qid)
    if not source:
        raise HTTPException(404, "Source question not found")
    subject = subject_or_404(db, source.paper.subject_id)
    candidate = db.get(Question, body.candidate_id)
    draft = candidate.paper if candidate else None
    generation = (draft.generation or {}) if draft else {}
    if not draft or not draft.generated or draft.subject_id != subject.id or generation.get("replacement_for_question_id") != qid:
        raise HTTPException(422, "Candidate is not a replacement suggestion for this question")
    if source.paper.archived or draft.archived or source.revision != generation.get("source_revision"):
        raise HTTPException(409, "The source paper or question changed after suggestions were generated")
    if not candidate.confirmed or candidate.unit != source.unit or candidate.marks != source.marks:
        raise HTTPException(422, "Replacement must retain the source unit and marks and have confirmed syllabus tags")
    comparison_id = generation.get("comparison_id")
    comparison = db.get(Comparison, comparison_id) if isinstance(comparison_id, int) else None
    if not comparison or comparison.status != "complete":
        raise HTTPException(409, "Replacement screening must finish before applying a candidate")
    candidate_results = [result for result in comparison.results if result.target_id == candidate.id]
    expected_refs = set(generation.get("reference_ids") or [])
    if not expected_refs or {result.reference_id for result in candidate_results} != expected_refs or any((reference := db.get(Question, result.reference_id)) is None or result.cache_key != pair_key(candidate, reference, subject.tree_version) for result in candidate_results):
        raise HTTPException(409, "Replacement screening is missing or stale; rerun it before applying this candidate")
    updated = await update_question(qid, QuestionInput(number=source.number, text=candidate.text, marks=source.marks, unit=source.unit, topic_ids=candidate.topic_ids, structure=candidate.structure), db)
    draft.generation = {**generation, "applied_candidate_id": candidate.id, "applied_source_revision": updated["revision"]}
    db.add(AuditEvent(action="replacement_applied", details={"source_question_id": qid, "candidate_id": candidate.id, "draft_paper_id": draft.id, "screening_comparison_id": comparison.id, "by": current_user.get().id}))
    db.commit()
    return {"success": True, "source_question": updated, "candidate_id": candidate.id, "draft_paper_id": draft.id, "screening_labels": {label: sum(result.label == label for result in candidate_results) for label in ("A", "B", "C", "review")}}


@app.get("/api/questions/{qid}/history", dependencies=[Depends(require_access)])
async def question_history(qid: int, db: Db = Depends(db_session)):
    question = db.get(Question, qid)
    subject_id = question.paper.subject_id if question else None
    if subject_id is None:
        deleted = db.scalars(select(AuditEvent).where(AuditEvent.action == "question_delete").order_by(AuditEvent.id.desc())).all()
        event = next((item for item in deleted if (item.details or {}).get("question_id") == qid), None)
        subject_id = (event.details or {}).get("subject_id") if event else None
    if subject_id is None:
        raise HTTPException(404, "Question history not found")
    subject_or_404(db, subject_id)
    actions = ("question_create", "question_edit", "question_auto_tag", "question_tree_invalidated", "question_delete")
    events = db.scalars(select(AuditEvent).where(
        AuditEvent.action.in_(actions),
        AuditEvent.details["question_id"].as_integer() == qid,
    ).order_by(AuditEvent.id.desc()).limit(100)).all()
    return [{"id": event.id, "action": event.action, "details": event.details, "created_at": event.created_at} for event in events]


def template_question_units(template: Template, questions: list[Question]) -> dict[int, int]:
    """Subquestions inherit their numbered parent's unit when that parent is unambiguous."""
    labels = {normalized_question_label(label): item["unit"]
              for item in (template.definition or {}).get("question_pattern", {}).get("units", [])
              for label in item.get("questions", [])}
    roots: dict[str, set[int]] = {}
    for label, unit in labels.items():
        match = re.match(r"\d+", label)
        if match:
            roots.setdefault(match[0], set()).add(unit)
    result = {}
    for question in questions:
        label = normalized_question_label(question.number)
        unit = labels.get(label)
        match = re.match(r"\d+", label)
        inherited = roots.get(match[0], set()) if match else set()
        if unit is None and len(inherited) == 1:
            unit = next(iter(inherited))
        if type(unit) is not int or unit not in range(1, 6):
            return {}
        result[question.id] = unit
    return result


def apply_tagging_pattern(db: Db, paper: Paper) -> dict:
    candidates = [(template, template_question_units(template, paper.questions))
                  for template in db.scalars(select(Template).where(
                      Template.kind == "comparison", or_(Template.subject_id == paper.subject_id, Template.subject_id.is_(None))))
                  if (template.definition or {}).get("question_pattern")]
    candidates = [(template, units) for template, units in candidates if units]
    # Only infer a pattern when every matching template agrees on the units.
    if not candidates or len({tuple(sorted(units.items())) for _, units in candidates}) != 1:
        return {}
    template, units = candidates[0]
    for question in paper.questions:
        unit = units[question.id]
        if question.unit == unit:
            continue
        before = question_out(question)
        question.unit = unit
        if any(unit_for_topic(topic) != unit for topic in question.topic_ids):
            question.topic_ids = []
        question.confirmed = bool(question.topic_ids)
        question.revision += 1
        db.add(AuditEvent(action="question_pattern_unit", details={"question_id": question.id, "template_id": template.id, "before": before, "after": question_out(question)}))
    db.commit()
    return {"template_id": template.id, "template_name": template.name, "units": {str(key): value for key, value in units.items()}}


def enqueue_paper_tagging(db: Db, paper: Paper):
    pattern = apply_tagging_pattern(db, paper)
    return enqueue_job(db, "tag_paper", {"paper_id": paper.id, "pattern": pattern}, total=len(paper.questions))


async def tag_paper(pid: int, job_id: int | None = None):
    with Session() as db:
        paper = paper_or_404(db, pid)
        subject = subject_or_404(db, paper.subject_id)
        job = db.get(BackgroundJob, job_id) if job_id else None
        pattern = (job.payload or {}).get("pattern") if job else None
        if not pattern:
            pattern = apply_tagging_pattern(db, paper)
        units = (pattern or {}).get("units", {})
        topics = [{"id": n["id"], "label": n["label"], "unit": unit_for_topic(n["id"]), "description": n.get("description", "")} for n in subject.tree["nodes"] if n.get("kind") == "topic"]
        pending = [q for q in paper.questions if not (q.confirmed and q.topic_ids)]
        total = len(paper.questions)
        completed = total - len(pending)
        failures = 0
        details = {"pattern": (pattern or {}).get("template_name"), "phase": "waiting", "questions": [], "needs_review": {}}

        def update_status():
            if job_id:
                with Session.begin() as status_db:
                    current = status_db.get(BackgroundJob, job_id)
                    if current:
                        current.payload = {**current.payload, "pattern": pattern, "tagging": dict(details)}
                set_progress(job_id, completed, total)

        update_status()
        # Keep each call within one template unit; share its topic list across four questions.
        for unit in dict.fromkeys(units.get(str(q.id)) for q in pending):
            group = [q for q in pending if units.get(str(q.id)) == unit]
            allowed = [t for t in topics if unit is None or t["unit"] == unit]
            valid = {t["id"]: t["unit"] for t in allowed}
            for offset in range(0, len(group), 4):
                if job_id:
                    ensure_not_cancelled(job_id)
                batch = group[offset:offset + 4]
                details.update(phase="waiting", questions=[q.number for q in batch])
                update_status()
                def activity(event):
                    if details["phase"] != "working" and event.get("event") in {"content", "thinking"}:
                        details["phase"] = "working"
                        update_status()
                try:
                    response = await model_json([
                        {"role": "system", "content": "Assign syllabus topics to each supplied exam question. The supplied unit, when present, is fixed by the paper template and MUST NOT change. Return JSON: {results:[{id:integer,topic_ids:[1 to 3 exact supplied IDs],unit:integer,uncertain:boolean,reason:string}]}. Include every question ID once. If no topic fits confidently, return topic_ids:[] and uncertain:true. Do not force a match."},
                        {"role": "user", "content": json.dumps({"unit": unit, "questions": [{"id": q.id, "text": q.text} for q in batch], "topics": allowed})},
                    ], think=False, on_event=activity)
                    rows = response.get("results", [])
                    if not isinstance(rows, list):
                        raise ValueError("Tagging response has no results list")
                except Exception as exc:
                    rows = []
                    batch_error = str(exc)[:500]
                else:
                    batch_error = "Missing, duplicated or invalid tagging result"
                if job_id:
                    ensure_not_cancelled(job_id)
                for question in batch:
                    matches = [row for row in rows if isinstance(row, dict) and type(row.get("id")) is int and row["id"] == question.id]
                    result = matches[0] if len(matches) == 1 else {}
                    ids = result.get("topic_ids")
                    tagged_unit = unit if unit is not None else result.get("unit")
                    uncertain = result.get("uncertain") is True and ids == []
                    accepted = (result.get("uncertain") is False and isinstance(ids, list) and 1 <= len(ids) <= 3
                                and type(tagged_unit) is int and tagged_unit in range(1, 6)
                                and all(isinstance(topic, str) and topic in valid and valid[topic] == tagged_unit for topic in ids))
                    if accepted:
                        before = question_out(question)
                        question.topic_ids = list(dict.fromkeys(ids))
                        question.unit = tagged_unit
                        question.confirmed = True
                        question.revision += 1
                        db.add(AuditEvent(action="question_auto_tag", details={"question_id": question.id, "subject_id": subject.id, "revision": question.revision, "before": before, "after": question_out(question)}))
                    else:
                        reason = str(result.get("reason") or ("No confident syllabus topic match" if uncertain else batch_error))[:500]
                        details["needs_review"][str(question.id)] = reason
                        db.add(AuditEvent(action="question_tag_uncertain" if uncertain else "question_tag_failed", details={"question_id": question.id, "reason": reason}))
                        if not uncertain:
                            failures += 1
                    db.commit()  # Checkpoint each question before updating visible progress.
                    completed += 1
                    update_status()
        details.update(phase="finished", questions=[])
        update_status()
        if failures:
            raise RuntimeError(f"Automatic tagging failed for {failures} question(s); untagged questions remain editable")


@app.post("/api/papers/{pid}/tag", dependencies=[Depends(require_access)])
async def start_tagging(pid: int, db: Db = Depends(db_session)):
    paper = paper_or_404(db, pid)
    apply_tagging_pattern(db, paper)
    if not MODEL:
        raise HTTPException(503, "Configure a local JUDGE_MODEL_NAME before automatic syllabus tagging")
    job = enqueue_paper_tagging(db, paper)
    return {"status": "queued", "job_id": job.id}


def unlink_if_unused(db: Db, digest: str | None):
    template_uses_hash = False
    if digest:
        for template in db.scalars(select(Template)):
            source = (template.definition or {}).get("source")
            if isinstance(source, dict) and source.get("sha256") == digest:
                template_uses_hash = True
                break
    if digest and not template_uses_hash and not db.scalar(select(Paper.id).where(Paper.file_hash == digest)) and not db.scalar(select(Material.id).where(Material.file_hash == digest)) and not db.scalar(select(SyllabusImport.id).where(SyllabusImport.file_hash == digest)):
        (ROOT / "data" / "uploads" / digest).unlink(missing_ok=True)


@app.delete("/api/papers/{pid}", dependencies=[Depends(require_access)])
async def delete_paper(pid: int, db: Db = Depends(db_session)):
    paper = paper_or_404(db, pid)
    ids = {q.id for q in paper.questions}
    affected = [c for c in db.scalars(select(Comparison)) if ids.intersection(c.target_ids + c.reference_ids)]
    affected_ids = {comparison.id for comparison in affected}
    if affected_ids:
        for child in db.scalars(select(Comparison).where(Comparison.rerun_of_id.in_(affected_ids))):
            child.rerun_of_id = None
    for comparison in affected:
        for result in db.scalars(select(PairResult).where(PairResult.comparison_id == comparison.id)):
            db.delete(result)
        db.delete(comparison)
    # Stop queued work for the paper and remove its completed job entries.
    # A running worker sees cancel_requested at its next checkpoint.
    job_ids = []
    for job in db.scalars(select(BackgroundJob).where(BackgroundJob.kind.in_(("tag_paper", "index_paper", "question_generation")))):
        if (job.payload or {}).get("paper_id") != paper.id:
            continue
        job_ids.append(job.id)
        if job.status == "running":
            job.status = "cancel_requested"
        else:
            db.delete(job)
    # Remove persisted upload cards for this paper so reopening an old chat
    # cannot resurrect a paper that was explicitly deleted.
    removed_chat_messages = 0
    removed_conversations = []
    touched_conversations = set()
    conversations = list(db.scalars(select(Conversation).where(Conversation.subject_id == paper.subject_id)))
    for conversation in conversations:
        messages = list(db.scalars(select(ChatMessage).where(ChatMessage.conversation_id == conversation.id).order_by(ChatMessage.id)))
        for index, message in enumerate(messages):
            upload = (message.payload or {}).get("upload")
            uploaded_paper = upload.get("paper") if isinstance(upload, dict) else None
            if not isinstance(uploaded_paper, dict) or uploaded_paper.get("id") != paper.id:
                continue
            touched_conversations.add(conversation.id)
            db.delete(message)
            removed_chat_messages += 1
            if index and messages[index - 1].role == "user" and paper.filename and paper.filename in messages[index - 1].text:
                db.delete(messages[index - 1])
                removed_chat_messages += 1
    db.flush()
    for conversation in conversations:
        if conversation.id not in touched_conversations:
            continue
        remaining = list(db.scalars(select(ChatMessage).where(ChatMessage.conversation_id == conversation.id).order_by(ChatMessage.id)))
        if not remaining:
            db.delete(conversation)
            removed_conversations.append(conversation.id)
        elif paper.filename and paper.filename in conversation.title:
            next_user_message = next((item.text.strip() for item in remaining if item.role == "user" and item.text.strip()), "Conversation")
            conversation.title = next_user_message[:80]
    for question in paper.questions:
        db.add(AuditEvent(action="question_delete", details={"question_id": question.id, "subject_id": paper.subject_id, "cascade_paper_id": paper.id, "by": current_user.get().id}))
    digest = paper.file_hash
    db.delete(paper)
    db.add(AuditEvent(action="paper_delete", details={"paper_id": paper.id, "subject_id": paper.subject_id, "comparisons_removed": [c.id for c in affected], "jobs_removed_or_cancelled": job_ids, "chat_upload_messages_removed": removed_chat_messages, "empty_conversations_removed": removed_conversations, "by": current_user.get().id}))
    db.commit()
    unlink_if_unused(db, digest)
    return {"deleted": pid, "comparisons_removed": [c.id for c in affected], "conversations_removed": removed_conversations}


@app.delete("/api/questions/{qid}", dependencies=[Depends(require_access)])
async def delete_question(qid: int, db: Db = Depends(db_session)):
    question = db.get(Question, qid)
    if not question:
        raise HTTPException(404, "Question not found")
    subject_or_404(db, question.paper.subject_id)
    affected = [c for c in db.scalars(select(Comparison)) if qid in c.target_ids + c.reference_ids]
    affected_ids = {comparison.id for comparison in affected}
    if affected_ids:
        for child in db.scalars(select(Comparison).where(Comparison.rerun_of_id.in_(affected_ids))):
            child.rerun_of_id = None
    for comparison in affected:
        for result in db.scalars(select(PairResult).where(PairResult.comparison_id == comparison.id)):
            db.delete(result)
        db.delete(comparison)
    db.add(AuditEvent(action="question_delete", details={"question_id": question.id, "subject_id": question.paper.subject_id, "comparisons_removed": [c.id for c in affected], "by": current_user.get().id}))
    db.delete(question)
    db.commit()
    return {"deleted": qid, "comparisons_removed": [c.id for c in affected]}


class CompareInput(BaseModel):
    subject_id: int
    title: str = Field(min_length=1, max_length=250)
    target_ids: list[int]
    reference_ids: list[int]


def resolved_pairs(db: Db, body: CompareInput):
    return resolved_pairs_with_units(db, body)


def resolved_pairs_with_units(db: Db, body: CompareInput, unit_overrides: dict[int, int] | None = None):
    subject_or_404(db, body.subject_id)
    ids = set(body.target_ids + body.reference_ids)
    questions = {q.id: q for q in db.scalars(select(Question).where(Question.id.in_(ids)))}
    if not ids or ids != set(questions):
        raise HTTPException(422, "Select existing target and reference questions")
    overrides = unit_overrides or {}
    if any(q.paper.subject_id != body.subject_id or q.paper.archived or
           (q.id in overrides and overrides[q.id] not in range(1, 6)) or
           (q.id not in overrides and (not q.confirmed or q.unit is None)) for q in questions.values()):
        raise HTTPException(422, "All selected questions need a valid pattern unit or confirmed syllabus unit")
    units = {qid: overrides.get(qid, q.unit) for qid, q in questions.items()}
    pairs, skipped, seen_pairs = [], [], set()
    for target_id in dict.fromkeys(body.target_ids):
        for reference_id in dict.fromkeys(body.reference_ids):
            if target_id == reference_id:
                continue
            a, b = questions[target_id], questions[reference_id]
            unordered_key = tuple(sorted((a.id, b.id)))
            if unordered_key in seen_pairs:
                continue
            seen_pairs.add(unordered_key)
            if units[a.id] != units[b.id]:
                skipped.append({"target_id": a.id, "reference_id": b.id, "reason": "different units"})
            else:
                pairs.append((a, b))
    return pairs, skipped


def comparison_question_with_unit(question: Question, unit: int):
    """Provide pattern-derived unit context to the classifier without editing the saved question."""
    return SimpleNamespace(
        id=question.id, revision=question.revision, text=question.text, marks=question.marks,
        unit=unit, topic_ids=question.topic_ids or [], structure=question.structure or {},
        number=question.number, paper=question.paper,
    )


@app.post("/api/comparisons/preview", dependencies=[Depends(require_access)])
async def preview_comparison(body: CompareInput, db: Db = Depends(db_session)):
    pairs, skipped = resolved_pairs(db, body)
    inventory = await inspect_local_models()
    models_ready = classifier_models_ready() and inventory.get("roles_available")
    return {"pairs": len(pairs), "skipped": skipped, "model_configured": bool(models_ready), "comparison_models": classifier_model_status(), "model_runtime": inventory, "subject_active": True}


class PaperComparisonInput(BaseModel):
    subject_id: int
    base_paper_id: int
    comparison_paper_ids: list[int] = Field(min_length=1, max_length=20)
    template_id: int
    preview: bool = False


def resolve_pattern_papers(db: Db, body: PaperComparisonInput):
    subject = subject_or_404(db, body.subject_id)
    template = db.get(Template, body.template_id)
    if not template or template.kind != "comparison" or template.subject_id not in (None, body.subject_id):
        raise HTTPException(422, "Choose a comparison pattern available to this subject")
    definition = template.definition or {}
    template_errors = validate_template("comparison", definition)
    if template_errors:
        raise HTTPException(409, template_errors)
    pattern = definition.get("question_pattern") or {}
    pattern_units = pattern.get("units") or []
    label_units = {normalized_question_label(label): unit.get("unit") for unit in pattern_units for label in unit.get("questions", [])}
    if not label_units:
        raise HTTPException(422, "The selected comparison pattern has no question-to-unit layout")
    paper_ids = [body.base_paper_id, *body.comparison_paper_ids]
    if len(paper_ids) != len(set(paper_ids)):
        raise HTTPException(422, "Choose each paper only once")
    found = {paper.id: paper for paper in db.scalars(select(Paper).where(Paper.id.in_(paper_ids), Paper.subject_id == body.subject_id, Paper.archived.is_(False)))}
    if len(found) != len(paper_ids):
        raise HTTPException(422, "Every selected paper must be active and belong to the current subject")
    if any(paper.generated for paper in found.values()):
        raise HTTPException(422, "Generated draft papers cannot be used in a paper comparison. Choose uploaded papers instead.")
    issues, target_ids, reference_ids, unit_map = [], [], [], {}

    def extract_for(paper: Paper):
        by_label: dict[str, list[Question]] = {}
        extras = []
        for question in paper.questions:
            label = normalized_question_label(question.number)
            if label in label_units:
                by_label.setdefault(label, []).append(question)
            else:
                extras.append(question.number)
        selected = []
        for label, unit in label_units.items():
            matches = by_label.get(label, [])
            if len(matches) != 1:
                issues.append(f"{paper.title}: Q{label} {'was not extracted' if not matches else 'was extracted more than once'}.")
                continue
            question = matches[0]
            selected.append(question.id)
            unit_map[question.id] = unit
        if extras:
            issues.append(f"{paper.title}: question labels outside this pattern: {', '.join(extras[:8])}.")
        return selected

    base = found[body.base_paper_id]
    references = [found[paper_id] for paper_id in body.comparison_paper_ids]
    target_ids = extract_for(base)
    for paper in references:
        reference_ids.extend(extract_for(paper))
    request = CompareInput(subject_id=body.subject_id, title=f"{base.title} vs {', '.join(paper.title for paper in references)}"[:250], target_ids=target_ids, reference_ids=reference_ids)
    pairs, skipped = (resolved_pairs_with_units(db, request, unit_map) if target_ids and reference_ids and not issues else ([], []))
    return subject, template, base, references, request, unit_map, pairs, skipped, issues


def group_comparison_blocks(pairs, unit_map):
    groups = {}
    for target, reference in pairs:
        unit = unit_map.get(target.id, target.unit)
        groups.setdefault((target.id, reference.paper_id, unit), (target, []))[1].append(reference)
    return [(target, references, unit) for (_, _, unit), (target, references) in groups.items()]


@app.post("/api/comparisons/by-papers", dependencies=[Depends(require_access)])
async def compare_papers_by_pattern(body: PaperComparisonInput, db: Db = Depends(db_session)):
    subject, template, base, references, request, unit_map, pairs, skipped, issues = resolve_pattern_papers(db, body)
    blocks = group_comparison_blocks(pairs, unit_map)
    if body.preview:
        inventory = await inspect_local_models()
        models_ready = bool(BLOCK_JUDGE_MODEL and inventory.get("judge_installed"))
        return {"checks": len(blocks), "base_paper": base.title, "comparison_papers": [paper.title for paper in references], "targets": len(set(request.target_ids)), "references": len(set(request.reference_ids)), "pairs": len(pairs), "skipped": skipped, "pattern_issues": issues, "model_configured": bool(models_ready), "model_runtime": inventory}
    if issues:
        raise HTTPException(422, "Pattern check needs attention: " + "; ".join(issues[:12]))
    if not pairs:
        raise HTTPException(422, "The selected pattern resolved to no same-unit question pairs")
    if not BLOCK_JUDGE_MODEL:
        raise HTTPException(503, "Configure local judge/helper models before starting comparisons")
    inventory = await inspect_local_models()
    if not inventory.get("judge_installed"):
        raise HTTPException(503, "Install the configured local Ollama model artifacts before comparison")
    snapshot = {"scope": "paper_pattern", "template_id": template.id, "template_name": template.name, "template_version": template.version,
                "base_paper_id": base.id, "comparison_paper_ids": [paper.id for paper in references], "strategy": BLOCK_VERSION,
                "question_units": {str(question_id): unit for question_id, unit in unit_map.items()}}
    comparison = Comparison(subject_id=subject.id, title=request.title, target_ids=request.target_ids, reference_ids=request.reference_ids,
                           template_snapshot=snapshot, total=len(blocks))
    db.add(comparison)
    db.flush()
    job = enqueue_job(db, "comparison", {"comparison_id": comparison.id}, total=len(blocks))
    db.add(AuditEvent(action="comparison_create", details={"comparison_id": comparison.id, "pairs": len(pairs), "checks": len(blocks), "template_id": template.id, "base_paper_id": base.id, "comparison_paper_ids": [paper.id for paper in references], "via": "paper_pattern_picker"}))
    db.commit()
    return {**comparison_out(comparison), "job_id": job.id}


async def run_block_comparison(db, comparison, pairs, unit_overrides, subject, job_id):
    blocks = group_comparison_blocks(pairs, unit_overrides)
    comparison.total = len(blocks)
    completed = {(r.target_id, r.reference_id) for r in db.scalars(select(PairResult).where(PairResult.comparison_id == comparison.id))}
    comparison.progress = len(completed)
    comparison.error = None
    db.commit()
    if job_id:
        set_progress(job_id, comparison.progress, comparison.total)

    def activity(target, reference, event):
        db.add(ComparisonActivity(comparison_id=comparison.id, target_id=target.id, reference_id=reference.id,
            agent=event.get("agent", "judge"), event=event["event"], payload=event.get("payload", {})))
        db.commit()

    try:
        for target, references, unit in blocks:
            if job_id:
                ensure_not_cancelled(job_id)
            first = references[0]
            if (target.id, first.id) in completed:
                continue
            activity(target, first, {"agent": "system", "event": "block_started", "payload": {"number": target.number, "paper": first.paper.title, "unit": unit, "questions": len(references)}})
            key = block_key(target, references, unit, subject.tree_version)
            cached = db.scalar(select(PairResult).where(PairResult.cache_key == key, PairResult.source == BLOCK_PROMPT_VERSION).order_by(PairResult.id.desc()))
            if cached:
                result = {"label": cached.label, "reason": cached.reason, "evidence": cached.evidence, "source": cached.source, "model": cached.model}
                activity(target, first, {"agent": "system", "event": "cache_hit", "payload": {"label": cached.label}})
            else:
                result = await classify_block(target, references, unit, on_event=lambda event: activity(target, first, {"agent": "judge", **event}), syllabus=subject.tree)
            if job_id:
                ensure_not_cancelled(job_id)
            db.add(PairResult(comparison_id=comparison.id, target_id=target.id, reference_id=first.id, cache_key=key, **result))
            comparison.progress += 1
            db.commit()
            activity(target, first, {"agent": "system", "event": "pair_finished", "payload": {"label": result["label"]}})
            if job_id:
                set_progress(job_id, comparison.progress, comparison.total)
        if job_id:
            ensure_not_cancelled(job_id)
        comparison.status = "complete"
        db.commit()
    except JobCancelled as exc:
        comparison.status, comparison.error = "interrupted", str(exc)
        db.commit()
        raise
    except Exception as exc:
        comparison.status, comparison.error = "failed", str(exc)
        db.commit()
        raise


async def run_comparison(cid: int, job_id: int | None = None):
    with Session() as db:
        comparison = comparison_or_404(db, cid)
        body = CompareInput(subject_id=comparison.subject_id, title=comparison.title, target_ids=comparison.target_ids, reference_ids=comparison.reference_ids)
        unit_overrides = {int(key): value for key, value in (comparison.template_snapshot or {}).get("question_units", {}).items()}
        pairs, _ = resolved_pairs_with_units(db, body, unit_overrides)
        subject = subject_or_404(db, body.subject_id)
        comparison.status = "running"
        if (comparison.template_snapshot or {}).get("strategy") == BLOCK_VERSION:
            await run_block_comparison(db, comparison, pairs, unit_overrides, subject, job_id)
            return
        comparison.total = len(pairs)
        completed = {(r.target_id, r.reference_id) for r in db.scalars(select(PairResult).where(PairResult.comparison_id == cid))}
        comparison.progress = len(completed)
        comparison.error = None
        db.commit()
        if job_id:
            set_progress(job_id, comparison.progress, comparison.total)
        def record_activity(target, reference, event):
            db.add(ComparisonActivity(
                comparison_id=cid,
                target_id=target.id,
                reference_id=reference.id,
                agent=str(event.get("agent", "system"))[:30],
                event=str(event.get("event", "activity"))[:40],
                payload=event.get("payload") if isinstance(event.get("payload"), dict) else {},
            ))
            db.commit()

        try:
            analysis_cache = {}
            active_pair = None
            for stored_a, stored_b in pairs:
                if job_id:
                    ensure_not_cancelled(job_id)
                if (stored_a.id, stored_b.id) in completed:
                    continue
                active_pair = (stored_a, stored_b)
                a = comparison_question_with_unit(stored_a, unit_overrides[stored_a.id]) if stored_a.id in unit_overrides else stored_a
                b = comparison_question_with_unit(stored_b, unit_overrides[stored_b.id]) if stored_b.id in unit_overrides else stored_b
                record_activity(a, b, {"agent": "system", "event": "pair_started", "payload": {"target_id": a.id, "reference_id": b.id}})
                key = pair_key(a, b, subject.tree_version)
                cached = db.scalar(select(PairResult).where(PairResult.cache_key == key, PairResult.source != "faculty_case_review").order_by(PairResult.id.desc()))
                if cached:
                    result = {"label": cached.label, "reason": cached.reason, "evidence": cached.evidence, "source": "cache", "model": cached.model}
                    record_activity(a, b, {"agent": "system", "event": "cache_hit", "payload": {"label": cached.label}})
                else:
                    missing_profiles = []
                    for question in (stored_a, stored_b):
                        if question.id not in analysis_cache:
                            analysis_key = analysis_cache_key(question)
                            cached_analysis = db.scalar(select(QuestionAnalysis).where(QuestionAnalysis.cache_key == analysis_key))
                            if cached_analysis:
                                analysis_cache[question.id] = cached_analysis.analysis
                            else:
                                missing_profiles.append((question, analysis_key))
                    if missing_profiles:
                        for question, _ in missing_profiles:
                            record_activity(a, b, {"agent": "extractor", "event": "started", "payload": {"question_id": question.id}})
                        extracted = await analyze_questions([question for question, _ in missing_profiles])
                        for (question, analysis_key), profile in zip(missing_profiles, extracted):
                            db.add(QuestionAnalysis(
                                question_id=question.id,
                                question_revision=question.revision,
                                analyzer_model=os.getenv("JUDGE_MODEL_NAME", os.getenv("MODEL_NAME", "")),
                                prompt_version=ANALYSIS_PROMPT_VERSION,
                                cache_key=analysis_key,
                                analysis=profile,
                            ))
                            analysis_cache[question.id] = profile
                        db.commit()
                        for question, _ in missing_profiles:
                            profile = analysis_cache[question.id]
                            record_activity(a, b, {"agent": "extractor", "event": "finished", "payload": {"question_id": question.id, "requirements": len(profile.get("requirements", [])), "concerns": profile.get("concerns", [])}})
                    profiles = [analysis_cache[question.id] for question in (a, b)]
                    result = await judge_pair(a, b, subject.tree, tuple(profiles), on_event=lambda event: record_activity(a, b, event))
                if job_id:
                    ensure_not_cancelled(job_id)
                db.add(PairResult(comparison_id=cid, target_id=a.id, reference_id=b.id, cache_key=key, **result))
                comparison.progress += 1
                db.commit()
                record_activity(a, b, {"agent": "system", "event": "pair_finished", "payload": {"label": result["label"], "source": result.get("source")}})
                if job_id:
                    set_progress(job_id, comparison.progress, comparison.total)
            if job_id:
                ensure_not_cancelled(job_id)
            comparison.status = "complete"
            if comparison.generated_paper_id:
                generated_paper = db.get(Paper, comparison.generated_paper_id)
                if generated_paper:
                    outcomes = list(db.scalars(select(PairResult).where(PairResult.comparison_id == comparison.id)))
                    counts = {label: sum(result.label == label for result in outcomes) for label in ("A", "B", "C", "review")}
                    generated_paper.generation = {**(generated_paper.generation or {}), "status": "checked", "comparison_id": comparison.id, "outcome_counts": counts, "claim": "provisional screening only; C results do not prove novelty"}
            db.commit()
        except JobCancelled as exc:
            comparison.status = "interrupted"
            comparison.error = str(exc)
            db.commit()
            raise
        except Exception as exc:
            if active_pair:
                try:
                    record_activity(*active_pair, {"agent": "system", "event": "failed", "payload": {"error": str(exc)[:500]}})
                except Exception:
                    logger.exception("Could not persist comparison failure event")
            comparison.status = "failed"
            comparison.error = str(exc)
            db.commit()
            raise


@app.post("/api/comparisons", dependencies=[Depends(require_access)])
async def create_comparison(body: CompareInput, db: Db = Depends(db_session)):
    pairs, _ = resolved_pairs(db, body)
    if not pairs:
        raise HTTPException(422, "No same-unit pairs were selected")
    if not classifier_models_ready():
        raise HTTPException(503, "Configure local JUDGE_MODEL_NAME and HELPER_MODEL_NAME values before starting comparisons")
    inventory = await inspect_local_models()
    if not inventory.get("roles_available"):
        raise HTTPException(503, "Install the configured Ollama model artifact before comparison")
    comparison = Comparison(**body.model_dump(), total=len(pairs))
    db.add(comparison)
    db.commit()
    job = enqueue_job(db, "comparison", {"comparison_id": comparison.id}, total=len(pairs))
    return {**comparison_out(comparison), "job_id": job.id}


@app.post("/api/comparisons/{cid}/resume", dependencies=[Depends(require_access)])
async def resume_comparison(cid: int, db: Db = Depends(db_session)):
    comparison = comparison_or_404(db, cid)
    if comparison.status not in ("interrupted", "failed"):
        raise HTTPException(409, "Only interrupted or failed comparisons can resume")
    if not (BLOCK_JUDGE_MODEL if (comparison.template_snapshot or {}).get("strategy") == BLOCK_VERSION else classifier_models_ready()):
        raise HTTPException(503, "Configure local JUDGE_MODEL_NAME and HELPER_MODEL_NAME values before resuming")
    inventory = await inspect_local_models()
    if not (inventory.get("judge_installed") if (comparison.template_snapshot or {}).get("strategy") == BLOCK_VERSION else inventory.get("roles_available")):
        raise HTTPException(503, "Install the configured Ollama model artifact before resuming")
    unit_overrides = {int(key): value for key, value in (comparison.template_snapshot or {}).get("question_units", {}).items()}
    resolved_pairs_with_units(db, CompareInput(subject_id=comparison.subject_id, title=comparison.title, target_ids=comparison.target_ids, reference_ids=comparison.reference_ids), unit_overrides)
    comparison.status = "queued"
    comparison.error = None
    db.commit()
    job = enqueue_job(db, "comparison", {"comparison_id": cid}, total=comparison.total)
    return {**comparison_out(comparison), "job_id": job.id}


@app.post("/api/comparisons/{cid}/rerun", dependencies=[Depends(require_access)])
async def rerun_comparison(cid: int, db: Db = Depends(db_session)):
    source = comparison_or_404(db, cid)
    inventory = await inspect_local_models()
    block_strategy = (source.template_snapshot or {}).get("strategy") == BLOCK_VERSION
    if not ((BLOCK_JUDGE_MODEL and inventory.get("judge_installed")) if block_strategy else (classifier_models_ready() and inventory.get("roles_available"))):
        raise HTTPException(503, "Install the configured local comparison model before refreshing this comparison")
    unit_overrides = {int(key): value for key, value in (source.template_snapshot or {}).get("question_units", {}).items()}
    pairs, _ = resolved_pairs_with_units(db, CompareInput(subject_id=source.subject_id, title=source.title, target_ids=source.target_ids, reference_ids=source.reference_ids), unit_overrides)
    if not pairs:
        raise HTTPException(422, "This comparison no longer contains valid same-unit pairs")
    total = len(group_comparison_blocks(pairs, unit_overrides)) if (source.template_snapshot or {}).get("strategy") == BLOCK_VERSION else len(pairs)
    refreshed = Comparison(
        subject_id=source.subject_id,
        title=f"Refresh of #{source.id} · {source.title}"[:250],
        target_ids=list(source.target_ids),
        reference_ids=list(source.reference_ids),
        rerun_of_id=source.id,
        template_snapshot=source.template_snapshot or {},
        status="queued",
        total=total,
    )
    db.add(refreshed)
    db.flush()
    job = enqueue_job(db, "comparison", {"comparison_id": refreshed.id}, total=total)
    db.add(AuditEvent(action="comparison_rerun", details={"source_comparison_id": source.id, "comparison_id": refreshed.id, "pair_count": len(pairs), "by": current_user.get().id}))
    db.commit()
    return {**comparison_out(refreshed), "job_id": job.id}


@app.get("/api/comparisons", dependencies=[Depends(require_access)])
async def comparisons(db: Db = Depends(db_session)):
    user = current_user.get()
    return [comparison_out(c) for c in db.scalars(select(Comparison).order_by(Comparison.id.desc())) if not user or user.role == "admin" or c.subject_id in user.subject_ids]


def result_block(result):
    return next((item for item in result.evidence if isinstance(item, dict) and item.get("id") == "block"), None)


def result_is_stale(db, result, tree_version):
    block = result_block(result)
    if not block:
        return result.cache_key != pair_key(db.get(Question, result.target_id), db.get(Question, result.reference_id), tree_version)
    base = db.get(Question, result.target_id)
    references = [db.get(Question, q["id"]) for q in block.get("questions", [])]
    return not base or not references or any(q is None for q in references) or result.cache_key != block_key(base, references, block["unit"], tree_version)


@app.get("/api/comparisons/{cid}", dependencies=[Depends(require_access)])
async def comparison_detail(cid: int, db: Db = Depends(db_session)):
    comparison = comparison_or_404(db, cid)
    results = db.scalars(select(PairResult).where(PairResult.comparison_id == cid).order_by(PairResult.id)).all()
    subject = subject_or_404(db, comparison.subject_id)
    job = next((item for item in db.scalars(select(BackgroundJob).where(BackgroundJob.kind == "comparison").order_by(BackgroundJob.id.desc())) if item.payload.get("comparison_id") == cid), None)
    return {
        **comparison_out(comparison),
        "job": {"id": job.id, "status": job.status, "progress": job.progress, "total": job.total, "error": job.error} if job else None,
        "results": [{"id": r.id, "target": question_out(db.get(Question, r.target_id)), "reference": question_out(db.get(Question, r.reference_id)), "label": r.label, "reason": r.reason, "evidence": r.evidence, "source": r.source, "block": result_block(r), "stale": result_is_stale(db, r, subject.tree_version)} for r in results],
    }


class ResultFeedbackInput(BaseModel):
    comment: str = Field(min_length=1, max_length=2000)
    suggested_label: str | None = None


class ResultResolutionInput(BaseModel):
    label: str
    reason: str = Field(min_length=1, max_length=2000)


@app.post("/api/results/{result_id}/feedback", dependencies=[Depends(require_access)])
async def record_result_feedback(result_id: int, body: ResultFeedbackInput, db: Db = Depends(db_session)):
    result = db.get(PairResult, result_id)
    if not result:
        raise HTTPException(404, "Comparison result not found")
    comparison = comparison_or_404(db, result.comparison_id)
    if not body.comment.strip():
        raise HTTPException(422, "Feedback comment cannot be blank")
    if body.suggested_label is not None and body.suggested_label not in ("A", "B", "C", "review"):
        raise HTTPException(422, "Suggested label must be A, B, C or review")
    feedback = AuditEvent(action="comparison_feedback", details={
        "comparison_id": comparison.id,
        "result_id": result.id,
        "target_id": result.target_id,
        "reference_id": result.reference_id,
        "original_label": result.label,
        "suggested_label": body.suggested_label,
        "comment": body.comment.strip(),
        "by": current_user.get().id,
    })
    db.add(feedback)
    db.commit()
    return {"recorded": True, "feedback_id": feedback.id, "original_label": result.label, "suggested_label": body.suggested_label}


@app.post("/api/results/{result_id}/resolve", dependencies=[Depends(require_access)])
async def resolve_result(result_id: int, body: ResultResolutionInput, db: Db = Depends(db_session)):
    """Resolve one review case; keep the model result in the audit and never propagate it."""
    if body.label not in ("A", "B", "C"):
        raise HTTPException(422, "A reviewed classification must be A, B or C")
    result = db.get(PairResult, result_id)
    if not result:
        raise HTTPException(404, "Comparison result not found")
    comparison = comparison_or_404(db, result.comparison_id)
    if result.label != "review" and result.source != "faculty_case_review":
        raise HTTPException(409, "Only an unresolved or previously reviewed case can be adjudicated")
    old_label, old_reason, old_source = result.label, result.reason, result.source
    result.label = body.label
    result.reason = body.reason.strip()
    result.source = "faculty_case_review"
    db.add(AuditEvent(action="comparison_case_resolution", details={
        "comparison_id": comparison.id, "result_id": result.id,
        "target_id": result.target_id, "reference_id": result.reference_id,
        "previous_label": old_label, "previous_reason": old_reason,
        "previous_source": old_source, "label": body.label,
        "reason": body.reason.strip(), "by": current_user.get().id,
        "propagation": "none",
    }))
    db.commit()
    return {"resolved": True, "result_id": result.id, "label": result.label, "reason": result.reason, "source": result.source, "propagation": "none"}


@app.get("/api/comparisons/{cid}/activity", dependencies=[Depends(require_access)])
async def comparison_activity(cid: int, after_id: int = 0, limit: int = 100, db: Db = Depends(db_session)):
    comparison_or_404(db, cid)
    events = db.scalars(select(ComparisonActivity).where(ComparisonActivity.comparison_id == cid, ComparisonActivity.id > max(0, after_id)).order_by(ComparisonActivity.id).limit(min(max(limit, 1), 200))).all()
    return [{"id": event.id, "comparison_id": event.comparison_id, "target_id": event.target_id, "reference_id": event.reference_id, "agent": event.agent, "event": event.event, "payload": event.payload, "created_at": event.created_at} for event in events]


@app.get("/api/comparisons/{cid}/events", dependencies=[Depends(require_access)])
async def stream_comparison(cid: int, request: Request, after_id: int = 0, db: Db = Depends(db_session)):
    comparison_or_404(db, cid)
    try:
        cursor = max(0, after_id, int(request.headers.get("last-event-id", "0")))
    except ValueError:
        raise HTTPException(422, "Last-Event-ID must be an integer")

    async def events():
        nonlocal cursor
        # Reconnect periodically to recheck session/subject authorization.
        for _ in range(1200):
            if await request.is_disconnected():
                break
            with Session() as session:
                comparison = session.get(Comparison, cid)
                if not comparison:
                    yield sse("complete", {"status": "deleted"})
                    break
                rows = await comparison_activity(cid, cursor, 200, session)
                status = comparison.status
            for row in rows:
                cursor = row["id"]
                yield sse("activity", row, cursor)
            if status not in {"queued", "running"} and len(rows) < 200:
                yield sse("complete", {"status": status})
                break
            if not rows:
                yield ": keepalive\n\n"
                await asyncio.sleep(0.5)

    return StreamingResponse(events(), media_type="text/event-stream", headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


class TemplateInput(BaseModel):
    subject_id: int | None = None
    kind: str
    name: str = Field(min_length=1, max_length=200)
    definition: dict


RESULT_TEMPLATE_COLUMNS = {
    "base_paper": "Base paper", "base_question_number": "Base question", "base_question": "Base question text",
    "reference_paper": "Reference paper", "reference_question_number": "Reference question",
    "reference_question": "Reference question text", "label": "Classification", "reason": "Reason",
    "evidence_ids": "Evidence IDs", "stale": "Stale",
}


def normalized_question_label(value: str) -> str:
    label = re.sub(r"[^a-z0-9]", "", value.casefold())
    if label.startswith("question"):
        return label[8:]
    if label.startswith("q") and len(label) > 1 and label[1].isdigit():
        return label[1:]
    return label


def validate_template(kind: str, definition: dict) -> list[str]:
    errors = []
    if not isinstance(definition, dict):
        return ["Template definition must be a JSON object"]
    if kind == "result":
        columns = definition.get("columns", list(RESULT_TEMPLATE_COLUMNS))
        if not isinstance(columns, list) or not columns or any(not isinstance(column, str) or column not in RESULT_TEMPLATE_COLUMNS for column in columns):
            errors.append("Result template columns must be a non-empty list of supported column keys")
        title = definition.get("title")
        if title is not None and (not isinstance(title, str) or len(title) > 200):
            errors.append("Result template title must be at most 200 characters")
    elif kind == "paper":
        sections = definition.get("sections")
        if sections is not None:
            if not isinstance(sections, list):
                errors.append("Paper template sections must be an array")
            else:
                names = [section.get("name") for section in sections if isinstance(section, dict)]
                if len(names) != len(sections) or any(not isinstance(name, str) or not name.strip() or len(name) > 150 for name in names):
                    errors.append("Every paper section needs a name")
                if any(not isinstance(section.get("slots", []), list) for section in sections if isinstance(section, dict)):
                    errors.append("Paper section slots must be an array")
                    return errors
                slots = [slot for section in sections if isinstance(section, dict) for slot in section.get("slots", [])]
                for slot in slots:
                    if not isinstance(slot, dict) or not isinstance(slot.get("count", 1), int) or isinstance(slot.get("count", 1), bool) or not 1 <= slot.get("count", 1) <= 100:
                        errors.append("Paper slots need an integer count from 1 through 100")
                        break
                    if slot.get("unit") is not None and (not isinstance(slot["unit"], int) or isinstance(slot["unit"], bool) or slot["unit"] not in range(1, 6)):
                        errors.append("Paper slot unit must be 1 through 5")
                        break
                    if slot.get("marks") is not None and (not isinstance(slot["marks"], int) or isinstance(slot["marks"], bool) or slot["marks"] <= 0):
                        errors.append("Paper slot marks must be a positive whole number")
                        break
                    if slot.get("name") is not None and (not isinstance(slot["name"], str) or len(slot["name"]) > 150):
                        errors.append("Paper slot names must be at most 150 characters")
                        break
                    topic_ids = slot.get("topic_ids", [])
                    if not isinstance(topic_ids, list) or len(topic_ids) > 3 or any(not isinstance(topic_id, str) for topic_id in topic_ids) or len(topic_ids) != len(set(topic_ids)):
                        errors.append("Paper slot topic_ids must contain up to three unique topic IDs")
                        break
                    if slot.get("choice_group") is not None and (not isinstance(slot["choice_group"], str) or len(slot["choice_group"]) > 80):
                        errors.append("Choice group names must be at most 80 characters")
                        break
                    if slot.get("choice_rule") not in (None, "all_required", "choose_any"):
                        errors.append("Paper slot choice_rule must be all_required or choose_any")
                        break
                    if slot.get("choice_rule") == "choose_any":
                        choose_count = slot.get("choose_count")
                        if not isinstance(slot.get("choice_group"), str) or not slot["choice_group"].strip() or not isinstance(choose_count, int) or isinstance(choose_count, bool) or not 1 <= choose_count <= slot["count"]:
                            errors.append("choose_any slots need a choice_group and choose_count between 1 and the slot count")
                            break
    elif kind == "question":
        fields = definition.get("required_fields", [])
        if not isinstance(fields, list) or any(not isinstance(field, str) or field not in {"number", "text", "marks", "unit", "topic_ids"} for field in fields):
            errors.append("Question required_fields may contain number, text, marks, unit and topic_ids")
    elif kind == "comparison":
        target = definition.get("target_rule", "selected_paper")
        reference = definition.get("reference_rule", "all_reference")
        if not isinstance(target, str) or target not in {"selected_paper", "selected_questions", "within_paper", "whole_bank"}:
            errors.append("Comparison target_rule must be selected_paper, selected_questions, within_paper or whole_bank")
        if not isinstance(reference, str) or reference not in {"all_reference", "last_n_reference", "bank", "same_paper", "all_bank"}:
            errors.append("Comparison reference_rule must be all_reference, last_n_reference, bank, same_paper or all_bank")
        lookback = definition.get("last_n", 5)
        if not isinstance(lookback, int) or isinstance(lookback, bool) or not 1 <= lookback <= 100:
            errors.append("Comparison last_n must be an integer from 1 through 100")
        if not isinstance(definition.get("include_within_target", False), bool):
            errors.append("include_within_target must be true or false")
        if not isinstance(definition.get("require_selected_references", False), bool):
            errors.append("require_selected_references must be true or false")
        pattern = definition.get("question_pattern")
        if pattern is not None:
            if not isinstance(pattern, dict) or pattern.get("granularity") not in {"question", "subquestion"}:
                errors.append("Question pattern needs a granularity of question or subquestion")
            else:
                units = pattern.get("units", [])
                if not isinstance(units, list) or len(units) > 5:
                    errors.append("Question pattern units must be a list of up to five units")
                else:
                    seen_units = set()
                    seen_question_labels = set()
                    for unit in units:
                        if not isinstance(unit, dict) or not isinstance(unit.get("unit"), int) or isinstance(unit.get("unit"), bool) or unit["unit"] not in range(1, 6) or unit["unit"] in seen_units:
                            errors.append("Question pattern units must be unique numbers from 1 through 5")
                            break
                        seen_units.add(unit["unit"])
                        questions = unit.get("questions", [])
                        if not isinstance(questions, list) or not questions or len(questions) > 100 or any(not isinstance(number, str) or not number.strip() or len(number) > 30 for number in questions):
                            errors.append("Every question pattern unit needs 1 through 100 question labels")
                            break
                        if len({number.strip().casefold() for number in questions}) != len(questions):
                            errors.append("Question labels must be unique within each pattern unit")
                            break
                        normalized_labels = {normalized_question_label(number) for number in questions}
                        if len(normalized_labels) != len(questions):
                            errors.append("Question labels must be unique within each pattern unit")
                            break
                        if seen_question_labels.intersection(normalized_labels):
                            errors.append("Question labels must belong to only one unit")
                            break
                        seen_question_labels.update(normalized_labels)
                        choices = unit.get("choices", [])
                        if not isinstance(choices, list) or any(not isinstance(group, list) or len(group) < 2 or any(not isinstance(option, list) or not option or any(not isinstance(number, str) or number not in questions for number in option) for option in group) for group in choices):
                            errors.append("Each OR group needs at least two alternatives made from labels in that unit")
                            break
        report_id = definition.get("report_template_id")
        if report_id is not None and (not isinstance(report_id, int) or isinstance(report_id, bool) or report_id < 1):
            errors.append("report_template_id must be a positive integer")
    return errors


def validate_template_topics(kind: str, definition: dict, subject: Subject | None) -> list[str]:
    if kind != "paper" or subject is None:
        return []
    sections = definition.get("sections", [])
    if not isinstance(sections, list):
        return []
    topic_units = {node["id"]: unit_for_topic(node["id"]) for node in subject.tree.get("nodes", []) if node.get("kind") == "topic"}
    errors = []
    for section_index, section in enumerate(sections, 1):
        if not isinstance(section, dict):
            continue
        for slot_index, slot in enumerate(section.get("slots", []), 1):
            if not isinstance(slot, dict) or not slot.get("topic_ids"):
                continue
            unit = slot.get("unit")
            if unit not in range(1, 6) or any(topic_id not in topic_units or topic_units[topic_id] != unit for topic_id in slot["topic_ids"]):
                errors.append(f"Section {section_index}, slot {slot_index} must select real topics that match its unit")
    return errors


@app.get("/api/templates", dependencies=[Depends(require_access)])
async def templates(db: Db = Depends(db_session)):
    user = current_user.get()
    return [{"id": t.id, "subject_id": t.subject_id, "kind": t.kind, "name": t.name, "definition": t.definition, "version": t.version} for t in db.scalars(select(Template).order_by(Template.kind, Template.name)) if not user or user.role == "admin" or t.subject_id is None or t.subject_id in user.subject_ids]


class ComparisonTemplateRun(BaseModel):
    paper_id: int | None = None
    question_ids: list[int] = Field(default_factory=list)
    reference_paper_ids: list[int] = Field(default_factory=list, max_length=20)
    title: str | None = Field(default=None, max_length=250)
    preview: bool = False


@app.post("/api/templates/import-report", dependencies=[Depends(require_access)])
async def import_report_template(file: UploadFile = File(...), subject_id: int | None = Form(None), db: Db = Depends(db_session)):
    if subject_id is not None:
        subject_or_404(db, subject_id)
    filename = Path(file.filename or "").name
    suffix = Path(filename).suffix.lower()
    if suffix not in {".txt", ".xlsx"}:
        raise HTTPException(422, "Choose a .txt or .xlsx report template")
    content = await file.read(2_000_001)
    if len(content) > 2_000_000:
        raise HTTPException(413, "Template must be 2 MB or smaller")
    try:
        if suffix == ".txt":
            lines = [line.strip() for line in content.decode("utf-8-sig").splitlines() if line.strip()]
            if not lines:
                raise ValueError("File is empty")
            fields = re.split(r"[|,\t]", lines[0]) if re.search(r"[|,\t]", lines[0]) else lines
        else:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            try:
                fields = list(next(workbook.active.iter_rows(min_row=1, max_row=1, values_only=True), ()))
            finally:
                workbook.close()
        lookup = {re.sub(r"[^a-z0-9]+", "", value.lower()): key for key, value in RESULT_TEMPLATE_COLUMNS.items()}
        fields = [str(value).strip() for value in fields if value is not None and str(value).strip()]
        columns = [field.lower() if field.lower() in RESULT_TEMPLATE_COLUMNS else lookup.get(re.sub(r"[^a-z0-9]+", "", field.lower())) for field in fields]
        if not columns or any(column is None for column in columns) or len(set(columns)) != len(columns):
            raise ValueError("Use unique supported report column names")
    except (ValueError, UnicodeError, OSError, KeyError, BadZipFile) as exc:
        raise HTTPException(422, f"Invalid report template: {exc}") from exc
    template = Template(subject_id=subject_id, kind="result", name=Path(filename).stem[:200] or "Report template", definition={"columns": columns})
    db.add(template)
    db.add(AuditEvent(action="template_create", details={"kind": "result", "name": template.name, "source": filename}))
    db.commit()
    return {"id": template.id, "name": template.name, "columns": columns}


@app.post("/api/templates/{tid}/run", dependencies=[Depends(require_access)])
async def run_comparison_template(tid: int, body: ComparisonTemplateRun, db: Db = Depends(db_session)):
    template = db.get(Template, tid)
    if not template or template.kind != "comparison":
        raise HTTPException(404, "Comparison template not found")
    definition = template.definition or {}
    errors = validate_template("comparison", definition)
    if errors:
        raise HTTPException(409, errors)
    if template.subject_id:
        subject_or_404(db, template.subject_id)
    subject_id = template.subject_id
    if body.paper_id:
        paper = paper_or_404(db, body.paper_id)
        if paper.archived:
            raise HTTPException(422, "Archived papers cannot be used in a new comparison scope")
        if subject_id and paper.subject_id != subject_id:
            raise HTTPException(422, "Paper does not belong to this template subject")
        subject_id = paper.subject_id
    elif subject_id is None:
        raise HTTPException(422, "Choose a paper so the shared template can resolve its subject")
    subject = subject_or_404(db, subject_id)
    papers = db.scalars(select(Paper).where(Paper.subject_id == subject_id, Paper.archived.is_(False))).all()
    by_id = {paper.id: paper for paper in papers}
    target_rule = definition.get("target_rule", "selected_paper")
    reference_rule = definition.get("reference_rule", "all_reference")
    if definition.get("require_selected_references") and not body.reference_paper_ids:
        raise HTTPException(422, "Select the comparison papers for this template run")
    if target_rule == "whole_bank":
        target_papers = [paper for paper in papers if paper.status == "bank"]
        target_ids = [q.id for paper in target_papers for q in paper.questions if q.confirmed]
    elif target_rule == "selected_questions":
        if not body.question_ids:
            raise HTTPException(422, "Choose question IDs to run this selected-questions template")
        target_ids = list(dict.fromkeys(body.question_ids))
        target_questions = db.scalars(select(Question).where(Question.id.in_(target_ids))).all()
        if len(target_questions) != len(target_ids) or any(q.paper.subject_id != subject_id or q.paper.archived for q in target_questions):
            raise HTTPException(422, "Every selected question must belong to this subject and an active paper")
    else:
        if not body.paper_id:
            raise HTTPException(422, "Choose a paper to run this comparison template")
        target_ids = [q.id for q in by_id[body.paper_id].questions if q.confirmed]
    if body.reference_paper_ids:
        if reference_rule not in {"all_reference", "last_n_reference"} or target_rule in {"within_paper", "whole_bank"}:
            raise HTTPException(422, "This comparison pattern does not accept selected reference papers")
        if len(set(body.reference_paper_ids)) != len(body.reference_paper_ids):
            raise HTTPException(422, "Choose each reference paper only once")
        reference_papers = [by_id.get(paper_id) for paper_id in body.reference_paper_ids]
        if any(paper is None or paper.status != "reference" or paper.archived for paper in reference_papers):
            raise HTTPException(422, "Every selected comparison paper must be an active reference paper in this subject")
        if body.paper_id and body.paper_id in body.reference_paper_ids:
            raise HTTPException(422, "The base paper cannot also be a comparison paper")
    elif reference_rule == "same_paper" or target_rule == "within_paper":
        if not body.paper_id:
            raise HTTPException(422, "Choose the paper to compare within")
        reference_papers = [by_id[body.paper_id]]
    elif reference_rule == "bank" or reference_rule == "all_bank":
        reference_papers = [paper for paper in papers if paper.status == "bank"]
    else:
        reference_papers = [paper for paper in papers if paper.status == "reference"]
        if reference_rule == "last_n_reference":
            reference_papers.sort(key=lambda item: (item.year or 0, item.id), reverse=True)
            reference_papers = reference_papers[:definition.get("last_n", 5)]
    if definition.get("include_within_target", False) and body.paper_id and target_rule not in {"within_paper", "whole_bank"}:
        if all(item.id != body.paper_id for item in reference_papers):
            reference_papers.append(by_id[body.paper_id])
    if not definition.get("include_within_target", False) and body.paper_id and target_rule != "within_paper":
        reference_papers = [paper for paper in reference_papers if paper.id != body.paper_id]
    reference_ids = [q.id for paper in reference_papers for q in paper.questions if q.confirmed]
    if target_rule == "whole_bank" and reference_rule == "all_bank":
        target_ids = reference_ids = [q.id for paper in reference_papers for q in paper.questions if q.confirmed]
    pattern_issues = []
    pattern = definition.get("question_pattern")
    if pattern and target_rule == "selected_paper" and body.paper_id:
        label_units = {
            normalized_question_label(label): unit["unit"]
            for unit in pattern.get("units", [])
            for label in unit.get("questions", [])
        }

        def matching_pattern_questions(paper):
            found = {}
            matching = []
            for question in paper.questions:
                key = normalized_question_label(question.number)
                expected_unit = label_units.get(key)
                if expected_unit is None:
                    if question.confirmed and question.unit in {unit["unit"] for unit in pattern.get("units", [])}:
                        pattern_issues.append(f"{paper.title}: Q{question.number} is outside the saved pattern")
                    continue
                found.setdefault(key, []).append(question)
                if not question.confirmed:
                    pattern_issues.append(f"{paper.title}: Q{question.number} needs topic confirmation")
                elif question.unit != expected_unit:
                    pattern_issues.append(f"{paper.title}: Q{question.number} is tagged Unit {question.unit}, expected Unit {expected_unit}")
                else:
                    matching.append(question.id)
            for label, expected_unit in label_units.items():
                if label not in found:
                    pattern_issues.append(f"{paper.title}: expected question {label} was not extracted (Unit {expected_unit})")
                elif len(found[label]) > 1:
                    pattern_issues.append(f"{paper.title}: question label {label} was extracted more than once")
            return matching

        target_ids = matching_pattern_questions(by_id[body.paper_id])
        reference_ids = [question_id for paper in reference_papers for question_id in matching_pattern_questions(paper)]
    if not target_ids or not reference_ids:
        if not body.preview:
            raise HTTPException(422, "The selected pattern has no confirmed target or reference questions")
    if len(target_ids) * len(reference_ids) > 10_000:
        raise HTTPException(422, "This scope exceeds 10,000 possible pairs; narrow the selected papers or questions")
    request = CompareInput(subject_id=subject_id, title=body.title or template.name, target_ids=target_ids, reference_ids=reference_ids)
    pairs, skipped = resolved_pairs(db, request) if target_ids and reference_ids else ([], [])
    if pattern_issues and not body.preview:
        raise HTTPException(422, "Pattern checks need attention: " + "; ".join(pattern_issues[:12]))
    if not pairs and not body.preview:
        raise HTTPException(422, "This template resolved to no same-unit pairs")
    block_strategy = bool(pattern and target_rule == "selected_paper" and body.paper_id and all(p.id != body.paper_id for p in reference_papers))
    checks = len(group_comparison_blocks(pairs, {})) if block_strategy else len(pairs)
    if body.preview:
        return {"checks": checks, "template_id": tid, "subject_id": subject_id, "targets": len(set(target_ids)), "references": len(set(reference_ids)), "pairs": len(pairs), "skipped": skipped, "pattern_issues": pattern_issues, "within_target": bool(definition.get("include_within_target", False))}
    if not (BLOCK_JUDGE_MODEL if block_strategy else classifier_models_ready()):
        raise HTTPException(503, "Configure the local comparison model before starting")
    inventory = await inspect_local_models()
    if not (inventory.get("judge_installed") if block_strategy else inventory.get("roles_available")):
        raise HTTPException(503, "Install the configured local comparison model before starting")
    report_snapshot = None
    report_id = definition.get("report_template_id")
    if report_id:
        report_template = db.get(Template, report_id)
        if not report_template or report_template.kind != "result" or report_template.subject_id not in (None, subject_id):
            raise HTTPException(422, "The comparison template refers to an unavailable report template")
        report_errors = validate_template("result", report_template.definition)
        if report_errors:
            raise HTTPException(409, "The linked report template is invalid")
        report_snapshot = {"id": report_template.id, "name": report_template.name, "version": report_template.version, "definition": report_template.definition}
    snapshot = {"strategy": BLOCK_VERSION if block_strategy else None, "comparison_template": {"id": template.id, "name": template.name, "version": template.version, "definition": definition}, "report_template": report_snapshot}
    comparison = Comparison(**request.model_dump(), total=checks, template_snapshot=snapshot)
    db.add(comparison)
    db.flush()
    job = enqueue_job(db, "comparison", {"comparison_id": comparison.id}, total=checks)
    db.add(AuditEvent(action="comparison_template_run", details={"template_id": tid, "version": template.version, "comparison_id": comparison.id, "subject_id": subject.id, "pairs": len(pairs)}))
    db.commit()
    return {**comparison_out(comparison), "job_id": job.id, "template_id": tid, "resolved_pairs": len(pairs), "checks": checks}


@app.post("/api/templates/import-paper-sample", dependencies=[Depends(require_access)])
async def import_paper_sample(subject_id: int = Form(...), file: UploadFile = File(...), db: Db = Depends(db_session)):
    subject_or_404(db, subject_id)
    filename = Path(file.filename or "paper-sample").name
    data = await file.read(25 * 1024 * 1024 + 1)
    try:
        text, warnings = await extract_text(filename, data)
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc
    digest = hashlib.sha256(data).hexdigest()
    for existing in db.scalars(select(Template).where(Template.subject_id == subject_id, Template.kind == "paper")):
        if (existing.definition or {}).get("source", {}).get("sha256") == digest:
            return {"id": existing.id, "subject_id": subject_id, "kind": existing.kind, "name": existing.name, "definition": existing.definition, "version": existing.version, "warnings": warnings, "duplicate": True}
    definition = suggest_paper_pattern(text, filename, digest)
    errors = validate_template("paper", definition)
    if errors:
        raise HTTPException(422, errors)
    storage = ROOT / "data" / "uploads"
    storage.mkdir(parents=True, exist_ok=True)
    target = storage / digest
    if not target.exists():
        target.write_bytes(data)
    template = Template(subject_id=subject_id, kind="paper", name=f"{Path(filename).stem[:170]} pattern", definition=definition)
    db.add(template)
    db.add(AuditEvent(action="paper_template_import", details={"subject_id": subject_id, "filename": filename, "hash": digest}))
    db.commit()
    db.refresh(template)
    return {"id": template.id, "subject_id": template.subject_id, "kind": template.kind, "name": template.name, "definition": template.definition, "version": template.version, "warnings": warnings, "duplicate": False}


@app.post("/api/templates", dependencies=[Depends(require_access)])
async def save_template(body: TemplateInput, db: Db = Depends(db_session)):
    if body.kind not in ("paper", "question", "result", "comparison"):
        raise HTTPException(422, "Template kind must be paper, question, result or comparison")
    if body.subject_id:
        subject_or_404(db, body.subject_id)
    errors = validate_template(body.kind, body.definition)
    errors.extend(validate_template_topics(body.kind, body.definition, db.get(Subject, body.subject_id) if body.subject_id else None))
    if errors:
        raise HTTPException(422, errors)
    template = Template(**body.model_dump())
    db.add(template)
    db.add(AuditEvent(action="template_create", details={"kind": body.kind, "name": body.name}))
    db.commit()
    return {"id": template.id, "version": template.version}


@app.put("/api/templates/{tid}", dependencies=[Depends(require_access)])
async def update_template(tid: int, body: TemplateInput, db: Db = Depends(db_session)):
    template = db.get(Template, tid)
    if not template:
        raise HTTPException(404, "Template not found")
    if template.subject_id:
        subject_or_404(db, template.subject_id)
    if body.subject_id:
        subject_or_404(db, body.subject_id)
    if body.kind not in ("paper", "question", "result", "comparison"):
        raise HTTPException(422, "Template kind must be paper, question, result or comparison")
    errors = validate_template(body.kind, body.definition)
    errors.extend(validate_template_topics(body.kind, body.definition, db.get(Subject, body.subject_id) if body.subject_id else None))
    if errors:
        raise HTTPException(422, errors)
    template.subject_id, template.kind, template.name, template.definition = body.subject_id, body.kind, body.name, body.definition
    template.version += 1
    db.add(AuditEvent(action="template_update", details={"id": tid, "version": template.version}))
    db.commit()
    return {"id": tid, "version": template.version}


@app.delete("/api/templates/{tid}", dependencies=[Depends(require_access)])
async def delete_template(tid: int, db: Db = Depends(db_session)):
    template = db.get(Template, tid)
    if not template:
        raise HTTPException(404, "Template not found")
    if template.subject_id:
        subject_or_404(db, template.subject_id)
    source = (template.definition or {}).get("source")
    digest = source.get("sha256") if isinstance(source, dict) else None
    db.add(AuditEvent(action="template_delete", details={"id": tid, "name": template.name, "definition": template.definition, "by": current_user.get().id}))
    db.delete(template)
    db.commit()
    unlink_if_unused(db, digest)
    return {"deleted": tid}


@app.get("/api/materials", dependencies=[Depends(require_access)])
async def materials(subject_id: int | None = None, db: Db = Depends(db_session)):
    query = select(Material).order_by(Material.id.desc())
    if subject_id:
        subject_or_404(db, subject_id)
        query = query.where(Material.subject_id == subject_id)
    user = current_user.get()
    return [{"id": m.id, "subject_id": m.subject_id, "filename": m.filename, "warnings": m.warnings, "proposals": m.proposals, "enrichment": {key: value for key, value in (m.enrichment or {}).items() if key != "_section_cache"}, "pages": len(m.extracted_text.split("\f"))} for m in db.scalars(query) if not user or user.role == "admin" or m.subject_id in user.subject_ids]


@app.get("/api/materials/{mid}/review", dependencies=[Depends(require_access)])
async def material_review(mid: int, db: Db = Depends(db_session)):
    material = db.get(Material, mid)
    if not material:
        raise HTTPException(404, "Material not found")
    subject = subject_or_404(db, material.subject_id)
    labels = {node["id"]: node["label"] for node in subject.tree["nodes"]}
    pages = material.extracted_text.split("\f")
    return {
        "id": material.id, "filename": material.filename, "warnings": material.warnings,
        "pages": [{"page": number, "preview": page.strip()[:1200], "characters": len(page.strip())} for number, page in enumerate(pages, 1)],
        "proposals": [{**proposal, "topic_id": topic_id, "topic_label": labels.get(topic_id, topic_id)} for topic_id, proposal in (material.proposals or {}).items()],
        "enrichment": {key: value for key, value in (material.enrichment or {}).items() if key != "_section_cache"},
    }


@app.get("/api/materials/{mid}/events", dependencies=[Depends(require_access)])
async def stream_material_activity(mid: int, request: Request, after_id: int = 0, db: Db = Depends(db_session)):
    material = db.get(Material, mid)
    if not material:
        raise HTTPException(404, "Material not found")
    subject_or_404(db, material.subject_id)
    try:
        cursor = max(0, after_id, int(request.headers.get("last-event-id", "0")))
    except ValueError:
        raise HTTPException(422, "Last-Event-ID must be an integer")

    async def events():
        nonlocal cursor
        for _ in range(1200):
            if await request.is_disconnected():
                break
            with Session() as session:
                current = session.get(Material, mid)
                if not current:
                    yield sse("complete", {"status": "deleted"})
                    break
                state = current.enrichment or {}
                rows = [item for item in state.get("activity", []) if item.get("id", 0) > cursor]
                status = state.get("status", "not_started")
            for row in rows:
                cursor = row["id"]
                yield sse("activity", row, cursor)
            if status not in {"queued", "running"}:
                yield sse("complete", {"status": status})
                break
            if not rows:
                yield ": keepalive\n\n"
                await asyncio.sleep(0.6)
    return StreamingResponse(events(), media_type="text/event-stream", headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


@app.get("/api/materials/{mid}/pages/{page}", dependencies=[Depends(require_access)])
async def material_page(mid: int, page: int, db: Db = Depends(db_session)):
    material = db.get(Material, mid)
    if not material:
        raise HTTPException(404, "Material not found")
    subject_or_404(db, material.subject_id)
    pages = material.extracted_text.split("\f")
    if page < 1 or page > len(pages):
        raise HTTPException(404, "Page not found")
    return {"page": page, "text": pages[page - 1]}


class MaterialProposalRevision(BaseModel):
    topic_id: str
    summary: str | None = None
    included: bool | None = None


@app.post("/api/materials/{mid}/proposal/revise", dependencies=[Depends(require_access)])
async def revise_material_proposal(mid: int, body: MaterialProposalRevision, db: Db = Depends(db_session)):
    material = db.get(Material, mid)
    if not material:
        raise HTTPException(404, "Material not found")
    subject_or_404(db, material.subject_id)
    if (material.enrichment or {}).get("status") != "review_ready":
        raise HTTPException(409, "Prepare a mapping preview before revising it")
    proposals = json.loads(json.dumps(material.proposals or {}))
    proposal = proposals.get(body.topic_id)
    if not proposal:
        raise HTTPException(404, "Mapped topic not found")
    if body.summary is not None:
        summary = body.summary.strip()
        if len(summary) > 500 or len(summary.split()) > 45:
            raise HTTPException(422, "Suggestion must be at most 45 words and 500 characters")
        proposal["summaries"] = [summary] if summary else []
    if body.included is not None:
        proposal["included"] = body.included
    material.proposals = proposals
    db.add(AuditEvent(action="material_mapping_revised", details={"material_id": mid, "topic_id": body.topic_id, "included": proposal.get("included", True), "by": current_user.get().id}))
    db.commit()
    return {"id": mid, "topic_id": body.topic_id, "proposal": proposal}


@app.post("/api/materials/{mid}/apply", dependencies=[Depends(require_access)])
async def apply_material_mapping(mid: int, db: Db = Depends(db_session)):
    material = db.get(Material, mid)
    if not material:
        raise HTTPException(404, "Material not found")
    subject = subject_or_404(db, material.subject_id)
    enrichment = material.enrichment or {}
    if enrichment.get("status") != "review_ready":
        raise HTTPException(409, "Prepare a mapping before applying it")
    result = _apply_material_mapping(db, material, subject, current_user.get().id)
    db.commit()
    return result


def _apply_material_mapping(db, material, subject, actor_id=None):
    mid = material.id
    enrichment = material.enrichment or {}
    if enrichment.get("source_tree_version") != subject.tree_version:
        raise HTTPException(409, "The syllabus changed after this mapping was prepared; prepare it again")
    proposals = {key: value for key, value in (material.proposals or {}).items() if value.get("included", True)}
    if not proposals:
        raise HTTPException(422, "No included topic mappings to apply")
    tree = json.loads(json.dumps(subject.tree))
    nodes = {node["id"]: node for node in tree["nodes"] if node.get("kind") == "topic"}
    pages = material.extracted_text.split("\f")
    for topic_id, proposal in proposals.items():
        if topic_id not in nodes:
            raise HTTPException(409, f"Topic {topic_id} is no longer in the syllabus")
        citations = proposal.get("citations", [])
        if not citations:
            raise HTTPException(422, f"Mapping for {topic_id} has no source citation")
        for citation in citations:
            page = citation.get("page")
            quote = citation.get("quote", "")
            if not isinstance(page, int) or page < 1 or page > len(pages) or not isinstance(quote, str) or not quote.strip() or " ".join(re.findall(r"[\w]+", quote.casefold())) not in " ".join(re.findall(r"[\w]+", pages[page - 1].casefold())):
                raise HTTPException(422, f"Mapping for {topic_id} has an unsupported source quote")
        node = nodes[topic_id]
        for field in ("context_words", "scope_terms", "supporting_examples"):
            node[field] = list({json.dumps(item, sort_keys=True): item for item in node.get(field, []) + proposal.get(field, [])}.values())
        for summary in proposal.get("summaries", []):
            note = {"text": summary, "citations": citations, "source": material.filename, "material_id": mid}
            notes = node.setdefault("material_notes", [])
            if not any(isinstance(existing, dict) and existing.get("material_id") == mid and existing.get("text") == summary for existing in notes):
                notes.append(note)
        node["material_citations"] = list({json.dumps(item, sort_keys=True): item for item in node.get("material_citations", []) + citations}.values())
        node.setdefault("review_status", "enriched")
    db.add(TreeRevision(subject_id=subject.id, version=subject.tree_version, tree=subject.tree, reason=f"Before teaching material {mid}"))
    subject.tree = tree
    subject.tree_version += 1
    db.add(TreeRevision(subject_id=subject.id, version=subject.tree_version, tree=tree, reason=f"Teaching material {mid} mapping"))
    material.enrichment = {**enrichment, "status": "applied", "tree_version": subject.tree_version}
    db.add(AuditEvent(action="material_mapping_applied", details={"material_id": mid, "topics": list(proposals), "tree_version": subject.tree_version, "by": actor_id}))
    return {"id": mid, "status": "applied", "tree_version": subject.tree_version, "topics": list(proposals)}


@app.post("/api/subjects/{sid}/materials", dependencies=[Depends(require_access)])
async def upload_material(sid: int, file: UploadFile = File(...), db: Db = Depends(db_session)):
    subject_or_404(db, sid)
    data = await file.read(25 * 1024 * 1024 + 1)
    try:
        text, warnings = await extract_text(file.filename or "", data)
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc
    digest = hashlib.sha256(data).hexdigest()
    duplicate = db.scalar(select(Material.id).where(Material.subject_id == sid, Material.file_hash == digest))
    if duplicate:
        raise HTTPException(409, f"This material was already uploaded for this subject (material {duplicate})")
    storage = ROOT / "data" / "uploads"
    storage.mkdir(parents=True, exist_ok=True)
    target = storage / digest
    if not target.exists():
        target.write_bytes(data)
    material = Material(subject_id=sid, filename=Path(file.filename or "upload").name, file_hash=digest, extracted_text=text, warnings=warnings, proposals={})
    db.add(material)
    db.commit()
    return {"id": material.id, "subject_id": sid, "filename": material.filename, "warnings": warnings, "pages": len(text.split("\f")), "proposals": {}, "enrichment": {"status": "not_started"}}


@app.post("/api/materials/{mid}/enrich", dependencies=[Depends(require_access)])
async def start_material_enrichment(mid: int, db: Db = Depends(db_session)):
    material = db.get(Material, mid)
    if not material:
        raise HTTPException(404, "Material not found")
    return await _prepare_material_mapping(db, material)


async def _prepare_material_mapping(db, material):
    subject = subject_or_404(db, material.subject_id)
    if (material.enrichment or {}).get("status") in ("queued", "running"):
        raise HTTPException(409, "This material already has a mapping in progress")
    topics = [{"id": node["id"], "label": node["label"]} for node in subject.tree["nodes"] if node.get("kind") == "topic"]
    result = await map_material(material.extracted_text, topics, material.filename)
    material.proposals = {topic_id: {**proposal, "included": True} for topic_id, proposal in result["proposals"].items()}
    material.enrichment = {key: value for key, value in result.items() if key != "proposals"} | {
        "status": "review_ready", "source_tree_version": subject.tree_version, "error": None,
    }
    db.add(AuditEvent(action="material_mapping_prepared", details={"material_id": material.id, "topics": list(material.proposals), "sections_processed": result["sections_processed"], "source_tree_version": subject.tree_version}))
    db.commit()
    return {"id": material.id, "material_id": material.id, "filename": material.filename, "status": "review_ready", "mapped_topics": len(material.proposals), "topic_matches": [
        {"topic": topic["label"], "pages": [citation["page"] for citation in material.proposals[topic["id"]]["citations"]]}
        for topic in topics if topic["id"] in material.proposals
    ]}


@app.delete("/api/materials/{mid}", dependencies=[Depends(require_access)])
async def delete_material(mid: int, db: Db = Depends(db_session)):
    material = db.get(Material, mid)
    if not material:
        raise HTTPException(404, "Material not found")
    subject_or_404(db, material.subject_id)
    digest = material.file_hash
    db.add(AuditEvent(action="material_delete", details={"id": mid, "filename": material.filename, "hash": digest, "by": current_user.get().id}))
    db.delete(material)
    db.commit()
    unlink_if_unused(db, digest)
    return {"deleted": mid}


class MaterialTopic(BaseModel):
    topic_id: str


@app.post("/api/materials/{mid}/propose", dependencies=[Depends(require_access)])
async def propose_material(mid: int, body: MaterialTopic, db: Db = Depends(db_session)):
    material = db.get(Material, mid)
    if not material:
        raise HTTPException(404, "Material not found")
    subject = subject_or_404(db, material.subject_id)
    topic = next((n for n in subject.tree["nodes"] if n["id"] == body.topic_id), None)
    if not topic:
        raise HTTPException(422, "Select an existing syllabus topic")
    if not MODEL:
        raise HTTPException(503, "Configure a local JUDGE_MODEL_NAME to draft material enrichment")
    pages = material.extracted_text.split("\f")
    terms = [topic["label"].lower()] + [x.lower() for x in topic.get("synonyms", [])]
    ranked = sorted(enumerate(pages, 1), key=lambda item: max((item[1].lower().count(term) for term in terms if term), default=0), reverse=True)
    excerpts = [{"page": number, "text": content[:3500]} for number, content in ranked[:4] if content.strip()]
    result = await model_json([
        {"role": "system", "content": "Draft a narrowly grounded syllabus enrichment proposal. Return JSON with keys description (optional string), context_words (string array), scope_terms (string array), supporting_examples (string array), cited_pages (integer array), uncertainty (string). Use only the supplied material, preserve official syllabus wording, and never turn broad topic knowledge into mandatory answer requirements."},
        {"role": "user", "content": json.dumps({"topic": topic, "material": material.filename, "excerpts": excerpts}, default=str)},
    ])
    cited = result.get("cited_pages", [])
    supplied_pages = {item["page"] for item in excerpts}
    if not isinstance(cited, list) or not cited or any(not isinstance(p, int) or p not in supplied_pages for p in cited):
        raise HTTPException(422, "Model proposal did not cite valid pages")
    description = str(result.get("description", "")).strip()[:500]
    if len(description.split()) > 45:
        raise HTTPException(422, "Description exceeds the 45 word limit")
    proposal = {
        "topic_id": body.topic_id, "material_id": mid, "included": True,
        "summaries": [description] if description else [],
        "uncertainties": [str(result.get("uncertainty", ""))[:500]] if result.get("uncertainty") else [],
        "citations": [{"file": material.filename, "page": page, "section": 1,
                       "quote": next(item["text"].strip()[:350] for item in excerpts if item["page"] == page)} for page in cited],
    }
    for field in ("context_words", "scope_terms", "supporting_examples"):
        values = result.get(field, [])
        if not isinstance(values, list) or any(not isinstance(x, str) for x in values):
            raise HTTPException(422, f"Invalid {field} in model proposal")
        proposal[field] = [x[:200] for x in values[:20]]
    if not any(citation["quote"] for citation in proposal["citations"]):
        raise HTTPException(422, "No source text was available for the cited pages")
    material.proposals = {**(material.proposals or {}), body.topic_id: proposal}
    material.enrichment = {**(material.enrichment or {}), "status": "review_ready", "source_tree_version": subject.tree_version}
    db.add(AuditEvent(action="material_mapping_prepared", details={"material_id": mid, "topic_id": body.topic_id}))
    db.commit()
    return proposal


def report_row_values(comparison: Comparison, result: PairResult, db: Db) -> dict[str, str]:
    a, b = db.get(Question, result.target_id), db.get(Question, result.reference_id)
    snapshots = {e.get("id"): e.get("text") for e in result.evidence if isinstance(e, dict)}
    stale = result_is_stale(db, result, subject_or_404(db, comparison.subject_id).tree_version)
    block = result_block(result)
    return {
        "base_paper": a.paper.title, "base_question_number": a.number, "base_question": snapshots.get("q1", a.text),
        "reference_paper": b.paper.title, "reference_question_number": ", ".join(q["number"] for q in block["questions"]) if block else b.number, "reference_question": " | ".join(f"Q{q['number']}: {q['text']}" for q in block["questions"]) if block else snapshots.get("q2", b.text),
        "label": result.label, "reason": result.reason,
        "evidence_ids": ", ".join(e.get("id", "?") for e in result.evidence if e.get("cited")),
        "stale": "Yes" if stale else "No",
    }


def report_text(comparison: Comparison, results: list[PairResult], db: Db, template: Template | None = None) -> str:
    definition = template.definition if template else {}
    columns = definition.get("columns", list(RESULT_TEMPLATE_COLUMNS))
    title = definition.get("title") or comparison.title
    lines = [title, f"Status: {comparison.status}", f"Targets: {comparison.target_ids}", f"References: {comparison.reference_ids}", f"Checks completed: {comparison.progress}/{comparison.total}", ""]
    lines.append(" | ".join(RESULT_TEMPLATE_COLUMNS[key] for key in columns))
    lines.append(" | ".join("---" for _ in columns))
    for result in results:
        values = report_row_values(comparison, result, db)
        cells = [values[key].replace("|", "\\|").replace("\r", " ").replace("\n", " ") for key in columns]
        lines.append(" | ".join(cells))
    if comparison.status != "complete":
        lines += ["", "Incomplete or failed results are shown above; this report is not final."]
    return "\n".join(lines)


@app.get("/api/comparisons/{cid}/export", dependencies=[Depends(require_access)])
async def export_comparison(cid: int, format: str = "md", template_id: int | None = None, db: Db = Depends(db_session)):
    comparison = comparison_or_404(db, cid)
    results = db.scalars(select(PairResult).where(PairResult.comparison_id == cid).order_by(PairResult.id)).all()
    template = db.get(Template, template_id) if template_id else None
    if template and (template.kind != "result" or template.subject_id not in (None, comparison.subject_id)):
        raise HTTPException(422, "Select a result template available to this subject")
    saved_report = (comparison.template_snapshot or {}).get("report_template") or {}
    report_definition = template.definition if template else saved_report.get("definition", {})
    columns = report_definition.get("columns", list(RESULT_TEMPLATE_COLUMNS))
    template_errors = validate_template("result", {"columns": columns})
    if template_errors:
        raise HTTPException(409, "The selected result template is invalid")
    report_for_export = template if template else SimpleNamespace(definition=report_definition) if saved_report else None
    text = report_text(comparison, results, db, report_for_export)
    if format in ("md", "txt"):
        data = text.encode()
        media = "text/markdown" if format == "md" else "text/plain"
    elif format == "pdf":
        document = pymupdf.open()
        lines = [part for line in text.splitlines() for part in (textwrap.wrap(line, width=94, break_long_words=True) or [""])]
        font_path = Path("/usr/share/fonts/ttf-readex-pro/ReadexPro-Regular.ttf")
        for start in range(0, len(lines), 48):
            page = document.new_page()
            if font_path.exists():
                page.insert_font(fontname="ReportFont", fontfile=str(font_path))
            y = 55
            for line in lines[start:start + 48]:
                page.insert_text((45, y), line, fontsize=10, fontname="ReportFont" if font_path.exists() else "helv")
                y += 14
        data = document.tobytes()
        media = "application/pdf"
    elif format == "xlsx":
        book = Workbook()
        sheet = book.active
        sheet.title = "Comparison"
        sheet.append([RESULT_TEMPLATE_COLUMNS[key] for key in columns])
        for r in results:
            values = report_row_values(comparison, r, db)
            safe = lambda value: "'" + value if isinstance(value, str) and value.startswith(("=", "+", "-", "@")) else value
            sheet.append([safe(values[key]) for key in columns])
        info = book.create_sheet("Report status")
        info.append(["Comparison", comparison.title])
        info.append(["Status", comparison.status])
        info.append(["Completed pairs", f"{comparison.progress}/{comparison.total}"])
        output = io.BytesIO()
        book.save(output)
        data = output.getvalue()
        media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        raise HTTPException(422, "Format must be md, txt, pdf or xlsx")
    return Response(data, media_type=media, headers={"Content-Disposition": f'attachment; filename="comparison-{cid}.{format}"'})


class ExportJobInput(BaseModel):
    format: str = "md"
    template_id: int | None = None


@app.post("/api/comparisons/{cid}/exports", dependencies=[Depends(require_access)])
async def queue_comparison_export(cid: int, body: ExportJobInput, db: Db = Depends(db_session)):
    comparison = comparison_or_404(db, cid)
    if body.format not in {"md", "txt", "pdf", "xlsx"}:
        raise HTTPException(422, "Format must be md, txt, pdf or xlsx")
    template = db.get(Template, body.template_id) if body.template_id else None
    if template and (template.kind != "result" or template.subject_id not in (None, comparison.subject_id)):
        raise HTTPException(422, "Select a result template available to this subject")
    saved_report = (comparison.template_snapshot or {}).get("report_template") or {}
    definition = template.definition if template else saved_report.get("definition", {})
    if validate_template("result", {"columns": definition.get("columns", list(RESULT_TEMPLATE_COLUMNS))}):
        raise HTTPException(409, "The selected result template is invalid")
    job = enqueue_job(db, "comparison_export", {"comparison_id": cid, "format": body.format, "template_id": body.template_id}, total=1)
    db.add(AuditEvent(action="comparison_export_queued", details={"comparison_id": cid, "job_id": job.id, "format": body.format, "by": current_user.get().id}))
    db.commit()
    return {"id": job.id, "kind": job.kind, "status": job.status, "progress": job.progress, "total": job.total}


async def _comparison_job(payload: dict, job_id: int):
    await run_comparison(int(payload["comparison_id"]), job_id)


async def _comparison_export_job(payload: dict, job_id: int):
    ensure_not_cancelled(job_id)
    comparison_id, format_name = int(payload["comparison_id"]), str(payload["format"])
    with Session() as db:
        response = await export_comparison(comparison_id, format_name, payload.get("template_id"), db)
        data, media_type = response.body, response.media_type or "application/octet-stream"
    ensure_not_cancelled(job_id)
    export_root = ROOT / "data" / "exports"
    export_root.mkdir(parents=True, exist_ok=True)
    filename = f"comparison-{comparison_id}-job-{job_id}.{format_name}"
    path, temporary = export_root / filename, export_root / f".{job_id}.tmp"
    temporary.write_bytes(data)
    os.replace(temporary, path)
    try:
        ensure_not_cancelled(job_id)
        with Session.begin() as db:
            job = db.get(BackgroundJob, job_id)
            job.payload = {**(job.payload or {}), "artifact": filename, "media_type": media_type, "bytes": len(data)}
        set_progress(job_id, 1, 1)
    except Exception:
        path.unlink(missing_ok=True)
        raise


async def _tag_paper_job(payload: dict, job_id: int):
    await tag_paper(int(payload["paper_id"]), job_id)


async def _preferred_generation_job(payload: dict, job_id: int):
    """Draft, screen by reference paper/unit, and retry repeat candidates up to twice."""
    paper_id, reference_ids, slots = int(payload["paper_id"]), payload["reference_ids"], payload["slots"]
    total = sum(int(slot["count"]) for slot in slots)
    with Session() as db:
        paper = db.get(Paper, paper_id)
        if not paper:
            raise RuntimeError("Draft paper was deleted")
        subject = subject_or_404(db, paper.subject_id)
        nodes = {node["id"]: node for node in subject.tree.get("nodes", [])}
        references = [db.get(Question, ref_id) for ref_id in reference_ids]
        if any(ref is None for ref in references):
            raise RuntimeError("A selected reference question was deleted")
        blocks = {}
        for ref in references:
            blocks.setdefault((ref.unit, ref.paper_id), []).append(SimpleNamespace(id=ref.id, revision=ref.revision, paper_id=ref.paper_id, number=ref.number, text=ref.text))
        paper.generation = {**(paper.generation or {}), "status": "running", "total": total}
        db.commit()
    completed = 0
    for slot_index, slot in enumerate(slots, 1):
        unit_blocks = [refs for (unit, _), refs in blocks.items() if unit == slot["unit"]]
        topics = [nodes[topic_id] for topic_id in slot["topic_ids"] if topic_id in nodes]
        context = [{"id": ref.id, "number": ref.number, "text": ref.text} for refs in unit_blocks for ref in refs]
        for index in range(1, int(slot["count"]) + 1):
            ensure_not_cancelled(job_id)
            number = slot.get("number") if slot.get("number") and int(slot["count"]) == 1 else f"G{slot_index}.{index}"
            with Session() as db:
                paper = db.get(Paper, paper_id)
                if any(q.number == number for q in paper.questions):
                    completed += 1
                    continue
                existing = [q.text for q in paper.questions]
            attempts, best = [], None
            for attempt in range(3):
                ensure_not_cancelled(job_id)
                with Session.begin() as db:
                    paper = db.get(Paper, paper_id)
                    paper.generation = {**paper.generation, "stage": f"Drafting {number} · attempt {attempt + 1}/3", "progress": completed}
                slot_input = {**slot, "count": 1, "constraints": (slot.get("constraints", "") + " Aim for B or C against past questions: change assessed scope or instruction depth, or assess another selected syllabus topic. Avoid repeats.").strip()}
                try:
                    candidate = (await draft_slot(slot_input, topics, context, existing + [item["text"] for item in attempts]))[0]
                except ValueError:
                    if attempt == 2 and best is None:
                        raise
                    continue
                candidate_question = SimpleNamespace(number=number, text=candidate["text"])
                checks = []
                for refs in unit_blocks:
                    ensure_not_cancelled(job_id)
                    checks.append(await classify_block(candidate_question, refs, slot["unit"]))
                ensure_not_cancelled(job_id)
                candidate["checks"] = checks
                attempts.append(candidate)
                score = (sum(check["label"] in ("B", "C") for check in checks), -sum(check["label"] == "A" for check in checks))
                if best is None or score > best[0]:
                    best = (score, candidate)
                if checks and all(check["label"] in ("B", "C") for check in checks):
                    break
            chosen = best[1]
            with Session.begin() as db:
                paper = db.get(Paper, paper_id)
                db.add(Question(paper_id=paper_id, number=number, text=chosen["text"], marks=slot["marks"], unit=slot["unit"], topic_ids=chosen["topic_ids"], confirmed=True))
                screens = dict(paper.generation.get("screens") or {})
                screens[number] = {"checks": chosen["checks"], "attempts": len(attempts)}
                completed += 1
                paper.generation = {**paper.generation, "screens": screens, "progress": completed, "stage": f"Checked {completed}/{total} questions"}
            set_progress(job_id, completed, total)
    ensure_not_cancelled(job_id)
    with Session.begin() as db:
        paper = db.get(Paper, paper_id)
        subject = subject_or_404(db, paper.subject_id)
        if paper.generation.get("comparison_id"):
            return
        screens = paper.generation.get("screens", {})
        comparison = Comparison(subject_id=paper.subject_id, title=f"Draft check · {paper.title}", target_ids=[q.id for q in paper.questions], reference_ids=reference_ids, generated_paper_id=paper.id, status="complete", template_snapshot={"strategy": BLOCK_VERSION})
        db.add(comparison)
        db.flush()
        counts = {label: 0 for label in ("A", "B", "C", "review")}
        for question in paper.questions:
            for result in screens[question.number]["checks"]:
                block = next(item for item in result["evidence"] if item["id"] == "block")
                refs = [db.get(Question, entry["id"]) for entry in block["questions"]]
                if any(ref is None for ref in refs):
                    raise RuntimeError("A selected reference was deleted during checking")
                counts[result["label"]] += 1
                db.add(PairResult(comparison_id=comparison.id, target_id=question.id, reference_id=refs[0].id, cache_key=block_key(question, refs, block["unit"], subject.tree_version), **result))
        comparison.progress = comparison.total = sum(counts.values())
        paper.generation = {**paper.generation, "status": "checked", "comparison_id": comparison.id, "outcome_counts": counts, "stage": "Finished", "claim": "B/C preference is based on model checks; review the resulting questions."}
        db.add(AuditEvent(action="question_generation_complete", details={"paper_id": paper.id, "comparison_id": comparison.id, "outcome_counts": counts, "prefer_bc": True}))


async def _insem_batch_generation_job(payload: dict, job_id: int):
    paper_id, slots, reference_ids = int(payload["paper_id"]), payload["slots"], payload["reference_ids"]
    total = sum(slot["count"] for slot in slots)
    with Session() as db:
        paper = db.get(Paper, paper_id)
        if not paper:
            raise RuntimeError("Draft paper was deleted")
        subject = subject_or_404(db, paper.subject_id)
        nodes = {node["id"]: node for node in subject.tree.get("nodes", [])}
        references = [db.get(Question, qid) for qid in reference_ids]
        if any(q is None for q in references):
            raise RuntimeError("A selected past question was deleted")
        by_unit = {unit: [{"number": q.number, "text": q.text} for q in references if q.unit == unit] for unit in (1, 2)}
        completed_initial = len(paper.questions)
        paper.generation = {**paper.generation, "status": "running", "total": total, "progress": completed_initial, "error": None}
        db.commit()
    set_progress(job_id, completed_initial, total)
    for unit in (1, 2):
        unit_slots = [slot for slot in slots if slot["unit"] == unit]
        if not unit_slots:
            continue
        with Session() as db:
            completed_numbers = {q.number for q in db.get(Paper, paper_id).questions}
        if all(slot["number"] in completed_numbers for slot in unit_slots):
            continue
        unit_topics = [nodes[tid] for tid in dict.fromkeys(tid for slot in unit_slots for tid in slot["topic_ids"]) if tid in nodes]
        try:
            hints = await semantic_topic_hints(unit_topics, by_unit[unit])
        except ModelUnavailable as exc:
            logger.warning("Semantic context unavailable for generation: %s", exc)
            with Session.begin() as db:
                current = db.get(Paper, paper_id)
                current.generation = {**current.generation, "semantic_warning": "Semantic topic matching was unavailable; questions were generated without it."}
            hints = {}
        ensure_not_cancelled(job_id)
        for batch_index in range(0, len(unit_slots), 5):
            ensure_not_cancelled(job_id)
            with Session() as db:
                paper = db.get(Paper, paper_id)
                existing_numbers = {q.number for q in paper.questions}
                existing = [q.text for q in paper.questions if q.unit == unit]
            pending = [slot for slot in unit_slots[batch_index:batch_index+5] if slot["number"] not in existing_numbers]
            if not pending:
                continue
            if not by_unit[unit]:
                raise RuntimeError(f"The selected past paper has no tagged Unit {unit} questions")
            with Session.begin() as db:
                paper = db.get(Paper, paper_id)
                paper.generation = {**paper.generation, "stage": f"Drafting Unit {unit} batch {batch_index//5+1}/2"}
            topics = [{**nodes[tid], "related_past_questions": hints.get(tid, [])} for tid in dict.fromkeys(tid for slot in pending for tid in slot["topic_ids"]) if tid in nodes]
            candidates = await draft_batch_with_fallback(pending, topics, by_unit[unit], existing)
            ensure_not_cancelled(job_id)
            with Session.begin() as db:
                paper = db.get(Paper, paper_id)
                paper.generation = {**paper.generation, "stage": f"Checking Unit {unit} batch {batch_index//5+1}/2"}
            screening = await screen_batch([{"number": number, **candidate} for number, candidate in candidates.items()], by_unit[unit], existing)
            repeats = [slot for slot in pending if screening[slot["number"]]["status"] == "too_close"]
            if repeats:
                ensure_not_cancelled(job_id)
                with Session.begin() as db:
                    paper = db.get(Paper, paper_id)
                    paper.generation = {**paper.generation, "stage": f"Replacing {len(repeats)} repeated question(s) in Unit {unit}"}
                other = existing + [item["text"] for item in candidates.values()]
                try:
                    replacement = await draft_batch_with_fallback(repeats, topics, by_unit[unit], other)
                    replacement_screen = await screen_batch([{"number": number, **candidate} for number, candidate in replacement.items()], by_unit[unit], other)
                    candidates.update(replacement)
                    screening.update(replacement_screen)
                except ValueError:
                    pass  # Preserve the original repeat and its reason for faculty review.
            ensure_not_cancelled(job_id)
            with Session.begin() as db:
                paper = db.get(Paper, paper_id)
                screens = dict(paper.generation.get("quick_screens") or {})
                for slot in pending:
                    number = slot["number"]
                    candidate = candidates[number]
                    db.add(Question(paper_id=paper_id, number=number, text=candidate["text"], marks=slot["marks"], unit=unit, topic_ids=candidate["topic_ids"], confirmed=True))
                    screens[number] = screening[number]
                completed = len(existing_numbers) + len(pending)
                paper.generation = {**paper.generation, "quick_screens": screens, "progress": completed, "stage": f"Checked {completed}/{total} questions"}
            set_progress(job_id, completed, total)
    ensure_not_cancelled(job_id)
    with Session.begin() as db:
        paper = db.get(Paper, paper_id)
        screens = paper.generation.get("quick_screens") or {}
        counts = {status: sum(item.get("status") == status for item in screens.values()) for status in ("different", "too_close", "review")}
        paper.generation = {**paper.generation, "status": "checked", "stage": "Finished", "quick_counts": counts,
            "claim": "Quick screen against one past paper; faculty review is required."}
        db.add(AuditEvent(action="question_generation_complete", details={"paper_id": paper.id, "quick_counts": counts, "strategy": "insem_quick_v1"}))


async def _question_generation_job(payload: dict, job_id: int):
    with Session() as db:
        paper = db.get(Paper, int(payload["paper_id"]))
        quick_screen = paper and (paper.generation or {}).get("screen_strategy") == "insem_quick_v1"
    if quick_screen:
        return await _insem_batch_generation_job(payload, job_id)
    if payload.get("prefer_bc"):
        return await _preferred_generation_job(payload, job_id)
    paper_id = int(payload["paper_id"])
    reference_ids = [int(value) for value in payload["reference_ids"]]
    slots = payload["slots"]
    with Session() as db:
        paper = db.get(Paper, paper_id)
        if not paper:
            raise RuntimeError("Draft paper was deleted before generation")
        subject = subject_or_404(db, paper.subject_id)
        reference_questions = {q.id: q for q in db.scalars(select(Question).where(Question.id.in_(reference_ids)))}
        if len(reference_questions) != len(set(reference_ids)):
            raise RuntimeError("A selected reference question was deleted before generation")
        tree_nodes = {node["id"]: node for node in subject.tree.get("nodes", [])}
        paper.generation = {**(paper.generation or {}), "status": "running", "prompt_version": GENERATION_PROMPT_VERSION}
        db.commit()
    reference_context_data = []
    reference_units = {}
    with Session() as db:
        for question_id in reference_ids:
            question = db.get(Question, question_id)
            if question:
                reference_units[question.id] = question.unit
                reference_context_data.append({"id": question.id, "number": question.number, "paper": question.paper.title, "unit": question.unit, "text": question.text[:1200]})
    completed = 0
    for slot_index, slot in enumerate(slots, 1):
        ensure_not_cancelled(job_id)
        with Session() as db:
            paper = db.get(Paper, paper_id)
            if not paper:
                raise RuntimeError("Draft paper was deleted during generation")
            existing_slot = {question.number: question for question in paper.questions if question.number.startswith(f"G{slot_index}.")}
            required_count = int(slot["count"])
            missing_numbers = [index for index in range(1, required_count + 1) if f"G{slot_index}.{index}" not in existing_slot]
            old_candidate_texts = [question.text for question in paper.questions]
            topic_context = [tree_nodes[topic_id] for topic_id in slot["topic_ids"] if topic_id in tree_nodes]
            reference_context = [{key: value for key, value in question.items() if key != "unit"} for question in reference_context_data if question["unit"] == slot["unit"]][:12]
        if missing_numbers:
            generation_slot = {**slot, "count": len(missing_numbers)}
            generated = await draft_slot(generation_slot, topic_context, reference_context, old_candidate_texts)
            ensure_not_cancelled(job_id)
            with Session.begin() as db:
                paper = db.get(Paper, paper_id)
                if not paper:
                    raise RuntimeError("Draft paper was deleted while saving generated questions")
                for index, item in zip(missing_numbers, generated):
                    db.add(Question(paper_id=paper_id, number=f"G{slot_index}.{index}", text=item["text"], marks=slot["marks"], unit=slot["unit"], topic_ids=item["topic_ids"], confirmed=True))
                paper.generation = {**(paper.generation or {}), "status": "running", "progress": completed + len(missing_numbers), "total": sum(int(item["count"]) for item in slots)}
        completed += int(slot["count"])
        set_progress(job_id, completed, sum(int(item["count"]) for item in slots))
    ensure_not_cancelled(job_id)
    with Session.begin() as db:
        paper = db.get(Paper, paper_id)
        if not paper:
            raise RuntimeError("Draft paper was deleted before redundancy screening")
        prior_comparison_id = (paper.generation or {}).get("comparison_id")
        prior_comparison = db.get(Comparison, prior_comparison_id) if prior_comparison_id else None
        if prior_comparison:
            return
        target_ids = [question.id for question in paper.questions]
        if not target_ids:
            raise RuntimeError("Question generation produced no candidates")
        comparison = Comparison(subject_id=paper.subject_id, title=f"Redundancy screen · {paper.title}", target_ids=target_ids, reference_ids=reference_ids, generated_paper_id=paper.id, status="queued")
        db.add(comparison)
        db.flush()
        job = enqueue_job(db, "comparison", {"comparison_id": comparison.id}, total=sum(1 for target in paper.questions for ref_id in reference_ids if target.unit == reference_units.get(ref_id) and target.id != ref_id))
        paper.generation = {**(paper.generation or {}), "status": "checking", "comparison_id": comparison.id, "generated_question_ids": target_ids, "progress": completed, "total": completed}
        db.add(AuditEvent(action="question_generation_complete", details={"paper_id": paper.id, "comparison_id": comparison.id, "candidate_ids": target_ids, "prompt_version": GENERATION_PROMPT_VERSION}))


async def _material_enrichment_job(payload: dict, job_id: int):
    material_id = int(payload["material_id"])
    with Session() as db:
        material = db.get(Material, material_id)
        if not material:
            raise RuntimeError("Teaching material was deleted before processing")
        subject = subject_or_404(db, material.subject_id)
        start_tree_version = subject.tree_version
        prior_enrichment = material.enrichment or {}
        section_cache = prior_enrichment.get("_section_cache", {})
        if not isinstance(section_cache, dict):
            section_cache = {}
        material.enrichment = {**prior_enrichment, "status": "running", "progress": 0, "total": 0, "prompt_version": MATERIAL_PROMPT_VERSION, "error": None, "activity": [], "activity_seq": 0}
        db.commit()
        topics = [{
            "id": node["id"], "label": node["label"], "unit": unit_for_topic(node["id"]),
            "description": node.get("description", ""), "source_text": node.get("source_text", ""),
            "scope_terms": node.get("scope_terms", []), "synonyms": node.get("synonyms", []),
            "material_notes": node.get("material_notes", []),
        } for node in subject.tree["nodes"] if node.get("kind") == "topic"]
        filename, extracted_text = material.filename, material.extracted_text
    def checkpoint_chunk(cache_key: str, value: dict) -> None:
        with Session.begin() as db:
            current = db.get(Material, material_id)
            if not current:
                raise RuntimeError("Teaching material was deleted while mapping")
            current_enrichment = current.enrichment or {}
            cache = current_enrichment.get("_section_cache", {})
            if not isinstance(cache, dict):
                cache = {}
            cache[cache_key] = value
            if len(cache) > 500:
                cache = dict(list(cache.items())[-500:])
            current.enrichment = {**current_enrichment, "_section_cache": cache}
    def material_activity(event: str, payload: dict) -> None:
        with Session.begin() as db:
            current = db.get(Material, material_id)
            if not current:
                return
            state = dict(current.enrichment or {})
            sequence = int(state.get("activity_seq", 0)) + 1
            activity = list(state.get("activity", []))
            safe_payload = {key: value for key, value in payload.items() if key != "content"}
            if "text" in safe_payload:
                safe_payload["text"] = str(safe_payload["text"])[:800]
            activity.append({"id": sequence, "agent": "mapper", "event": event, "payload": safe_payload, "created_at": datetime.now(timezone.utc).isoformat()})
            state["activity_seq"] = sequence
            state["activity"] = activity[-300:]
            current.enrichment = state
    set_progress(job_id, 0, len(topics))
    def material_progress(current: int, total: int) -> None:
        ensure_not_cancelled(job_id)
        set_progress(job_id, current, total)

    result = await map_material(extracted_text, topics, filename, progress=material_progress, activity=material_activity)
    ensure_not_cancelled(job_id)
    with Session.begin() as db:
        material = db.get(Material, material_id)
        if not material:
            raise RuntimeError("Teaching material was deleted during processing")
        subject = subject_or_404(db, material.subject_id)
        if subject.tree_version != start_tree_version:
            raise RuntimeError("Syllabus changed while this material was being mapped; retry against the active tree")
        valid_topics = {node["id"] for node in subject.tree["nodes"] if node.get("kind") == "topic"}
        unknown = set(result["proposals"]) - valid_topics
        if unknown:
            raise RuntimeError("Syllabus topics changed during mapping; retry the material mapping")
        material.proposals = {topic_id: {**proposal, "included": True} for topic_id, proposal in result["proposals"].items()}
        material.enrichment = {**{key: value for key, value in result.items() if key != "proposals"}, "status": "review_ready", "source_tree_version": start_tree_version, "activity": (material.enrichment or {}).get("activity", []), "activity_seq": (material.enrichment or {}).get("activity_seq", 0)}
        db.add(AuditEvent(action="material_mapping_prepared", details={"material_id": material_id, "topics": list(material.proposals), "sections_processed": result["sections_processed"], "source_tree_version": start_tree_version}))


async def _syllabus_import_job(payload: dict, job_id: int):
    source_id = int(payload["source_id"])
    with Session() as db:
        source = db.get(SyllabusImport, source_id)
        if not source:
            raise RuntimeError("Syllabus upload was deleted before processing")
        already_active = source.status == "active"
        if not already_active:
            source.status, source.error = "processing", None
            filename, text = source.filename, source.extracted_text
            db.commit()
    if already_active:
        set_progress(job_id, 1, 1)
        return
    try:
        tree = await build_tree(text, filename)
        ensure_not_cancelled(job_id)
    except Exception as exc:
        with Session.begin() as db:
            source = db.get(SyllabusImport, source_id)
            if source:
                source.status = "failed"
                source.error = str(exc)[:2000]
        raise
    with Session.begin() as db:
        source = db.get(SyllabusImport, source_id)
        if not source:
            raise RuntimeError("Syllabus upload was deleted during processing")
        subject = activate_tree(db, tree, f"syllabus upload: {filename}")
        source.subject_id = subject.id
        source.status, source.error = "active", None
        uploader = db.get(User, source.submitted_by) if source.submitted_by else None
        if uploader and uploader.role == "teacher" and subject.id not in uploader.subject_ids:
            uploader.subject_ids = [*uploader.subject_ids, subject.id]
        db.add(AuditEvent(action="syllabus_tree_activate", details={"source_id": source.id, "subject_id": subject.id, "tree_version": subject.tree_version, "prompt_version": SYLLABUS_PROMPT_VERSION}))
    set_progress(job_id, 1, 1)


JOB_HANDLERS = {
    "index_paper": index_paper,
    "comparison": _comparison_job,
    "comparison_export": _comparison_export_job,
    "tag_paper": _tag_paper_job,
    "question_generation": _question_generation_job,
    "material_enrichment": _material_enrichment_job,
    "syllabus_import": _syllabus_import_job,
}


DIST = ROOT / "frontend" / "dist"
if DIST.exists():
    app.mount("/assets", StaticFiles(directory=DIST / "assets"), name="assets")

    @app.get("/")
    async def workspace():
        return FileResponse(DIST / "index.html")
