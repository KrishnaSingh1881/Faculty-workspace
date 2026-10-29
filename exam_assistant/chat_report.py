"""Excel report for comparisons saved in a conversation."""
import io
import re
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def _safe(value):
    value = str(value or '')
    return "'" + value if value.startswith(('=', '+', '-', '@')) else value


def _key(number):
    return re.sub(r'[^a-z0-9]', '', str(number).casefold()).removeprefix('q')


def build_chat_report(comparisons, db, question_model, result_model, report_row_values):
    from sqlalchemy import select
    book = Workbook()
    book.remove(book.active)
    groups = defaultdict(list)
    details = []
    for comparison in comparisons:
        results = db.scalars(select(result_model).where(result_model.comparison_id == comparison.id).order_by(result_model.id)).all()
        for result in results:
            target, reference = db.get(question_model, result.target_id), db.get(question_model, result.reference_id)
            if not target or not reference:
                continue
            groups[target.paper_id].append((comparison, result, target, reference, report_row_values(comparison, result, db)))
            details.append((comparison, result, target, reference, report_row_values(comparison, result, db)))
    border = Border(left=Side(style='thin', color='B8B8B8'), right=Side(style='thin', color='B8B8B8'), top=Side(style='thin', color='B8B8B8'), bottom=Side(style='thin', color='B8B8B8'))
    fills = {'A':'D9EAD3','B':'FFE599','C':'F4CCCC','review':'E7E6E6'}
    for index, (base_id, entries) in enumerate(groups.items(), 1):
        sheet = book.create_sheet(f'Comparison {index}')
        sheet.freeze_panes = 'C3'
        base_title = entries[0][2].paper.title
        reference_papers = list(dict.fromkeys(entry[3].paper_id for entry in entries))
        base_questions = list(dict.fromkeys(entry[2].id for entry in entries))
        base_by_id = {entry[2].id: entry[2] for entry in entries}
        ref_titles = {entry[3].paper_id: entry[3].paper.title for entry in entries}
        sheet.append(['Q. No', base_title] + [value for pid in reference_papers for value in ('Q. No', ref_titles[pid], 'Classification')])
        sheet.append(['', 'Rows align by question number. Classifications use the saved comparison scope, which may cover the full unit block.'])
        sheet.column_dimensions['A'].width = 10
        sheet.column_dimensions['B'].width = 55
        for offset, _ in enumerate(reference_papers):
            start = 3 + offset * 3
            sheet.column_dimensions[get_column_letter(start)].width = 10
            sheet.column_dimensions[get_column_letter(start + 1)].width = 55
            sheet.column_dimensions[get_column_letter(start + 2)].width = 17
        for question_id in base_questions:
            target = base_by_id[question_id]
            row = [target.number, next(entry[4]['base_question'] for entry in entries if entry[2].id == question_id)]
            for paper_id in reference_papers:
                matching = [entry for entry in entries if entry[2].id == question_id and entry[3].paper_id == paper_id]
                references = []
                for entry in matching:
                    block = next((item for item in entry[1].evidence if isinstance(item, dict) and item.get('id') == 'block'), None)
                    if block:
                        references.extend(q for q in block.get('questions', []) if _key(q['number']) == _key(target.number))
                    elif _key(entry[3].number) == _key(target.number):
                        references.append({'number': entry[3].number, 'text': entry[4]['reference_question']})
                unique = {str(q['number']): q['text'] for q in references}
                labels = list(dict.fromkeys(entry[1].label for entry in matching))
                row.extend([', '.join(unique), '\n'.join(unique.values()), ', '.join(labels) if labels else '—'])
            sheet.append([_safe(value) for value in row])
            sheet.row_dimensions[sheet.max_row].height = min(409, max(35, 16 * max(2, max((len(str(v)) // 46 + str(v).count('\n') + 1 for v in row), default=2))))
        summary_row = sheet.max_row + 2
        sheet.cell(summary_row, 2, 'Classification totals')
        for offset, paper_id in enumerate(reference_papers):
            col = 5 + offset * 3
            sheet.cell(summary_row, col, ref_titles[paper_id])
            for n, label in enumerate(('A', 'B', 'C', 'review'), 1):
                sheet.cell(summary_row + n, 2, 'Needs review' if label == 'review' else f'Count of Type "{label}" Questions')
                sheet.cell(summary_row + n, col, sum(entry[1].label == label for entry in entries if entry[3].paper_id == paper_id))
        for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, max_col=2 + 3*len(reference_papers)):
            for cell in row:
                cell.font = Font(name='Times New Roman', size=11, bold=cell.row == 1 or cell.row >= summary_row)
                cell.alignment = Alignment(vertical='center', wrap_text=True, horizontal='center' if cell.column == 1 or (cell.column >= 3 and (cell.column - 3) % 3 != 1) else 'left')
                if cell.row != 2:
                    cell.border = border
                if cell.row == 1:
                    cell.fill = PatternFill('solid', fgColor='E9E9E9')
                elif cell.row >= summary_row:
                    cell.fill = PatternFill('solid', fgColor='F4CCCC')
                elif cell.column >= 5 and (cell.column - 5) % 3 == 0 and cell.value in fills:
                    cell.fill = PatternFill('solid', fgColor=fills[cell.value])
        sheet.row_dimensions[1].height = 30
        sheet.row_dimensions[2].height = 45
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
    sheet = book.create_sheet('Reasons')
    sheet.append(['Comparison', 'Base paper', 'Q. No', 'Base question', 'Past paper', 'Classification', 'Reason', 'Past-paper unit block', 'Status', 'Stale'])
    for comparison, result, target, reference, values in details:
        sheet.append([_safe(v) for v in (comparison.title, target.paper.title, target.number, values['base_question'], reference.paper.title, result.label, result.reason, values['reference_question'], comparison.status, values['stale'])])
    widths = [25, 23, 10, 55, 23, 16, 60, 65, 15, 10]
    for i, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    sheet.freeze_panes = 'D2'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.orientation = 'landscape'
    sheet.print_title_rows = '1:1'
    sheet.auto_filter.ref = sheet.dimensions
    for row in sheet.iter_rows():
        for cell in row:
            cell.font = Font(name='Times New Roman', size=11, bold=cell.row == 1)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            cell.border = border
            if cell.row == 1:
                cell.fill = PatternFill('solid', fgColor='E9E9E9')
            elif cell.column == 6 and cell.value in fills:
                cell.fill = PatternFill('solid', fgColor=fills[cell.value])
        if row[0].row > 1:
            sheet.row_dimensions[row[0].row].height = min(409, max(35, 16 * (len(str(row[7].value or '')) // 56 + str(row[7].value or '').count('\n') + 1)))
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
